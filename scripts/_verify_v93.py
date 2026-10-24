# -*- coding: utf-8 -*-
import zipfile, re
z = zipfile.ZipFile(r'C:/Users/LENOVO/Desktop/余慧玲9月月度汇报_v9_3.pptx')
xml = z.read('ppt/charts/chart5.xml').decode('utf-8','ignore')
sers = re.findall(r'<c:ser>.*?</c:ser>', xml, re.S)
for ser in sers:
    nm = re.search(r'<c:tx>.*?<c:v>([^<]+)</c:v>', ser, re.S).group(1)
    dl = re.search(r'<c:dLbls>.*?</c:dLbls>', ser, re.S).group(0)
    sv = re.search(r'<c:showVal val="(\d)"/>', dl).group(1)
    pos = re.search(r'<c:dLblPos val="(\w+)"/>', dl).group(1)
    print(f"series {nm:>6} showVal={sv} pos={pos}")
# full package validation
import posixpath
names = set(z.namelist())
ct = z.read('[Content_Types].xml').decode('utf-8','ignore')
overrides = set(re.findall(r'PartName="([^"]+)"', ct))
defaults = set(re.findall(r'Extension="([^"]+)"', ct))
probs = []
for o in overrides:
    if o.lstrip('/') not in names: probs.append('CT override missing: '+o)
for n in names:
    if n == '[Content_Types].xml': continue
    ext = n.rsplit('.',1)[-1].lower()
    if ('/'+n) not in overrides and ext not in defaults: probs.append('part not declared: '+n)
    if n.endswith('.rels'):
        base = posixpath.dirname(posixpath.dirname(n))
        for rid, tgt in re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read(n).decode('utf-8','ignore')):
            if tgt.startswith('http'): continue
            full = tgt.lstrip('/') if tgt.startswith('/') else posixpath.normpath(posixpath.join(base, tgt))
            if full not in names: probs.append(f'{n}: {rid}->{tgt} MISSING {full}')
print("PACKAGE PROBLEMS:", probs if probs else "none")
# data check
from pptx import Presentation
prs = Presentation(r'C:/Users/LENOVO/Desktop/余慧玲9月月度汇报_v9_3.pptx')
for sh in prs.slides[4].shapes:
    if sh.has_chart:
        for se in sh.chart.plots[0].series:
            print(se.name, "->", [round(v) for v in se.values])
