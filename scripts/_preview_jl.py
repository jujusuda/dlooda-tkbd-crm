# -*- coding: utf-8 -*-
"""按真实几何生成第 15 页 SVG 预览"""
from pptx import Presentation
from pptx.oxml.ns import qn

OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_达人建联修改.pptx"
prs = Presentation(OUT)
s = prs.slides[14]
SW, SH = prs.slide_width, prs.slide_height
EMU = 914400.0

def fill_of(sh):
    try:
        if sh.fill.type is not None and sh.fill.type == 1:
            return "#" + str(sh.fill.fore_color.rgb)
    except Exception:
        pass
    try:
        spPr = sh._element.find(qn('p:spPr'))
        if spPr is None:
            spPr = sh._element.find(qn('p:spPr'))
        sf = spPr.find(qn('a:solidFill'))
        if sf is not None:
            c = sf.find(qn('a:srgbClr'))
            if c is not None:
                return "#" + c.get('val')
    except Exception:
        pass
    return None

def geom_points(sh):
    spPr = sh._element.find(qn('p:spPr'))
    if spPr is None:
        return None
    cg = spPr.find(qn('a:custGeom'))
    if cg is None:
        return None
    path = cg.find(qn('a:pathLst')).find(qn('a:path'))
    pts = []
    for node in path:
        if node.tag in (qn('a:moveTo'), qn('a:lnTo')):
            pt = node.find(qn('a:pt'))
            pts.append((int(pt.get('x')), int(pt.get('y'))))
    return pts

def txt_lines(sh):
    out = []
    if not sh.has_text_frame:
        return out
    for p in sh.text_frame.paragraphs:
        t = p.text
        if not t.strip():
            continue
        sz, col, bold = 12, "#3D3D3D", False
        for r in p.runs:
            if r.font.size:
                sz = r.font.size.pt
            if r.font.color and r.font.color.type is not None:
                col = "#" + str(r.font.color.rgb)
            if r.font.bold:
                bold = True
            break
        al = str(p.alignment)
        out.append((t, sz, col, bold, al))
    return out

def anchor_y(sh, n):
    """粗略：整体垂直居中时逐行排布"""
    return None

parts = []
for sh in s.shapes:
    L, T, W, H = sh.left, sh.top, sh.width, sh.height
    f = fill_of(sh)
    pts = geom_points(sh)
    if pts:
        poly = " ".join(f"{L+x},{T+y}" for x, y in pts)
        parts.append(f'<polygon points="{poly}" fill="{f or "none"}"/>')
    else:
        # 圆角/椭圆/矩形
        is_ell = sh._element.find(qn('p:spPr')) is not None and \
                 sh._element.find(qn('p:spPr')).find(qn('a:prstGeom')) is not None and \
                 sh._element.find(qn('p:spPr')).find(qn('a:prstGeom')).get('prst') == 'ellipse'
        if f is None:
            continue
        if is_ell:
            parts.append(f'<ellipse cx="{L+W/2}" cy="{T+H/2}" rx="{W/2}" ry="{H/2}" fill="{f}"/>')
        else:
            r = min(73152, W / 12)
            parts.append(f'<rect x="{L}" y="{T}" width="{W}" height="{H}" rx="{r}" fill="{f}"/>')
    lines = txt_lines(sh)
    n = max(1, len(lines))
    for i, (t, sz, col, bold, al) in enumerate(lines):
        px = L if 'LEFT' in al else (L + W / 2 if 'CENTER' in al else L + W)
        anc = "start" if 'LEFT' in al else ("middle" if 'CENTER' in al else "end")
        py = T + H * (i + 0.5) / n
        parts.append(f'<text x="{px}" y="{py}" font-size="{sz*12700}" fill="{col}" '
                     f'text-anchor="{anc}" dominant-baseline="middle" '
                     f'font-weight="{"700" if bold else "400"}">{t}</text>')

# 表格
for sh in s.shapes:
    if not sh.has_table:
        continue
    tb = sh.table
    y = sh.top
    for ri, row in enumerate(tb.rows):
        x = sh.left
        h = row.height
        for ci, cell in enumerate(row.cells):
            w = tb.columns[ci].width
            f = None
            tcPr = cell._tc.find(qn('a:tcPr'))
            if tcPr is not None:
                sf = tcPr.find(qn('a:solidFill'))
                if sf is not None:
                    c = sf.find(qn('a:srgbClr'))
                    if c is not None:
                        f = "#" + c.get('val')
            parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{f or "#FFFFFF"}" stroke="#E6E1E3" stroke-width="9144"/>')
            txt = cell.text
            col, bold, sz = "#3D3D3D", False, 12.5
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.color and r.font.color.type is not None:
                        col = "#" + str(r.font.color.rgb)
                    bold = bool(r.font.bold)
                    if r.font.size:
                        sz = r.font.size.pt
                    break
                break
            parts.append(f'<text x="{x+w/2}" y="{y+h/2}" font-size="{sz*12700}" fill="{col}" '
                         f'text-anchor="middle" dominant-baseline="middle" '
                         f'font-weight="{"700" if bold else "400"}">{txt}</text>')
            x += w
        y += h

W = int(SW / EMU * 100)
H = int(SH / EMU * 100)
sc = 680.0 / SW
HH = SH * sc
inner = ("<g transform='scale(%f)'>" % sc) + "".join(parts) + "</g>"
svg = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 680 {HH:.1f}" width="100%">'
       f'<rect width="680" height="{HH:.1f}" fill="#FFFFFF"/>' + inner + '</svg>')
p = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\_preview_jl.svg"
open(p, 'w', encoding='utf-8').write(svg)
print("ok", p)
