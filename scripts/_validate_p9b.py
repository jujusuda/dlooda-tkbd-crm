# -*- coding: utf-8 -*-
import zipfile, posixpath
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(9).pptx"
OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_v9_1(9)_P9修改.pptx"
ok = lambda b: "OK " if b else "FAIL"
def c(t): return qn('c:' + t)

z = zipfile.ZipFile(OUT)
print(ok(z.testzip() is None), "zip integrity")
bad = []
for n in z.namelist():
    if n.endswith('.rels'):
        base = posixpath.dirname(posixpath.dirname(n))
        for rel in etree.fromstring(z.read(n)):
            if rel.get('TargetMode') == 'External':
                continue
            if posixpath.normpath(posixpath.join(base, rel.get('Target'))) not in z.namelist():
                bad.append((n, rel.get('Target')))
print(ok(not bad), "rels resolve", bad[:4])
for n in z.namelist():
    if n.endswith('.xml') or n.endswith('.rels'):
        etree.fromstring(z.read(n))
print("OK  all xml well-formed")

src, out = Presentation(SRC), Presentation(OUT)
print(ok(len(src.slides) == len(out.slides)), f"slides {len(out.slides)}")
diff = [i + 1 for i in range(len(out.slides)) if i != 8 and
        etree.tostring(out.slides[i]._element) != etree.tostring(src.slides[i]._element)]
print(ok(not diff), "other slides untouched", diff)

sl = out.slides[8]
pics = [sh for sh in sl.shapes if sh.shape_type == 13]
print(ok(not pics), f"pictures on slide 9 -> {len(pics)} (must be 0)")
types = {}
for sh in sl.shapes:
    types[sh.shape_type] = types.get(sh.shape_type, 0) + 1
print("shape types:", {str(k): v for k, v in types.items()})
print("\nshapes:")
for i, sh in enumerate(sl.shapes):
    txt = (sh.text_frame.text if sh.has_text_frame else "").replace("\n", " | ")
    print(f"[{i:2d}] {sh.name:<20} L={sh.left:>8} T={sh.top:>8} W={sh.width:>8} H={sh.height:>8} {txt[:58]}")

# --- combo chart
ch = [sh for sh in sl.shapes if sh.has_chart and sh.width > 5000000][0].chart
pa = ch._chartSpace.plotArea
print("\nplotArea:", [e.tag.split('}')[1] for e in pa])
catId = pa.find(c('catAx')).find(c('axId')).get('val')
for va in pa.findall(c('valAx')):
    ttl = va.find(c('title'))
    t = None
    if ttl is not None:
        t = ttl.find(c('tx')).find(c('rich')).find(qn('a:p')).find(qn('a:r')).find(qn('a:t')).text
    nf = va.find(c('numFmt'))
    print(f"valAx id={va.find(c('axId')).get('val')} pos={va.find(c('axPos')).get('val')} "
          f"numFmt={nf.get('formatCode')} title={t!r} grid={va.find(c('majorGridlines')) is not None}")
for nm in ('barChart', 'lineChart'):
    p = pa.find(c(nm))
    print(f"{nm} axIds={[x.get('val') for x in p.findall(c('axId'))]}")
    for s in p.findall(c('ser')):
        name = s.find(c('tx')).find(c('strRef')).find(c('strCache')).find(c('pt')).find(c('v')).text
        vals = [x.find(c('v')).text for x in
                s.find(c('val')).find(c('numRef')).find(c('numCache')).findall(c('pt'))]
        dl = s.find(c('dLbls'))
        lab = (dl.find(c('dLblPos')).get('val'), dl.find(c('numFmt')).get('formatCode'),
               'whitefill' if dl.find(c('spPr')).find(qn('a:solidFill')) is not None else 'nofill')
        ln = s.find(c('spPr')).find(qn('a:ln'))
        dash = None
        if ln is not None:
            pd = ln.find(qn('a:prstDash'))
            dash = pd.get('val') if pd is not None else None
        mk = s.find(c('marker'))
        sym = mk.find(c('symbol')).get('val') if mk is not None else None
        print(f"   {name:<22} {vals} lbl={lab} dash={dash} marker={sym}")
print("legend:", ch.has_legend, ch.legend.position, ch.legend.include_in_layout)
print("title:", ch.chart_title.text_frame.text)

# --- one-decimal charts
def rd(chart):
    lc = chart._chartSpace.plotArea.find(c('lineChart'))
    out = []
    for s in lc.findall(c('ser')):
        vals = [x.find(c('v')).text for x in s.find(c('val')).find(c('numRef')).find(c('numCache')).findall(c('pt'))]
        dl = s.find(c('dLbls'))
        fmt = dl.find(c('numFmt')).get('formatCode') if dl is not None and dl.find(c('numFmt')) is not None else None
        out.append((vals, fmt))
    return out
for sh in sl.shapes:
    if sh.has_chart and sh.width < 5000000:
        print("\n", sh.name, rd(sh.chart))

# --- cards text
print("\ncard texts:")
for sh in sl.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() and sh.width > 5000000 and sh.height < 500000:
        print("  ", repr(sh.text_frame.text))
    if sh.has_text_frame and sh.text_frame.text.startswith(('结果', '内容端', '承接端', '达人端', '判断')):
        print("   row:", repr(sh.text_frame.text))

# --- overlap of graphics
boxes = [(sh.name, sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)
         for sh in sl.shapes if sh.top < 3800000 and sh.left > 400000]
print("\noverlaps (top band):")
for i in range(len(boxes)):
    for j in range(i + 1, len(boxes)):
        _, l1, t1, r1, b1 = boxes[i]
        _, l2, t2, r2, b2 = boxes[j]
        if l1 < r2 and l2 < r1 and t1 < b2 and t2 < b1:
            print("   OVERLAP", boxes[i][0], boxes[j][0])
print("   done")
