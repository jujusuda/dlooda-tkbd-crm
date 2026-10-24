# -*- coding: utf-8 -*-
from pptx import Presentation
F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_2.pptx'
prs = Presentation(F)
s = prs.slides[4]
def emu(v): return round(v/914400,2) if v is not None else None
print("=== NEW P5 shapes ===")
for i, sh in enumerate(s.shapes):
    L=emu(sh.left); T=emu(sh.top); W=emu(sh.width); H=emu(sh.height)
    extra=""
    if sh.has_text_frame: extra="TXT="+repr(sh.text_frame.text[:75])
    if sh.has_chart: extra="CHART="+str(sh.chart.chart_type)
    print(f"[{i}] L={L} T={T} W={W} H={H} {extra}")
print("\n=== CHART SERIES CHECK ===")
for sh in s.shapes:
    if sh.has_chart:
        ch=sh.chart
        for se in ch.plots[0].series:
            print(se.name, "->", [round(v) for v in se.values])
