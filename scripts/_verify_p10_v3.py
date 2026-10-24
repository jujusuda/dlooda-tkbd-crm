# -*- coding: utf-8 -*-
import zipfile, re, posixpath
from pptx import Presentation

F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_8.pptx'
prs = Presentation(F)
s = prs.slides[9]

# charts
charts = [sh for sh in s.shapes if sh.has_chart]
print("charts:", len(charts))
for sh in charts:
    ch = sh.chart
    vals = {se.name: [round(v,2) for v in se.values] for se in ch.plots[0].series}
    print("  ", vals)

# table
tb = [sh for sh in s.shapes if sh.has_table][0].table
print("\n=== TABLE ===")
for r in range(len(tb.rows)):
    print("  | ".join(tb.cell(r,c).text for c in range(len(tb.columns))))

# bottom text raw
print("\n=== BOTTOM TEXT (full) ===")
for sh in s.shapes:
    if sh.has_text_frame and '内容转化型' in sh.text_frame.text or (sh.has_text_frame and '达人承接型' in sh.text_frame.text):
        pass
bottom_texts = []
for sh in s.shapes:
    if sh.has_text_frame:
        t = sh.text_frame.text
        if '三连升' in t or '缺供给' in t or '逐月加码' in t or '放量型' in t or '非转化崩' in t:
            bottom_texts.append(t)
for t in bottom_texts:
    print("  ---")
    print("  ", t.replace("\n"," / "))

# forbidden full-paragraph markers
joined = " ".join(bottom_texts)
for marker in ['数据：', '变化：', '原因：']:
    print(f"\ncontains '{marker}':", marker in joined)

# key phrases present
for kw in ['三连升','缺供给','逐月加码','放量型','非转化崩']:
    print(f"key phrase '{kw}' present:", kw in joined)

# dLblPos legality on all line charts in file
z = zipfile.ZipFile(F)
illegal = []
for n in sorted(z.namelist()):
    if re.match(r'ppt/(slides/)?charts/chart\d+\.xml$', n) and '<c:lineChart>' in z.read(n).decode('utf-8','ignore'):
        pos = re.findall(r'<c:dLblPos val="(\w+)"/>', z.read(n).decode('utf-8','ignore'))
        for p in pos:
            if p in ('outEnd','inEnd'): illegal.append((n,p))
print("\nILLEGAL line-chart dLblPos (P10 not responsible):", illegal if illegal else "none")

# package integrity
ct = z.read('[Content_Types].xml').decode('utf-8','ignore')
overrides = set(re.findall(r'PartName="([^"]+)"', ct))
defaults = set(re.findall(r'Extension="([^"]+)"', ct))
names = set(z.namelist())
probs = []
for o in overrides:
    if o.lstrip('/') not in names: probs.append('CT override missing: '+o)
for nm in names:
    if nm == '[Content_Types].xml': continue
    ext = nm.rsplit('.',1)[-1].lower()
    if ('/'+nm) not in overrides and ext not in defaults: probs.append('part not declared: '+nm)
    if nm.endswith('.rels'):
        base = posixpath.dirname(posixpath.dirname(nm))
        for rid, tgt in re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read(nm).decode('utf-8','ignore')):
            if tgt.startswith('http'): continue
            full = tgt.lstrip('/') if tgt.startswith('/') else posixpath.normpath(posixpath.join(base, tgt))
            if full not in names: probs.append(f'{nm}: {rid}->{tgt} MISSING {full}')
print("PACKAGE PROBLEMS:", probs if probs else "none")

# no independent 落点策略 block
print("independent '落点策略' text block present:", any(sh.has_text_frame and sh.text_frame.text.strip().startswith('落点策略') for sh in s.shapes))
