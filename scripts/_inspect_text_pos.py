from pptx import Presentation

SRC = r'E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-09\8be9710a0a0dd58ff4cb419316577cc7\8月余慧玲8月月度汇报(2).pptx'
p = Presentation(SRC)

def shape_fill_hex(sh):
    try:
        if sh.fill.type == 1:
            return '#'+str(sh.fill.fore_color.rgb)
    except Exception:
        pass
    return None

for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        top = sh.top/914400 if sh.top is not None else None
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                try:
                    if run.font.color and run.font.color.type is not None:
                        col = '#'+str(run.font.color.rgb)
                        if col in ('#FFFFFF', '#FCEAF1'):
                            fh = shape_fill_hex(sh)
                            print('P%d  col=%s  top=%.2f  fill=%s  txt=[%s]' % (i, col, top, fh, run.text[:30]))
                except Exception:
                    pass
