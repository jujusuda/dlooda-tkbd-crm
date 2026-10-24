# -*- coding: utf-8 -*-
import zipfile, os
from pptx import Presentation

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(9).pptx"
DST = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\_pics"
os.makedirs(DST, exist_ok=True)

prs = Presentation(SRC)
sl = prs.slides[8]
for i, sh in enumerate(sl.shapes):
    if sh.shape_type == 13:
        blob = sh.image.blob
        ext = sh.image.ext or 'png'
        fn = os.path.join(DST, f"pic{i}.{ext}")
        with open(fn, 'wb') as f:
            f.write(blob)
        print("wrote", fn, len(blob), "bytes", sh.image.content_type)
