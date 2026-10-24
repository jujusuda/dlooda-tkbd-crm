# -*- coding: utf-8 -*-
from pptx import Presentation
import re, zipfile, posixpath
F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_9.pptx'
prs = Presentation(F)
s = prs.slides[10]

print("=== P11 text dump ===")
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print(repr(sh.text_frame.text))

print("\n=== TABLES in P11 ===")
for sh in s.shapes:
    if sh.has_table:
        tb = sh.table
        print("--- table %d cols ---" % len(tb.columns))
        for r in range(len(tb.rows)):
            row=[]
            for c in range(len(tb.columns)):
                cell = tb.cell(r,c)
                t = cell.text
                # detect rose bold
                bold_rose = False
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.bold and run.font.color and run.font.color.type is not None:
                            try:
                                if str(run.font.color.rgb) == 'B06579': bold_rose = True
                            except: pass
                row.append((t, '★' if bold_rose else ''))
            print("  ", row)

# forbidden words in subtitle / title?
full = "\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
print("\n=== checks ===")
print("淘汰退出 in page:", '淘汰退出' in full)
print("缺市场信号 in page:", '缺市场信号' in full)
print("具体数字 in subtitle? subtitle text:", [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and '寄样量相近' in sh.text_frame.text])

# package integrity
z = zipfile.ZipFile(F)
names = set(z.namelist())
ct = z.read('[Content_Types].xml').decode('utf-8','ignore')
overrides = set(re.findall(r'PartName="([^"]+)"', ct))
defaults = set(re.findall(r'Extension="([^"]+)"', ct))
probs = []
for o in overrides:
    if o.lstrip('/') not in names: probs.append('CT override missing: '+o)
for n in names:
    if n == '[Content_Types].xml': continue
    ext = n.rsplit('.',1)[-1].lower()
    if ('/'+n) not in overrides and ext not in defaults: probs.append('part not declared: '+n)
    if n.endswith('.rels'):
        base = posixpath.dirname(posixpath.dirname(n))
        for rid, tgt in re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read(n).decode('utf-8','ignore')):
            if tgt.startswith('http'): continue
            fullp = tgt.lstrip('/') if tgt.startswith('/') else posixpath.normpath(posixpath.join(base, tgt))
            if fullp not in names: probs.append(f'{n}: {rid}->{tgt} MISSING {fullp}')
print("PACKAGE PROBLEMS:", probs if probs else "none")
# shape bounds overflow check
SW = 13.33; SH = 7.5
for sh in s.shapes:
    if sh.left is None: continue
    if (sh.left/914400 + sh.width/914400) > SW + 0.05 or (sh.top/914400 + sh.height/914400) > SH + 0.05:
        print("OVERFLOW:", sh.shape_type, round(sh.left/914400,2), round(sh.top/914400,2), round(sh.width/914400,2), round(sh.height/914400,2))
print("overflow check done")
