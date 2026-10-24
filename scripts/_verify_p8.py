# -*- coding: utf-8 -*-
import zipfile, re, posixpath
from pptx import Presentation

F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_4.pptx'
# ---- content check ----
prs = Presentation(F)
s = prs.slides[7]
print("=== P8 text scan ===")
for sh in s.shapes:
    if sh.has_text_frame:
        t = sh.text_frame.text.replace("\n"," / ")
        if t.strip():
            print(" TXT:", t[:160])
tb = None
for sh in s.shapes:
    if sh.has_table: tb = sh.table; break
print("\n=== TABLE ===")
for r in range(len(tb.rows)):
    cells=[]
    for c in range(len(tb.columns)):
        cells.append(repr(tb.cell(r,c).text.replace("\n","|")))
    print("R%d:"%r, " | ".join(cells))

# ---- package integrity ----
print("\n=== PACKAGE VALIDATION ===")
z = zipfile.ZipFile(F)
names = set(z.namelist())
ct = z.read('[Content_Types].xml').decode('utf-8','ignore')
overrides = set(re.findall(r'PartName="([^"]+)"', ct))
defaults = set(re.findall(r'Extension="([^"]+)"', ct))
probs = []
for o in overrides:
    if o.lstrip('/') not in names: probs.append('CT override missing: '+o)
for n in names:
    if n == '[Content_Types].xml': continue
    ext = n.rsplit('.',1)[-1].lower()
    if ('/'+n) not in overrides and ext not in defaults: probs.append('part not declared: '+n)
    if n.endswith('.rels'):
        base = posixpath.dirname(posixpath.dirname(n))
        for rid, tgt in re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read(n).decode('utf-8','ignore')):
            if tgt.startswith('http'): continue
            full = tgt.lstrip('/') if tgt.startswith('/') else posixpath.normpath(posixpath.join(base, tgt))
            if full not in names: probs.append('%s: %s->%s MISSING %s'%(n,rid,tgt,full))
print("PACKAGE PROBLEMS:", probs if probs else "none")
# line chart dLblPos check
print("\n=== line chart dLblPos check ===")
for n in sorted(names):
    if re.match(r'ppt/(slides/)?charts/chart\d+\.xml$', n):
        x = z.read(n).decode('utf-8','ignore')
        if '<c:lineChart>' in x:
            pos = re.findall(r'<c:dLblPos val="(\w+)"/>', x)
            bad = [p for p in pos if p in ('outEnd','inEnd')]
            print(f"  {n}: dLblPos={pos} -> {'BAD(outEnd/inEnd on line!)' if bad else 'ok'}")
