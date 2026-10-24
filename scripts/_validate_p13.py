# -*- coding: utf-8 -*-
import zipfile
from pptx import Presentation
from pptx.oxml.ns import qn

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(11).pptx"
OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_v9_1(11)_达人素材供给修改.pptx"
IDX = 13  # 第 14 页

ok = True


def chk(cond, msg):
    global ok
    print(("  PASS  " if cond else "  FAIL  ") + msg)
    if not cond:
        ok = False


# --- package integrity
z = zipfile.ZipFile(OUT)
print("ZIP testzip:", z.testzip())
a = Presentation(SRC)
b = Presentation(OUT)
chk(len(a.slides) == len(b.slides), f"页数一致 {len(b.slides)}")

# --- other slides byte-identical
src_zip, out_zip = zipfile.ZipFile(SRC), zipfile.ZipFile(OUT)
same = diff = 0
for n in src_zip.namelist():
    parts = n.split("/")
    if len(parts) == 3 and parts[0] == "ppt" and parts[1] == "slides" and parts[2].startswith("slide") and n.endswith(".xml"):
        num = int(parts[2][5:-4])
        if num == IDX + 1:
            continue
        ra = src_zip.read(n).split(b"?>")[-1]
        rb = out_zip.read(n).split(b"?>")[-1]
        CR = chr(13); LF = chr(10)

        def norm(x):
            s = x.decode("utf-8")
            return s.replace(CR + LF, LF).replace(CR, LF).encode("utf-8")

        if norm(ra) == norm(rb):
            same += 1
        else:
            diff += 1
            print("   DIFF", n)
chk(diff == 0, f"其余 {same} 页与源内容完全一致（忽略 XML 声明/换行差异）")

sl = b.slides[IDX]
pics = [s for s in sl.shapes if s.shape_type is not None and str(s.shape_type).startswith("PICTURE")]
chk(len(pics) == 0, f"本页图片数 = {len(pics)}（要求 0）")
charts = [s for s in sl.shapes if s.has_chart]
chk(len(charts) == 2, f"本页原生图表数 = {len(charts)}（要求 2）")

# --- chart 1: 整月人均
c1 = charts[0].chart
chk([str(x) for x in c1.plots[0].categories] == ["8月", "9月"], f"左图横轴 {[str(x) for x in c1.plots[0].categories]}")
v1 = list(c1.plots[0].series[0].values)
chk(v1 == [1.48, 1.90], f"左图数值 {v1}")
chk(c1.chart_title.text_frame.text.strip() == "整月人均素材（条/人）", f"左图标题 {c1.chart_title.text_frame.text!r}")
el1 = c1.plots[0]._element
dp = {}
for d in el1.iter(qn("c:dPt")):
    i = d.find(qn("c:idx")).get("val")
    clr = d.find(".//" + qn("a:srgbClr"))
    dp[i] = clr.get("val") if clr is not None else None
print("  左图逐点配色:", dp)
chk(dp.get("0") == "D9A5B5", "8月柱 = #D9A5B5")
chk(dp.get("1") == "B06579", "9月柱 = #B06579")
chk(not c1.has_legend, "左图无图例")

# --- chart 2: 四款分组
c2 = charts[1].chart
chk([str(x) for x in c2.plots[0].categories] == ["28173", "2190", "2178", "2189"],
    f"右图横轴 {[str(x) for x in c2.plots[0].categories]}")
names = [s.name for s in c2.plots[0].series]
vals = [list(s.values) for s in c2.plots[0].series]
print("  右图系列:", names, vals)
chk(names == ["8月", "9月"], "右图系列名 8月/9月")
chk(vals[0] == [1.52, 1.42, 1.21, 1.47], f"8月 {vals[0]}")
chk(vals[1] == [1.79, 1.69, 1.50, 1.48], f"9月 {vals[1]}")
chk(c2.chart_title.text_frame.text.strip() == "四款爆品人均素材（条/人）", f"右图标题 {c2.chart_title.text_frame.text!r}")
cols = []
for s in c2.plots[0].series:
    e = s._element.find(".//" + qn("a:srgbClr"))
    cols.append(e.get("val") if e is not None else None)
print("  右图系列配色:", cols)
chk(cols == ["D9A5B5", "B06579"], "8月浅豆沙 / 9月玫红")
chk(c2.has_legend and str(c2.legend.position).startswith("TOP"), f"右图图例 {c2.legend.position}")

# --- data labels
for tag, ch in (("左", c1), ("右", c2)):
    el = ch.plots[0]._element
    nfmt = [d.find(qn("c:numFmt")).get("formatCode") for d in el.iter(qn("c:dLbls")) if d.find(qn("c:numFmt")) is not None]
    sv = [d.find(qn("c:showVal")).get("val") for d in el.iter(qn("c:dLbls")) if d.find(qn("c:showVal")) is not None]
    pos = [d.find(qn("c:dLblPos")).get("val") for d in el.iter(qn("c:dLbls")) if d.find(qn("c:dLblPos")) is not None]
    print(f"  {tag}图 dLbls numFmt={nfmt} showVal={sv} pos={pos}")
    chk(all(f == "0.00" for f in nfmt) and nfmt, f"{tag}图标签两位小数")
    chk(all(v == "1" for v in sv) and sv, f"{tag}图显示数值")
    chk(all(p == "outEnd" for p in pos), f"{tag}图标签位置 outEnd（柱状合法）")

# --- text
alltext = "\n".join(s.text_frame.text for s in sl.shapes if s.has_text_frame)
for kw in ["因果", "结论", "动作", "1070→752", "723→395", "1.90", "+28%",
           "延续复投核心达人", "重点补 2189 素材", "稳定多次发布达人占比", "2+"]:
    chk(kw in alltext, f"含文案 {kw!r}")
chk("达人素材供给" in alltext, "主标题保留")
for bad in ["靠人头铺量", "近七成"]:
    chk(bad not in alltext, f"无残留旧措辞 {bad!r}")

# --- overlap / overflow
boxes = []
for s in sl.shapes:
    if s.left is None:
        continue
    boxes.append((s.shape_type, s.left / 914400, s.top / 914400, s.width / 914400, s.height / 914400))
W, H = b.slide_width / 914400, b.slide_height / 914400
for t, x, y, w, h in boxes:
    if x < -0.01 or y < -0.01 or x + w > W + 0.01 or y + h > H + 0.01:
        ok = False
        print(f"  FAIL  越界 {t} {x:.2f},{y:.2f},{w:.2f},{h:.2f}")
print("  边界检查完成（画布 %.2f x %.2f）" % (W, H))

print("\nRESULT:", "ALL PASS" if ok else "HAS FAILURES")
