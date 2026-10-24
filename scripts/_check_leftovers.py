# -*- coding: utf-8 -*-
import zipfile
from pptx import Presentation
from pptx.oxml.ns import qn

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(7).pptx"
OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_P9修改.pptx"

for tag, path in (('SRC', SRC), ('OUT', OUT)):
    z = zipfile.ZipFile(path)
    charts = [n for n in z.namelist() if '/charts/' in n and n.endswith('.xml')]
    xl = [n for n in z.namelist() if 'embeddings' in n]
    print(f"{tag}: size={len(z.read('ppt/presentation.xml'))} bytes, "
          f"pkg_entries={len(z.namelist())}, chartparts={len(charts)}, embeddings={len(xl)}")
    z.close()

prs = Presentation(OUT)
sl = prs.slides[8]
ch = sl.shapes[19].chart
chart_el = ch._chartSpace.chart
print("\nchart children:", [e.tag.split('}')[1] for e in chart_el])
lg = ch.has_legend, ch.legend.position, ch.legend.include_in_layout
print("legend:", lg)

# series -> confirm 4 lines + 2 bars exist and legend would list them
pa = ch._chartSpace.plotArea
for nm in ('barChart', 'lineChart'):
    p = pa.find(qn('c:' + nm))
    sers = p.findall(qn('c:ser'))
    print(f"\n{nm}: {len(sers)} series")
    for s in sers:
        tx = s.find(qn('c:tx')).find(qn('c:strRef')).find(qn('c:strCache')).find(qn('c:pt')).find(qn('c:v')).text
        idx = s.find(qn('c:idx')).get('val')
        order = s.find(qn('c:order')).get('val')
        kids = [k.tag.split('}')[1] for k in s]
        print(f"   idx={idx} order={order} {tx:<22} children={kids}")
