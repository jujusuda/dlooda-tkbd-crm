# -*- coding: utf-8 -*-
"""第 14 页（达人素材供给）重做：左=整月人均趋势柱图，右=四款分组柱图，下方=因果/结论/动作。
全部 PPT 原生可编辑，禁止图片。
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(11).pptx"
OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_v9_1(11)_达人素材供给修改.pptx"

FONT = "微软雅黑"
ROSE = RGBColor(0xB0, 0x65, 0x79)   # 玫红  #B06579
PINK = RGBColor(0xD9, 0xA5, 0xB5)   # 浅豆沙粉 #D9A5B5
INK = RGBColor(0x1F, 0x1F, 0x1F)
BODY = RGBColor(0x3D, 0x3D, 0x3D)
MUTE = RGBColor(0x9A, 0x9A, 0x9A)
CARD_BG = RGBColor(0xF7, 0xF4, 0xF5)


# ---------------- font helpers ----------------
def force_ea(font):
    """让中文字形也走微软雅黑（python-pptx 只写 a:latin）"""
    try:
        rPr = font._rPr
    except Exception:
        return
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        return
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = copy.deepcopy(latin)
        ea.tag = qn("a:ea")
        latin.addnext(ea)
    else:
        ea.set("typeface", latin.get("typeface"))


def style_font(font, size=None, bold=None, color=None, name=FONT):
    if name:
        font.name = name
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color
    force_ea(font)


def add_para(tf, parts, align=PP_ALIGN.LEFT, space_after=0, first=False, line_spacing=None):
    """parts: list of (text, size, bold, color)"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if space_after:
        p.space_after = Pt(space_after)
    if line_spacing:
        p.line_spacing = line_spacing
    for t, size, bold, color in parts:
        r = p.add_run()
        r.text = t
        style_font(r.font, size, bold, color)
    return p


def new_tb(slide, x, y, w, h, wrap=True, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf


# ---------------- chart helpers ----------------
def ensure_numfmt(dlbls, code):
    """确保 dLbls 里有 numFmt（位置：CT_DLbls 首元素）"""
    nf = dlbls.find(qn("c:numFmt"))
    if nf is None:
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import nsdecls
        nf = parse_xml(f'<c:numFmt {nsdecls("c")} formatCode="{code}" sourceLinked="0"/>')
        dlbls.insert(0, nf)
    else:
        nf.set("formatCode", code)
        nf.set("sourceLinked", "0")


def setup_labels(plot, fmt="0.00", size=10):
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.show_value = True
    try:
        dl.number_format = fmt
        dl.number_format_is_linked = False
    except Exception as e:
        print("  [warn] number_format api:", e)
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    style_font(dl.font, size, False, BODY)
    # 兜底：直接写 XML，确保每一处 dLbls 都带 numFmt
    for d in plot._element.iter(qn("c:dLbls")):
        ensure_numfmt(d, fmt)


def chart_font(chart, size=10, color=BODY):
    style_font(chart.font, size, False, color)


def style_axis(ax, size=10, color=MUTE, fmt=None, maximum=None, major=None):
    ax.has_major_gridlines = False
    tl = ax.tick_labels
    style_font(tl.font, size, False, color)
    if fmt is not None:
        try:
            tl.number_format = fmt
            tl.number_format_is_linked = False
        except Exception as e:
            print("  [warn] axis numfmt:", e)
    if maximum is not None:
        ax.maximum_scale = maximum
        ax.minimum_scale = 0
    if major is not None:
        ax.major_unit = major


def set_point_color(series, idx, rgb):
    pt = series.points[idx]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb
    pt.format.line.fill.background()


# =====================================================================
prs = Presentation(SRC)
slide = prs.slides[13]  # 第 14 页：达人素材供给

# ---- 删除原有内容（保留 0 强调条 / 1 主标题 / 2 PART 标签 / 3 分隔线）
keep = {0, 1, 2, 3}
for sh in list(slide.shapes):
    if list(slide.shapes).index(sh) in keep:
        continue
for idx in sorted([i for i in range(len(slide.shapes)) if i not in keep], reverse=True):
    el = slide.shapes[idx]._element
    el.getparent().remove(el)

print("shapes left:", len(slide.shapes))

# ---------------- 左：整月人均素材（8月 vs 9月） ----------------
CH1 = (0.6, 1.78, 5.4, 2.6)
cd = CategoryChartData()
cd.categories = ["8月", "9月"]
cd.add_series("人均素材", (1.48, 1.90))
gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                            Inches(CH1[0]), Inches(CH1[1]), Inches(CH1[2]), Inches(CH1[3]), cd)
ch1 = gf.chart
ch1.has_legend = False
ch1.has_title = True
tr = ch1.chart_title.text_frame
tr.text = "整月人均素材（条/人）"
tr.paragraphs[0].alignment = PP_ALIGN.LEFT
style_font(tr.paragraphs[0].runs[0].font, 11, True, INK)
force_ea(ch1.chart_title.text_frame.paragraphs[0].runs[0].font)
chart_font(ch1, 10, BODY)
pl = ch1.plots[0]
pl.gap_width = 130
s0 = pl.series[0]
s0.format.fill.solid()
s0.format.fill.fore_color.rgb = ROSE
s0.format.line.fill.background()
set_point_color(s0, 0, PINK)   # 8月 浅豆沙粉
set_point_color(s0, 1, ROSE)   # 9月 玫红
setup_labels(pl, "0.00", 10)
style_axis(ch1.category_axis, 11, BODY)
style_axis(ch1.value_axis, 10, MUTE, fmt="0.0", maximum=2.5, major=0.5)

# +28% 标注
tb, tf = new_tb(slide, 2.55, 2.14, 1.5, 0.30)
add_para(tf, [("+28%", 16, True, ROSE)], align=PP_ALIGN.CENTER, first=True)

# 图下一句结论
tb, tf = new_tb(slide, 0.6, 4.46, 5.4, 0.32)
add_para(tf, [("9月人均素材较8月回升（1.48→1.90，", 12, False, BODY),
              ("+28%", 12, True, ROSE),
              ("）", 12, False, BODY)], first=True)

# ---------------- 右：四款爆品人均素材（8月 vs 9月） ----------------
CH2 = (6.35, 1.78, 5.98, 2.6)
cd2 = CategoryChartData()
cd2.categories = ["28173", "2190", "2178", "2189"]
cd2.add_series("8月", (1.52, 1.42, 1.21, 1.47))
cd2.add_series("9月", (1.79, 1.69, 1.50, 1.48))
gf2 = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                             Inches(CH2[0]), Inches(CH2[1]), Inches(CH2[2]), Inches(CH2[3]), cd2)
ch2 = gf2.chart
ch2.has_title = True
tr2 = ch2.chart_title.text_frame
tr2.text = "四款爆品人均素材（条/人）"
tr2.paragraphs[0].alignment = PP_ALIGN.LEFT
style_font(tr2.paragraphs[0].runs[0].font, 11, True, INK)
chart_font(ch2, 10, BODY)
pl2 = ch2.plots[0]
pl2.gap_width = 55
pl2.overlap = -12
for ser, col in zip(pl2.series, [PINK, ROSE]):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col
    ser.format.line.fill.background()
setup_labels(pl2, "0.00", 10)
style_axis(ch2.category_axis, 11, BODY)
style_axis(ch2.value_axis, 10, MUTE, fmt="0.0", maximum=2.4, major=0.5)
ch2.has_legend = True
lg = ch2.legend
lg.position = XL_LEGEND_POSITION.TOP
lg.include_in_layout = False
lg.font.size = Pt(10)
lg.font.name = FONT
lg.font.color.rgb = BODY
force_ea(lg.font)

# ---------------- 下方：因果 / 结论 / 动作 ----------------
CARDS = [
    ("因果", [("9月素材总量下降（1070→752），但达人更精炼（723→395），人均反升；四款普遍上升，", 12, False, BODY),
              ("是整体供给效率提升，非单款拉动", 12, True, BODY)]),
    ("结论", [("9月人均素材 1.90，较8月回升（+28%），达人端素材供给效率提升，", 12, False, BODY),
              ("核心款人均普遍 1.5 以上", 12, True, BODY)]),
    ("动作", None),
]
CX, CY, CW, CH = 0.6, 4.92, 3.87, 1.62
GAP = 0.26
ACTION = [
    "① 延续复投核心达人（产出更高效）",
    "② 重点补 2189 素材（唯一持平款）",
    "③ 稳定多次发布达人占比，让人均往 2+ 走",
]

for i, (title, parts) in enumerate(CARDS):
    x = CX + i * (CW + GAP)
    # 底板
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(CY), Inches(CW), Inches(CH))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.fill.background()
    card.shadow.inherit = False
    card.text_frame.text = ""
    # 左侧强调条
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(CY + 0.12), Inches(0.055), Inches(CH - 0.24))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ROSE
    bar.line.fill.background()
    bar.shadow.inherit = False
    bar.text_frame.text = ""
    # 标题
    tb, tf = new_tb(slide, x + 0.22, CY + 0.16, CW - 0.36, 0.3)
    add_para(tf, [(title, 14, True, ROSE)], first=True)
    # 正文
    tb, tf = new_tb(slide, x + 0.22, CY + 0.52, CW - 0.36, CH - 0.66)
    if parts is not None:
        add_para(tf, parts, first=True, line_spacing=1.25)
    else:
        for j, a in enumerate(ACTION):
            if a.startswith("②"):
                add_para(tf, [(a[:2], 12, True, ROSE), (a[2:], 12, False, BODY)],
                         first=(j == 0), space_after=3, line_spacing=1.15)
            else:
                add_para(tf, [(a, 12, False, BODY)],
                         first=(j == 0), space_after=3, line_spacing=1.15)

prs.save(OUT)
print("saved:", OUT)
