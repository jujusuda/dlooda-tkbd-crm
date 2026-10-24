# -*- coding: utf-8 -*-
from pptx import Presentation
F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_1.pptx'
prs = Presentation(F)
s = prs.slides[4]  # P5
def emu(v):
    return round(v/914400,2) if v is not None else None
print("=== SLIDE 5 shapes ===")
for i, sh in enumerate(s.shapes):
    L=emu(sh.left); T=emu(sh.top); W=emu(sh.width); H=emu(sh.height)
    kind = sh.shape_type
    extra=""
    if sh.has_text_frame:
        extra = "TXT="+repr(sh.text_frame.text[:70])
    if sh.has_chart:
        extra = "CHART="+str(sh.chart.chart_type)
    if sh.has_table:
        extra = "TABLE rows=%d cols=%d"%(len(sh.table.rows), len(sh.table.columns))
    print(f"[{i}] L={L} T={T} W={W} H={H} {kind} {extra}")
