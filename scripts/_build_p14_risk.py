from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData
import collections

SRC = r'E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-09\8be9710a0a0dd58ff4cb419316577cc7\8月余慧玲8月月度汇报(4).pptx'
OUT = r'C:\Users\LENOVO\Desktop\8月余慧玲8月月度汇报_第14页达人结构风险_v2.pptx'

p = Presentation(SRC)
print('Slides:', len(p.slides), 'Size: %.2fx%.2f in' % (p.slide_width/914400, p.slide_height/914400))

# ===== 从源文件提取配色 =====
fills = collections.Counter()
for s in p.slides:
    for sh in s.shapes:
        try:
            if sh.fill.type == 1: fills['#'+str(sh.fill.fore_color.rgb)] += 1
        except: pass
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    try:
                        if c.fill.type == 1: fills['#'+str(c.fill.fore_color.rgb)] += 1
                    except: pass
        if sh.has_chart:
            for ser in sh.chart.series:
                try:
                    if ser.format.fill.type == 1: fills['#'+str(ser.format.fill.fore_color.rgb)] += 1
                except: pass

print('Source fill colors:', dict(fills))

# ===== 配色（沿用原文件艳粉主题，与 PPT 其余页一致）=====
MAIN      = RGBColor.from_string('C94F82')   # 深玫 标题栏/表头（原文件主色）
CARD_BG   = RGBColor.from_string('FCEAF1')   # 浅粉底 卡片/隔行
BODY      = RGBColor.from_string('4A3D42')   # 正文深色
WHITE     = RGBColor.from_string('FFFFFF')
GRAY      = RGBColor.from_string('7A6A70')   # 辅助灰
RED       = RGBColor.from_string('E54C4C')   # 预警/极高（原文件红）
YELLOW    = RGBColor.from_string('E0A83C')   # 中风险（原文件橙黄）
GREEN     = RGBColor.from_string('4CAF50')   # 低风险/正向（原文件绿）
CHART_BAR = RGBColor.from_string('C94F82')   # 图表柱=深玫
LINE_SOFT = RGBColor.from_string('F5A0BC')   # 边框=亮粉

# ===== 数据（严格不修改）=====
SKUS = ['2190','2178','2722','2189','2789','28173']  # 按风险降序
DATA = {
    '2190': {'total': 16, 'old_pct': 99.5, 'sample_pct': 18.8, 'risk': '极高', 'risk_color': RED},
    '2178': {'total': 19, 'old_pct': 96.9, 'sample_pct': 26.3, 'risk': '极高', 'risk_color': RED},
    '2722': {'total': 27, 'old_pct': 97.2, 'sample_pct': 44.4, 'risk': '极高', 'risk_color': RED},
    '2189': {'total': 19, 'old_pct': 85.8, 'sample_pct': 42.1, 'risk': '中',   'risk_color': YELLOW},
    '2789': {'total': 12, 'old_pct': 71.8, 'sample_pct': 58.3, 'risk': '中',   'risk_color': YELLOW},
    '28173':{'total': 33, 'old_pct': 61.0, 'sample_pct': 75.8, 'risk': '低',   'risk_color': GREEN},
}

STRATEGY = {
    '2190': '99.5% 订单来自老达人，新寄样达人几乎无产出。优先维护存量头部老达人，紧急批量储备备份新达人，解决梯队断层风险',
    '2178': '订单高度依赖老达人，新寄样达人单产很低。定向拓展同赛道垂类达人，持续补充新达人池，分散头部依赖风险',
    '2722': '寄样达人数量不少，但订单贡献极低，新人转化弱。复盘达人筛选标准，淘汰低效达人，更换高潜力账号寄样',
    '2189': '老达人仍是订单主力，新寄样达人有基础产出。维持现有寄样节奏，跟进已寄达人二次视频，持续提升新达人订单占比',
    '2789': '寄样达人数量已经过半，但整体订单体量偏小。评估品效，减少无效寄样，重点孵化已经出单的潜力新人',
    '28173':'新寄样达人贡献接近 4 成订单，达人池厚度充足。作为主力扩量款，扩大寄样规模，把这款的达人运营经验复制到其他 SKU',
}

# ===== 找到第 14 页（index 13）并清空重建 =====
slide_idx = 13  # 0-based
s = p.slides[slide_idx]

# 删除现有形状（从后往前删避免索引错）
shapes_to_remove = [sh for sh in s.shapes]
for sh in reversed(shapes_to_remove):
    sp = sh._element
    sp.getparent().remove(sp)

# ===== 辅助函数 =====
def rect(slide, l, t, w, h, fill=None, line=None, lw=0):
    sh = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))  # 1=rectangle
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh

def tb(slide, l, t, w, h, runs, bold=False, color=BODY, size=10, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, fontname='微软雅黑'):
    sh = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = sh.text_frame; tf.word_wrap = True; tf.auto_size = None
    tf.paragraphs[0].alignment = align
    if anchor:
        tf.paragraphs[0].space_before = Pt(0); tf.paragraphs[0].space_after = Pt(0)
    for i, r in enumerate(runs):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        if isinstance(r, dict):
            txt = r.get('t', ''); sz = r.get('s', size); c = r.get('c', color)
            b = r.get('b', bold); fn = r.get('f', fontname)
            al = r.get('a', align)
            p.alignment = al
        else:
            txt = r; sz = size; c = color; b = bold; fn = fontname
        run = p.add_run(); run.text = txt
        run.font.size = Pt(sz); run.font.bold = b
        run.font.color.rgb = c if isinstance(c, RGBColor) else RGBColor.from_string(c)
        run.font.name = fn
    # set anchor on frame
    try: tf.anchor = anchor
    except: pass
    return sh

def table(slide, l, t, w, col_widths, rows, fsize=9, rowh=0.28, headh=0.32,
          head_fill=MAIN, head_text=WHITE, alt_fill=CARD_BG):
    nrows = len(rows); ncols = len(col_widths)
    tbl_sh = slide.shapes.add_table(nrows, ncols, Inches(l), Inches(t),
                                      Inches(w), Inches(rowh*nrows + headh))
    tbl = tbl_sh.table if hasattr(tbl_sh, 'table') else tbl_sh
    # set column widths
    for ci, cw in enumerate(col_widths):
        tbl.columns[ci].width = Inches(cw)
    for ri, row_data in enumerate(rows):
        rh = headh if ri == 0 else rowh
        # set row height via cell text frame padding is tricky; use minimum height
        for ci, cell_val in enumerate(row_data):
            cell = tbl.cell(ri, ci)
            cell.text = ''
            if isinstance(cell_val, dict):
                txt = cell_val.get('t', '')
                c = cell_val.get('c', BODY)
                b = cell_val.get('b', False)
            else:
                txt = str(cell_val); c = BODY; b = False
            cell.text = txt
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(fsize); p.font.name = '微软雅黑'
                p.font.color.rgb = c if isinstance(c, RGBColor) else RGBColor.from_string(c)
                p.font.bold = b
                p.alignment = PP_ALIGN.CENTER
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = head_fill
            elif ri % 2 == 0 and alt_fill:
                cell.fill.solid(); cell.fill.fore_color.rgb = alt_fill
            else:
                cell.fill.background()
    return tbl_sh

# ============================================================
# 布局参数（英寸）
# ============================================================
PAGE_W = 13.333  # 宽屏 16:9
PAGE_H = 7.5

# 顶部标题区
TITLE_H = 0.95

# 三栏布局
LEFT_W = 5.8       # 左侧图表区
RIGHT_W = 7.2      # 右侧双表区
GAP = 0.35
LEFT_L = 0.45
RIGHT_L = LEFT_L + LEFT_W + GAP

# 右侧上下两表
RIGHT_TOP_H = 2.6   # 上表（明细）
RIGHT_GAP = 0.2
RIGHT_BOT_H = 2.5   # 下表（策略）

# 左侧图表
CHART_T = TITLE_H + 0.15
CHART_H = RIGHT_TOP_H + RIGHT_GAP + RIGHT_BOT_H - 0.1

# ============================================================
# 1) 顶部标题栏
# ============================================================
rect(s, 0, 0, PAGE_W, TITLE_H, fill=MAIN)
tb(s, 0.35, 0.12, 11, 0.45,
   [{'t': '达人结构集中度风险分析', 's': 22, 'b': True, 'c': WHITE}],
   size=22, align=PP_ALIGN.LEFT)

# Part 03 标签（右上角）
rect(s, 12.0, 0.25, 1.0, 0.42, fill=CARD_BG)
tb(s, 12.0, 0.27, 1.0, 0.38,
   [{'t': 'Part 03', 's': 10, 'b': True, 'c': MAIN}],
   size=10, align=PP_ALIGN.CENTER)

# 导语
tb(s, 0.45, TITLE_H - 0.32, 12.0, 0.3,
   [{'t': '核心判断：多款爆款高度依赖存量老达人，新寄样达人订单承接能力不足；一旦头部老达人断更，大盘订单会受到明显冲击，各 SKU 新老达人梯队健康度差异很大。',
     's': 9, 'c': GRAY}],
   size=9, color=GRAY)

# ============================================================
# 2) 左侧：横向条形图 —— 各 SKU 订单对老达人依赖度
# ============================================================
chart_l = LEFT_L
chart_t = CHART_T + 0.35
chart_w = LEFT_W - 0.1
chart_h = CHART_H - 0.55

chart_data = CategoryChartData()
chart_data.categories = SKUS
chart_data.add_series('老达人订单占比', [DATA[k]['old_pct'] for k in SKUS])

ch = s.shapes.add_chart(
    XL_CHART_TYPE.BAR_CLUSTERED, Inches(chart_l), Inches(chart_t),
    Inches(chart_w), Inches(chart_h), chart_data
).chart

# 图表样式
ch.has_legend = False
try:
    plot = ch.plot_area
    if plot.has_fill: plot.has_fill = False
except Exception:
    pass  # 某些版本不支持，跳过

# 系列颜色：统一主色（风险差异在右侧表格中用颜色标出）
series = ch.series[0]
series.format.fill.solid()
series.format.fill.fore_color.rgb = CHART_BAR

# 坐标轴
ax = ch.value_axis
ax.has_major_gridlines = True
ax.major_gridlines.format.line.color.rgb = RGBColor.from_string('E0D0D5')
ax.major_gridlines.format.line.width = Pt(0.5)
ax.maximum_scale = 100
ax.minimum_scale = 0

cat_ax = ch.category_axis
cat_ax.tick_labels.font.size = Pt(9)
cat_ax.tick_labels.font.color.rgb = BODY
ax.tick_labels.font.size = Pt(8.5)
ax.tick_labels.font.color.rgb = GRAY

# 图表标题
tb(s, chart_l, CHART_T + 0.02, chart_w, 0.3,
   [{'t': '各 SKU 订单对老达人依赖度', 's': 12, 'b': True, 'c': MAIN}],
   size=12, bold=True, color=MAIN)

# 图表解读
tb(s, chart_l, chart_t + chart_h + 0.05, chart_w, 0.35,
   [{'t': '▸ 条形越长 → 订单越依赖老达人 → 新寄样达人产出薄弱 → 达人结构风险越高',
     's': 8.5, 'c': GRAY}],
   size=8.5, color=GRAY)

# ============================================================
# 3) 右上：SKU 风险明细表格
# ============================================================
tbl_t = CHART_T + 0.02
detail_rows = [['SKU', '总出单达人', '老达人订单占比', '寄样达人占比', '风险等级']]
for sku in SKUS:
    d = DATA[sku]
    detail_rows.append([
        sku,
        str(d['total']),
        '%.1f%%' % d['old_pct'],
        '%.1f%%' % d['sample_pct'],
        {'t': d['risk'], 'c': d['risk_color'], 'b': True}
    ])

table(s, RIGHT_L, tbl_t, RIGHT_W,
      [0.65, 1.15, 1.55, 1.35, 0.9],
      detail_rows, fsize=9, rowh=0.30, headh=0.34)

# 表标题
tb(s, RIGHT_L, tbl_t - 0.28, RIGHT_W, 0.26,
   [{'t': 'SKU 风险明细（按风险降序）', 's': 11, 'b': True, 'c': MAIN}],
   size=11, bold=True, color=MAIN)

# ============================================================
# 4) 右下：按风险制定达人运营动作表格
# ============================================================
strat_t = tbl_t + RIGHT_TOP_H + RIGHT_GAP - 0.05
strat_rows = [['SKU', '风险', '下月动作']]
for sku in SKUS:
    d = DATA[sku]
    strat_rows.append([
        sku,
        {'t': d['risk'], 'c': d['risk_color'], 'b': True},
        STRATEGY[sku]
    ])

table(s, RIGHT_L, strat_t, RIGHT_W,
      [0.60, 0.50, 5.85],
      strat_rows, fsize=8.2, rowh=0.36, headh=0.34)

# 表标题
tb(s, RIGHT_L, strat_t - 0.28, RIGHT_W, 0.26,
   [{'t': '按风险制定达人运营动作', 's': 11, 'b': True, 'c': MAIN}],
   size=11, bold=True, color=MAIN)

# ============================================================
# 5) 右下角口径说明框
# =================================================#
note_l = RIGHT_L + RIGHT_W - 3.3
note_t = strat_t + 6 * 0.36 + 0.38 + 0.1  # after strategy table
note_w = 3.4
note_h = 1.55

# 口径框背景
rect(s, note_l, note_t, note_w, note_h, fill=CARD_BG, line=LINE_SOFT, lw=1.0)

tb(s, note_l + 0.1, note_t + 0.06, note_w - 0.2, note_h - 0.1,
   [
       {'t': '📋 口径说明', 's': 9, 'b': True, 'c': MAIN},
       {'t': '', 's': 4},
       {'t': '• 寄样达人：寄样时间落在 2026.7.15–8.15 区间的达人', 's': 7.5, 'c': BODY},
       {'t': '• 老达人：寄样时间早于 2026.7.15 的存量达人', 's': 7.5, 'c': BODY},
       {'t': '• 订单统计口径：8 月全月 SKU 总订单，包含历史老视频回单', 's': 7.5, 'c': BODY},
       {'t': '', 's': 3},
       {'t': '⚠️ 不点名任何达人 ID，只看结构趋势', 's': 7.5, 'c': GRAY},
   ],
   size=9)

# ============================================================
# 6) 页码（右下角）
# ============================================================
tb(s, 12.6, 7.15, 0.5, 0.3,
   [{'t': '14', 's': 9, 'c': GRAY}],
   size=9, color=GRAY, align=PP_ALIGN.RIGHT)

# ============================================================
# 保存
# ============================================================
p.save(OUT)
print('\nSAVED:', OUT)
print('Slide 14 rebuilt: 达人结构集中度风险分析')
print('Data verified:')
for sku in SKUS:
    d = DATA[sku]
    print('  %s: total=%d old=%.1f%% sample=%.1f%% risk=%s' % (
        sku, d['total'], d['old_pct'], d['sample_pct'], d['risk']))
