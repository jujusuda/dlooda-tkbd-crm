from pptx import Presentation
from pptx.util import Emu
import collections

SRC = r'E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-09\8be9710a0a0dd58ff4cb419316577cc7\8月余慧玲8月月度汇报(3).pptx'

p = Presentation(SRC)
print('SLIDE COUNT:', len(p.slides))
print('SIZE: %.2f x %.2f in' % (p.slide_width/914400, p.slide_height/914400))

# slide backgrounds
print('\n=== SLIDE BACKGROUNDS ===')
for i, s in enumerate(p.slides, 1):
    try:
        ft = s.background.fill.type
        if ft == 1:
            print('  P%d bg=#%s' % (i, str(s.background.fill.fore_color.rgb)))
        else:
            print('  P%d bg=type%s (none/theme)' % (i, ft))
    except Exception as e:
        print('  P%d bg=err' % i)

# collect colors
fills = collections.Counter()
textcols = collections.Counter()
lines = collections.Counter()
for s in p.slides:
    for sh in s.shapes:
        try:
            if sh.fill.type == 1:
                fills['#'+str(sh.fill.fore_color.rgb)] += 1
        except Exception:
            pass
        try:
            if sh.line and sh.line.color and sh.line.color.type == 1:
                lines['#'+str(sh.line.color.rgb)] += 1
        except Exception:
            pass
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    try:
                        if c.fill.type == 1:
                            fills['#'+str(c.fill.fore_color.rgb)] += 1
                    except Exception:
                        pass
        if sh.has_chart:
            for ser in sh.chart.series:
                try:
                    if ser.format.fill.type == 1:
                        fills['#'+str(ser.format.fill.fore_color.rgb)] += 1
                except Exception:
                    pass
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        if run.font.color and run.font.color.type is not None:
                            textcols['#'+str(run.font.color.rgb)] += 1
                    except Exception:
                        pass

print('\n=== FILL COLORS (shape+table+chart) ===')
for c, n in fills.most_common():
    print('  %s : %d' % (c, n))
print('\n=== TEXT COLORS ===')
for c, n in textcols.most_common():
    print('  %s : %d' % (c, n))
print('\n=== LINE COLORS ===')
for c, n in lines.most_common():
    print('  %s : %d' % (c, n))

# chart series detail per slide
print('\n=== CHART SERIES DETAIL ===')
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if sh.has_chart:
            ch = sh.chart
            print('  P%d chart:' % i)
            try:
                for ser in ch.series:
                    try:
                        nm = ser.name
                    except Exception:
                        nm = '?'
                    fc = '?'
                    try:
                        if ser.format.fill.type == 1:
                            fc = '#%s' % str(ser.format.fill.fore_color.rgb)
                    except Exception:
                        pass
                    print('     series=%s fill=%s' % (nm, fc))
            except Exception as e:
                print('     err', e)
            # axis / plot area
            try:
                if ch.category_axis and ch.category_axis.tick_labels:
                    pass
            except Exception:
                pass
