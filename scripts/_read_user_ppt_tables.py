# -*- coding: utf-8 -*-
from pptx import Presentation

path = r"E:/xwechat_files/wxid_od6e7r1klxs422_54c2/temp/RWTemp/2026-09/8be9710a0a0dd58ff4cb419316577cc7/余慧玲8月月度汇报.pptx"
prs = Presentation(path)

for idx, slide in enumerate(prs.slides, start=1):
    has_content = False
    for shp in slide.shapes:
        if shp.has_table:
            has_content = True
            print(f"===== Slide {idx} TABLE: {shp.name} =====")
            tbl = shp.table
            for r in tbl.rows:
                cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
                print("  | " + " | ".join(cells))
            print()
        if shp.has_chart:
            has_content = True
            ch = shp.chart
            print(f"===== Slide {idx} CHART: {shp.name} ({ch.chart_type}) =====")
            try:
                cats = list(ch.plots[0].categories)
                print("  cats:", cats)
            except Exception as e:
                print("  cats err:", e)
            for s in ch.plots[0].series:
                try:
                    print(f"  series '{s.name}':", list(s.values))
                except Exception as e:
                    print(f"  series err:", e)
            print()
    if has_content:
        print()
