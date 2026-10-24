# -*- coding: utf-8 -*-
"""v8 -> v9 修改脚本（按《Dlooda9月汇报_PPT修改指令_v9》执行）
全局：标题24pt/副标题15pt/大数字34pt/落点策略①②③首句加粗
逐页：P4策略①②③、P5标题、P8测品履约率留空、P9重建柱状图、P10/P11标题+卡片宽度、
      P11重建对比图、P12删履约率卡重排、P15环形图标签修复
"""
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.enum.shapes import MSO_SHAPE

INK   = RGBColor.from_string('1F1F1F')
BODY  = RGBColor.from_string('3D3D3D')
MUTE  = RGBColor.from_string('9A9A9A')
ACC   = RGBColor.from_string('B06579')
LINE_C= RGBColor.from_string('EAE2E5')
GRAY_D= RGBColor.from_string('5A5A5A')
FONT  = '微软雅黑'

SRC = r'C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\work\v8.pptx'
OUT = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9.pptx'

prs = Presentation(SRC)

def set_box_text(sh, text, size, bold, color, align=None):
    tf = sh.text_frame; tf.clear()
    p = tf.paragraphs[0]
    if align is not None: p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color; r.font.name = FONT

def rebuild_strategy(sh, newtext=None):
    tf = sh.text_frame
    full = newtext if newtext else tf.text
    body = full.replace('落点策略', '').strip('　 ').strip()
    tf.clear()
    p = tf.paragraphs[0]
    def add(t, bold, color):
        r = p.add_run(); r.text = t
        r.font.size = Pt(14); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = FONT
    add('落点策略', True, ACC); add('　', False, BODY)
    for seg in [s for s in re.split(r'(?=[②③④])', body) if s]:
        m = re.search(r'[，、：（(]', seg)
        bi = m.end() if m else len(seg)
        add(seg[:bi], True, INK)
        if bi < len(seg): add(seg[bi:], False, BODY)

def add_col_chart(slide, x, y, w, h, cats, series, colors, fmt='General', gap=80):
    cd = CategoryChartData(); cd.categories = cats
    for (nm, vals) in series: cd.add_series(nm, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                                Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(10); ch.legend.font.name = FONT
    pl = ch.plots[0]; pl.gap_width = gap
    for i, c in enumerate(colors):
        se = pl.series[i]
        se.format.fill.solid(); se.format.fill.fore_color.rgb = c
        se.format.line.color.rgb = c
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.show_value = True; dl.number_format = fmt; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(10); dl.font.bold = True; dl.font.name = FONT; dl.font.color.rgb = INK
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE_C
    va.major_gridlines.format.line.width = Pt(0.5)
    va.tick_labels.font.size = Pt(10); va.tick_labels.font.name = FONT
    va.tick_labels.font.color.rgb = MUTE
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(11); ca.tick_labels.font.name = FONT
    ca.tick_labels.font.color.rgb = GRAY_D
    ca.format.line.color.rgb = LINE_C
    return ch

def delete_shape(sh):
    sh._element.getparent().remove(sh._element)

# 内容页（1-based 页码）
CONTENT = [4,5,6,8,9,10,11,12,13,15,16,17,19,20]

# 逐页标题/副标题（页码 -> (主标题, 副标题)）
TITLE_SUB = {
 5:  ('9 月常卖款月度出单',
      '2190 领跑、2189 补量成功 ｜ 2190 三连升登顶、2189 三连升翻倍、2178 回落、28173 高点回落'),
 9:  ('2178 + 2190｜少寄高转化 vs 回升缺量',
      '→ 差异在款与达人匹配效率 ｜ 2190 以最少寄样做出最高出单率与最多订单；2178 出单率回升但量没跟上'),
 10: ('2189 + 28173｜内容转化强 vs 达人承接强',
      '→ 一个缺供给、一个可放量 ｜ 2189 视频出单率老品最高、月度三连升；28173 寄样逐月加码、出单稳定'),
 11: ('2881 + 28102｜同为测款，一强一弱',
      '→ 差异在内容转化（质），非量 ｜ 2881 质（内容转化）更强、加投可放大；28102 缺市场信号、淘汰退出'),
 12: ('2L06｜测款数据全维度健康',
      '→ 款型有真实需求，为明年春夏备货 ｜ 转化 / 自然承接均为测品最优'),
}
STRATEGY_NEW = {4: '落点策略　① 延续爆品扩量 ② 秋冬当季主推款放量 ③ 巩固 Q4 开局'}

# ---------------- 全局排版 ----------------
for no in CONTENT:
    s = prs.slides[no-1]
    for sh in list(s.shapes):
        if sh.left is None: continue
        L = sh.left/914400; T = sh.top/914400; W = sh.width/914400
        if not sh.has_text_frame: continue
        txt = sh.text_frame.text.strip()
        if not txt: continue
        # 主标题：顶部左侧大字 -> 24pt 加粗近黑
        if 0.15 <= T <= 0.6 and L < 2 and W > 5 and 'PART' not in txt:
            if no in TITLE_SUB:
                set_box_text(sh, TITLE_SUB[no][0], 24, True, INK)
            else:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = INK
        # 副标题：PART 下说明行 -> 15pt 常规灰
        elif 1.15 <= T <= 1.30 and L < 2 and W > 5:
            if no in TITLE_SUB:
                set_box_text(sh, TITLE_SUB[no][1], 15, False, BODY)
            else:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(15); r.font.bold = False; r.font.color.rgb = BODY
        # 落点策略：①②③ 首句加粗
        elif txt.startswith('落点策略'):
            rebuild_strategy(sh, STRATEGY_NEW.get(no))
            continue
        # 大数字 38pt -> 34pt
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size is not None and r.font.size.pt == 38:
                    r.font.size = Pt(34)

# ---------------- s8 · 全景表：测品履约率留空 ----------------
s8 = prs.slides[7]
for sh in s8.shapes:
    if sh.has_table:
        tbl = sh.table
        for ri in range(1, len(tbl.rows)):
            sku = tbl.cell(ri, 0).text.strip()
            if sku in ('2881', '2L06', '28102'):
                c = tbl.cell(ri, 4)
                tf = c.text_frame; tf.clear()
                p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
                r = p.add_run(); r.text = '—'
                r.font.size = Pt(12); r.font.color.rgb = MUTE; r.font.name = FONT
        break
for sh in s8.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip().startswith('口径说明'):
        p0 = sh.text_frame.paragraphs[0]
        r = p0.add_run()
        r.text = '；测品（2881 / 2L06 / 28102）不统计履约率，留空标「—」'
        ref = p0.runs[0].font
        r.font.size = ref.size; r.font.bold = ref.bold
        try: r.font.color.rgb = ref.color.rgb
        except Exception: r.font.color.rgb = BODY
        r.font.name = FONT
        break

# ---------------- s9 · 重建柱状图 + 左卡文字宽度 ----------------
s9 = prs.slides[8]
for sh in list(s9.shapes):
    if sh.has_chart and abs(sh.left/914400-5.00) < 0.3 and abs(sh.top/914400-2.55) < 0.3:
        delete_shape(sh)
add_col_chart(s9, 5.00, 2.55, 3.33, 2.60,
              ['2190', '2178'],
              [('寄样量', [33, 46]), ('达人出单率(%)', [35.71, 17.95])],
              [GRAY_D, ACC], fmt='General', gap=80)
for sh in s9.shapes:
    if sh.has_text_frame and abs(sh.top/914400-3.09) < 0.06 and abs(sh.left/914400-0.86) < 0.06:
        sh.width = Inches(3.79)   # 原宽 6.37 溢出卡片

# ---------------- s10 · 左卡文字宽度 ----------------
s10 = prs.slides[9]
for sh in s10.shapes:
    if sh.has_text_frame and abs(sh.top/914400-3.09) < 0.06 and abs(sh.left/914400-0.86) < 0.06:
        sh.width = Inches(3.55)   # 原宽 5.17 溢出卡片

# ---------------- s11 · 重建对比图 ----------------
s11 = prs.slides[10]
for sh in list(s11.shapes):
    if sh.has_chart and abs(sh.left/914400-7.05) < 0.3 and abs(sh.top/914400-2.55) < 0.3:
        delete_shape(sh)
add_col_chart(s11, 7.05, 2.55, 5.68, 2.20,
              ['2881', '28102'],
              [('达人出单率(%)', [57.14, 2.70]), ('视频出单率(%)', [64.71, 2.78])],
              [ACC, GRAY_D], fmt='0.00"%"', gap=90)

# ---------------- s12 · 删履约率卡 + 重排两卡 + 文案去履约 ----------------
s12 = prs.slides[11]
card_fill = None; accent_c = None
band = []
for sh in list(s12.shapes):
    if sh.left is None: continue
    T = sh.top/914400; H = sh.height/914400
    if 2.50 < T < 3.70:
        if abs(H-1.10) < 0.05 and card_fill is None:
            try: card_fill = sh.fill.fore_color.rgb
            except Exception: pass
        if H < 0.10 and accent_c is None:
            try: accent_c = sh.fill.fore_color.rgb
            except Exception: pass
        band.append(sh)
for sh in band: delete_shape(sh)
if card_fill is None: card_fill = RGBColor.from_string('F9F5F6')
if accent_c is None:  accent_c  = RGBColor.from_string('C97F95')

def add_card12(x, label, value, sub):
    r1 = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.55), Inches(5.84), Inches(1.10))
    r1.fill.solid(); r1.fill.fore_color.rgb = card_fill; r1.line.fill.background(); r1.shadow.inherit = False
    r2 = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.55), Inches(5.84), Inches(0.06))
    r2.fill.solid(); r2.fill.fore_color.rgb = accent_c; r2.line.fill.background(); r2.shadow.inherit = False
    tb = s12.shapes.add_textbox(Inches(x+0.12), Inches(2.69), Inches(5.60), Inches(0.28))
    set_box_text(tb, label, 13, False, MUTE)
    tb = s12.shapes.add_textbox(Inches(x+0.12), Inches(2.95), Inches(5.60), Inches(0.44))
    set_box_text(tb, value, 28, True, ACC)
    tb = s12.shapes.add_textbox(Inches(x+0.12), Inches(3.39), Inches(5.60), Inches(0.26))
    set_box_text(tb, sub, 12, False, MUTE)

add_card12(0.60, '视频出单率', '30.00%', '内容侧健康')
add_card12(6.89, '自然单占比', '66.00%', '自然承接稳')

# 文案去除履约表述
for sh in s12.shapes:
    if not sh.has_text_frame: continue
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if '履约 88% 最高；' in r.text:
                r.text = r.text.replace('款型有真实需求；履约 88% 最高；', '款型有真实需求；')
            if '履约与转化双健康' in r.text:
                r.text = r.text.replace('履约与转化双健康', '转化与承接双健康')

# ---------------- s15 · 环形图：关自动标签 + 图例 ----------------
s15 = prs.slides[14]
for sh in s15.shapes:
    if sh.has_chart and 'DOUGHNUT' in str(sh.chart.chart_type):
        ch = sh.chart
        ch.plots[0].has_data_labels = False
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(9); ch.legend.font.name = FONT

# ---------------- s13 · 归因收束：字号/加粗归一 ----------------
s13 = prs.slides[12]
for sh in s13.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip().startswith('归因收束'):
        tf = sh.text_frame
        paras = [p.text for p in tf.paragraphs]
        tf.clear()
        for i, t in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            m = re.match(r'^([^：]{2,8}：)', t)
            head = m.group(1) if m else ''
            rest = t[len(head):]
            if head:
                r = p.add_run(); r.text = head
                r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = INK; r.font.name = FONT
            r = p.add_run(); r.text = rest
            r.font.size = Pt(14); r.font.bold = False; r.font.color.rgb = BODY; r.font.name = FONT
        break

prs.save(OUT)
print('SAVED:', OUT, 'slides:', len(prs.slides._sldIdLst))
