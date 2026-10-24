# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches

prs = Presentation(r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\Dlooda_8月月度汇报_v6.pptx")
slide = prs.slides[11]  # slide 12
for i, shp in enumerate(slide.shapes):
    if not shp.has_text_frame:
        continue
    t = shp.text_frame.text.strip()[:60].replace('\n',' ')
    print(f"{i}: ({shp.left/Inches(1):.2f},{shp.top/Inches(1):.2f},{shp.width/Inches(1):.2f},{shp.height/Inches(1):.2f}) | {t!r}")
