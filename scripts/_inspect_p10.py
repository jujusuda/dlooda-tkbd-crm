# -*- coding: utf-8 -*-
from pptx import Presentation
import re
F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_5.pptx'
prs = Presentation(F)
s = prs.slides[9]
def emu(v): return round(v/914400,2) if v is not None else None
print("=== P10 shapes (count=%d) ===" % len(s.shapes))
for i, sh in enumerate(s.shapes):
    L=emu(sh.left); T=emu(sh.top); W=emu(sh.width); H=emu(sh.height)
    extra=""
    if sh.has_text_frame: extra="TXT="+repr(sh.text_frame.text[:110].replace("\n"," / "))
    if sh.has_chart: extra="CHART="+str(sh.chart.chart_type)
    if sh.has_table: extra="TABLE rows=%d cols=%d"%(len(sh.table.rows), len(sh.table.columns))
    print(f"[{i}] L={L} T={T} W={W} H={H} {extra}")
for sh in s.shapes:
    if sh.has_chart:
        ch=sh.chart
        print("\n--- CHART", ch.chart_type, "---")
        try:
            print("cats:", ch.plots[0].categories)
            for se in ch.plots[0].series:
                print("  ", se.name, "->", [round(v,3) if isinstance(v,float) else v for v in se.values])
        except Exception as e:
            print("chart read err", e)
