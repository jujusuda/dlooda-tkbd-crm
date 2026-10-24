# -*- coding: utf-8 -*-
import zipfile, re, posixpath
from pptx import Presentation

F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_7.pptx'
prs = Presentation(F)
s = prs.slides[9]

print("=== P10 shape count:", len(s.shapes), "===")
# charts
charts = [sh for sh in s.shapes if sh.has_chart]
print("charts:", len(charts))
expected = {
 '达人出单率': [('2189',[56.00,22.86,31.11]),('28173',[30.19,33.78,31.58])],
 '出单达人':   [('2189',[14,8,14]),('28173',[16,25,24])],
 '点击率':     [('2189',[3.49,2.96,5.17]),('28173',[4.22,3.33,4.06])],
 '转化率':     [('2189',[2.37,2.08,1.13]),('28173',[2.11,1.73,0.90])],
}
ci = 0
for sh in charts:
    ch = sh.chart
    vals = {}
    for se in ch.plots[0].series:
        vals[se.name] = [round(v,2) for v in se.values]
    print(" chart", ci, "series:", vals)
    ci += 1

# table
tb = [sh for sh in s.shapes if sh.has_table][0].table
print("\n=== TABLE ===")
rows = []
for r in range(len(tb.rows)):
    rows.append([tb.cell(r,c).text for c in range(len(tb.columns))])
for row in rows:
    print(" | ".join(row))

# dLblPos legality on all line charts in the file
z = zipfile.ZipFile(F)
print("\n=== dLblPos on line charts (whole file) ===")
bad = []
for n in sorted(z.namelist()):
    if re.match(r'ppt/(slides/)?charts/chart\d+\.xml$', n):
        x = z.read(n).decode('utf-8','ignore')
        if '<c:lineChart>' in x:
            pos = re.findall(r'<c:dLblPos val="(\w+)"/>', x)
            il = re.findall(r'<c:lineChart>', x)
            print(" ", n, "dLblPos=", pos)
            for p in pos:
                if p in ('outEnd','inEnd'):
                    bad.append((n,p))
print("ILLEGAL line-chart dLblPos:", bad if bad else "none")

# package integrity
ct = z.read('[Content_Types].xml').decode('utf-8','ignore')
overrides = set(re.findall(r'PartName="([^"]+)"', ct))
defaults = set(re.findall(r'Extension="([^"]+)"', ct))
names = set(z.namelist())
probs = []
for o in overrides:
    if o.lstrip('/') not in names: probs.append('CT override missing: '+o)
for nm in names:
    if nm == '[Content_Types].xml': continue
    ext = nm.rsplit('.',1)[-1].lower()
    if ('/'+nm) not in overrides and ext not in defaults: probs.append('part not declared: '+nm)
    if nm.endswith('.rels'):
        base = posixpath.dirname(posixpath.dirname(nm))
        for rid, tgt in re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read(nm).decode('utf-8','ignore')):
            if tgt.startswith('http'): continue
            full = tgt.lstrip('/') if tgt.startswith('/') else posixpath.normpath(posixpath.join(base, tgt))
            if full not in names: probs.append(f'{nm}: {rid}->{tgt} MISSING {full}')
print("PACKAGE PROBLEMS:", probs if probs else "none")

# check no separate bottom strategy block on slide10 (no text starting with 落点策略)
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip().startswith('落点策略'):
        print("WARN: 独立落点策略块仍存在!")
