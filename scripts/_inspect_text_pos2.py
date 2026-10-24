from pptx import Presentation

SRC = r'E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-09\8be9710a0a0dd58ff4cb419316577cc7\8月余慧玲8月月度汇报(3).pptx'
p = Presentation(SRC)

WATCH = {'#C94F82', '#FCEAF1', '#7A6A70', '#B8AEB2', '#FFFFFF', '#E54C4C', '#4CAF50', '#E879A6', '#B25B7E'}
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        top = sh.top/914400 if sh.top is not None else None
        left = sh.left/914400 if sh.left is not None else None
        # shape fill behind
        bgr = ''
        try:
            if sh.fill.type == 1: bgr = '#%s' % str(sh.fill.fore_color.rgb)
        except Exception:
            pass
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                try:
                    if run.font.color and run.font.color.type is not None:
                        col = '#%s' % str(run.font.color.rgb)
                    else:
                        continue
                except Exception:
                    continue
                if col in WATCH and run.text.strip():
                    print('P%d top=%.2f L=%.2f bg=%s COL=%s TXT=%s' % (i, top, left, bgr, col, run.text.strip()[:30]))
