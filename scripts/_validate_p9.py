# -*- coding: utf-8 -*-
import zipfile, sys, posixpath
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(7).pptx"
OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_P9修改.pptx"

def c(t): return qn('c:' + t)

ok = lambda b: "OK " if b else "FAIL"

# ---- package integrity
z = zipfile.ZipFile(OUT)
print(ok(z.testzip() is None), "zip integrity")
bad = []
for n in z.namelist():
    if n.endswith('.rels'):
        base = posixpath.dirname(posixpath.dirname(n))      # <dir>/_rels/x.rels -> <dir>
        root = etree.fromstring(z.read(n))
        for rel in root:
            if rel.get('TargetMode') == 'External':
                continue
            p = posixpath.normpath(posixpath.join(base, rel.get('Target')))
            if p not in z.namelist():
                bad.append((n, rel.get('Target')))
print(ok(not bad), "rels resolve", bad[:5])

prs_src = Presentation(SRC)
prs = Presentation(OUT)
print(ok(len(prs.slides) == len(prs_src.slides)), f"slide count {len(prs.slides)} vs {len(prs_src.slides)}")

# ---- other slides untouched
diff = []
for i in range(len(prs.slides)):
    if i == 8:
        continue
    a = etree.tostring(prs.sliderefs[i] if False else prs.slides[i]._element)
    b = etree.tostring(prs_src.slides[i]._element)
    if a != b:
        diff.append(i + 1)
print(ok(not diff), "slides other than 9 unchanged", diff)

sl = prs.slides[8]
print("\n--- SLIDE 9 shapes ---")
for i, sh in enumerate(sl.shapes):
    box = (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)
    extra = ""
    if sh.has_chart:
        extra = "  <CHART>"
    elif sh.has_table:
        extra = "  <TABLE>"
    txt = (sh.text_frame.text if sh.has_text_frame else "").replace("\n", " | ")
    print(f"[{i:2d}] {sh.name:<22} L={sh.left:>8} T={sh.top:>8} W={sh.width:>8} H={sh.height:>8}{extra} {txt[:70]}")

# ---- combo chart structure
def combo_report(idx, label):
    sh = sl.shapes[idx]
    pa = sh.chart._chartSpace.plotArea
    order = [e.tag.split('}')[1] for e in pa]
    print(f"\n--- {label} plotArea order: {order}")
    bar = pa.find(c('barChart'))
    line = pa.find(c('lineChart'))
    valaxes = pa.findall(c('valAx'))
    cat = pa.find(c('catAx'))
    catId = cat.find(c('axId')).get('val')
    print("  catAx id", catId)
    for va in valaxes:
        aid = va.find(c('axId')).get('val')
        pos = va.find(c('axPos')).get('val')
        nf = va.find(c('numFmt'))
        print(f"  valAx id={aid} axPos={pos} crossAx={va.find(c('crossAx')).get('val')} "
              f"numFmt={nf.get('formatCode') if nf is not None else None} delete={va.find(c('delete')).get('val')}")
    def series_of(plot, tag):
        if plot is None:
            return None
        axids = [x.get('val') for x in plot.findall(c('axId'))]
        out = []
        for s in plot.findall(c('ser')):
            nm = s.find(c('tx')).find(c('strRef')).find(c('strCache')).find(c('pt')).find(c('v')).text
            vals = [pt.find(c('v')).text for pt in s.find(c('val')).find(c('numRef')).find(c('numCache')).findall(c('pt'))]
            pos = None
            for dl in s.findall(c('dLbls')):
                p = dl.find(c('dLblPos'))
                nf = dl.find(c('numFmt'))
                if p is not None:
                    pos = (p.get('val'), nf.get('formatCode') if nf is not None else None)
            out.append((nm, vals, pos, [x.get('val') for x in s.findall(c('axId'))]))
        return tag, axids, out
    for info in (series_of(bar, 'barChart'), series_of(line, 'lineChart')):
        tag, axids, rows = info
        print(f"  {tag} plot axIds={axids}")
        for nm, vals, pos, saxes in rows:
            print(f"      {nm:<22} vals={vals} label={pos} serAxIds={saxes}")

combo_report(19, "merged combo chart (shape 19)")

# ---- click / conversion: values + number format
for idx, label, expA, expB in [(6, "点击率", ['3.5','3.0','5.2'], ['4.2','3.3','4.1']),
                               (8, "转化率", ['2.4','2.1','1.1'], ['2.1','1.7','0.9'])]:
    pa = sl.shapes[idx].chart._chartSpace.plotArea
    lc = pa.find(c('lineChart'))
    sers = lc.findall(c('ser'))
    got = []
    fmts = []
    for s in sers:
        vals = [pt.find(c('v')).text for pt in s.find(c('val')).find(c('numRef')).find(c('numCache')).findall(c('pt'))]
        got.append(vals)
        nf = s.find(c('dLbls')).find(c('numFmt')) if s.find(c('dLbls')) is not None else None
        fmts.append(nf.get('formatCode') if nf is not None else None)
    print(f"\n{label}: {got}  numFmt={fmts}")
    print("   ", ok(got == [expA, expB]), "values = one-decimal target")

# ---- table unchanged
def table_dump(prs, idx):
    t = prs.slides[idx].shapes[13].table if False else None
    return None
for i, sh in enumerate(sl.shapes):
    if sh.has_table:
        rows = [[sh.table.cell(r, c).text for c in range(len(sh.table.columns))] for r in range(len(sh.table.rows))]
        print("\nTABLE (unchanged?) rows:")
        for r in rows:
            print("   ", r)
src_table = None
for sh in prs_src.slides[8].shapes:
    if sh.has_table:
        src_table = [[sh.table.cell(r, c).text for c in range(len(sh.table.columns))] for r in range(len(sh.table.rows))]
cur_table = None
for sh in sl.shapes:
    if sh.has_table:
        cur_table = [[sh.table.cell(r, c).text for c in range(len(sh.table.columns))] for r in range(len(sh.table.rows))]
print("   ", ok(src_table == cur_table), "table identical to source")

# ---- overlap check for top-row graphics
boxes = []
for sh in sl.shapes:
    if sh.top < 4400000 and sh.left > 500000:
        nm = sh.name
        txt = (sh.text_frame.text if sh.has_text_frame else nm)[:18]
        boxes.append((txt, sh.left, sh.top, sh.left + sh.width, sh.top + sh.height))
print("\noverlap check (graphics row):")
for i in range(len(boxes)):
    for j in range(i + 1, len(boxes)):
        n1, l1, t1, r1, b1 = boxes[i]
        n2, l2, t2, r2, b2 = boxes[j]
        if l1 < r2 and l2 < r1 and t1 < b2 and t2 < b1:
            print("   OVERLAP:", n1, n2)
print("   done")
