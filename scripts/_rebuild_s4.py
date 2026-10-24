# -*- coding: utf-8 -*-
"""Rebuild slide 4 (9月店铺整体表现) of the 9-month report per v-latest spec.
Keeps deck size, rose/豆沙粉 palette, white bg, 微软雅黑. Only modifies this page."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.chart.data import CategoryChartData

FONT = '微软雅黑'
INK   = RGBColor.from_string('1F1F1F')   # 标题近黑
BODY  = RGBColor.from_string('3D3D3D')   # 正文深灰
MUTE  = RGBColor.from_string('9A9A9A')   # 浅灰（副标题/标签）
ACC   = RGBColor.from_string('C97F95')   # 豆沙粉主色（图表/强调线/编号）
ACC_D = RGBColor.from_string('B06579')   # 深豆沙粉（KPI大数字/重点）
CARD  = RGBColor.from_string('F9F5F6')   # 极浅卡底
LINE  = RGBColor.from_string('E3DEE0')   # 细分隔线
LINE_P= RGBColor.from_string('E9DFE3')   # 卡片浅粉边框

F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_1.pptx'
prs = Presentation(F)
s = prs.slides[3]  # slide 4

# ---- remove all existing shapes ----
for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)

# ---- helpers ----
def rect(x, y, w, h, fill, line=None, line_w=Pt(1)):
    sp = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))  # 1=RECTANGLE
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = line_w
    return sp

def txt(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get('align', align)
        if 'space_after' in p: para.space_after = Pt(p['space_after'])
        if 'space_before' in p: para.space_before = Pt(p['space_before'])
        if 'line_spacing' in p: para.line_spacing = p['line_spacing']
        for (t, sz, bold, color) in p['runs']:
            r = para.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    return tb

def add_bar_chart(x, y, w, h, cats, vals, label_fmt):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series('', vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.has_data_labels = False  # default off; enable below
    plot = ch.plots[0]
    plot.gap_width = 90
    # single series color
    ser = plot.series[0]
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = ACC
    ser.format.line.fill.background()
    # data labels
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = label_fmt
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(11); dl.font.bold = True; dl.font.name = FONT; dl.font.color.rgb = BODY
    # value axis from 0
    va = ch.value_axis
    va.minimum_scale = 0
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor.from_string('EDE7E9')
    va.major_gridlines.format.line.width = Pt(0.75)
    va.has_title = False
    va.tick_labels.font.size = Pt(9); va.tick_labels.font.name = FONT; va.tick_labels.font.color.rgb = MUTE
    va.tick_labels.number_format = '#,##0'; va.tick_labels.number_format_is_linked = False
    # category axis
    ca = ch.category_axis
    ca.has_title = False
    ca.tick_labels.font.size = Pt(11); ca.tick_labels.font.name = FONT; ca.tick_labels.font.color.rgb = BODY
    ca.format.line.color.rgb = LINE
    ca.format.line.width = Pt(0.75)
    return ch

# ============ LAYER 1 · TITLE ============
rect(0.60, 0.42, 0.10, 0.66, ACC)                                  # accent vertical line
txt(0.86, 0.40, 9.6, 0.66,
    [{'runs': [('9 月店铺表现延续增长，Q3 GMV 与订单同步回升', 24, True, INK)]}],
    anchor=MSO_ANCHOR.MIDDLE)
txt(9.30, 0.48, 3.43, 0.50,
    [{'runs': [('PART 01 · 店铺整体情况', 13, False, MUTE)]}],
    align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
rect(0.60, 1.18, 12.13, 1/72, LINE)                               # divider

# ============ LAYER 2 · SUBTITLE + KPI CARDS ============
txt(0.86, 1.30, 11.8, 0.34,
    [{'runs': [('9 月 GMV 环比增长 3.35%，订单增长 3.66%；Q3 季度表现较 Q2 明显回升', 15, False, MUTE)]}])

# left card (GMV)
rect(0.60, 1.82, 5.96, 1.15, CARD, line=LINE_P)
txt(0.90, 1.98, 5.36, 0.30,
    [{'runs': [('9 月 GMV（USD）', 14, False, MUTE)]}])
txt(0.90, 2.24, 5.36, 0.58,
    [{'runs': [('$157,089.96', 38, True, ACC_D)]}], anchor=MSO_ANCHOR.MIDDLE)
txt(0.90, 2.86, 5.36, 0.22,
    [{'runs': [('环比 ', 13, False, MUTE), ('+3.35%', 13, True, ACC_D)]}])
# right card (订单)
rect(6.77, 1.82, 5.96, 1.15, CARD, line=LINE_P)
txt(6.77+0.30, 1.98, 5.36, 0.30,
    [{'runs': [('9 月订单（单）', 14, False, MUTE)]}])
txt(6.77+0.30, 2.24, 5.36, 0.58,
    [{'runs': [('7,514', 38, True, ACC_D)]}], anchor=MSO_ANCHOR.MIDDLE)
txt(6.77+0.30, 2.86, 5.36, 0.22,
    [{'runs': [('环比 ', 13, False, MUTE), ('+3.66%', 13, True, ACC_D)]}])

# ============ LAYER 3 · QUARTER CHARTS ============
txt(0.60, 3.05, 5.96, 0.28,
    [{'runs': [('季度 GMV（USD）', 14, True, BODY)]}])
txt(6.77, 3.05, 5.96, 0.28,
    [{'runs': [('季度订单（单）', 14, True, BODY)]}])
add_bar_chart(0.60, 3.32, 5.96, 2.55, ['Q1', 'Q2', 'Q3'], [365220, 338093, 425760], '#,##0')
add_bar_chart(6.77, 3.32, 5.96, 2.55, ['Q1', 'Q2', 'Q3'], [17290, 18108, 20421], '#,##0')

# Q3 vs Q2 line
txt(0.60, 5.96, 12.13, 0.30,
    [{'runs': [('Q3 较 Q2：GMV ', 14, False, BODY), ('+25.9%', 14, True, ACC_D),
               (' ｜ 订单 ', 14, False, BODY), ('+12.8%', 14, True, ACC_D)]}],
    align=PP_ALIGN.CENTER)

# ============ LAYER 4 · CORE FINDINGS ============
txt(0.60, 6.28, 6.0, 0.32,
    [{'runs': [('核心经营发现', 16, True, INK)]}])
col_w = 3.90
findings = [
    ('① ', 'Q3 季度回升：GMV 达 $425,760，订单达 20,421 单，均高于 Q1、Q2。'),
    ('② ', '9 月延续增长：GMV 与订单环比均实现正增长，但增幅低于 Q3 整体环比增幅。'),
    ('③ ', '后续分析重点：进一步拆解核心 SKU 的 GMV 贡献、环比变化及达人出单表现，识别主要增长驱动因素。'),
]
for i, (num, body) in enumerate(findings):
    x = 0.60 + i * (col_w + 0.20)
    txt(x, 6.58, col_w, 0.78,
        [{'runs': [(num, 14, True, ACC_D), (body, 13, False, BODY)],
          'line_spacing': 1.08, 'space_after': 0}],
        anchor=MSO_ANCHOR.TOP)

prs.save(F)
print('SAVED ->', F)
