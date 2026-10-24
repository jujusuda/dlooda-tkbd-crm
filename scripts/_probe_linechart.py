from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(7).pptx"
prs = Presentation(SRC)
slide = prs.slides[8]

cd = CategoryChartData()
cd.categories = ['7月','8月','9月']
for nm, v in [('A',(14,8,14)),('B',(16,25,24)),('C',(56.0,22.86,31.11)),
              ('D',(30.19,33.78,31.58)),('E',(33.33,28.41,27.85)),('F',(18.58,23.84,23.03))]:
    cd.add_series(nm, v)

gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.6), Inches(1.82), Inches(4.2), Inches(2.66), cd)
ch = gf.chart
plotArea = ch._chartSpace.plotArea
print("LINE_MARKERS plotArea children:")
for el in plotArea:
    tag = el.tag.split('}')[1]
    print("  ", tag)
    if tag.endswith('Chart'):
        for sub in el:
            stag = sub.tag.split('}')[1]
            print("      ", stag, dict(sub.attrib) if sub.attrib else '')
            if stag == 'ser':
                kids = [s.tag.split('}')[1] for s in sub]
                print("          ser children:", kids)
print()
# also check whether a legend exists and where
chart_el = ch._chartSpace.chart
print("chart element children:", [e.tag.split('}')[1] for e in chart_el])
valAx = plotArea.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}valAx')
print("valAx children:", [e.tag.split('}')[1] for e in valAx])
catAx = plotArea.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}catAx')
print("catAx children:", [e.tag.split('}')[1] for e in catAx])
