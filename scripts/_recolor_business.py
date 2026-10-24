# -*- coding: utf-8 -*-
"""Global color re-theme of the final PPT -> low-saturation gray-tone pale-pink business style.
Keeps ALL text, data, layout, fonts, positions, chart structure 100% unchanged; only colors change.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor

SRC = r'E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-09\8be9710a0a0dd58ff4cb419316577cc7\8月余慧玲8月月度汇报(3).pptx'
OUT = r'C:\Users\LENOVO\Desktop\8月余慧玲8月月度汇报_商务淡粉版.pptx'

def C(h): return RGBColor.from_string(h)

# ---- palette ----
MAIN       = 'B87C8D'  # title bars / table headers / strong blocks / section bg
MAIN_LIGHT = 'D7B2BE'  # secondary accent blocks (white text on them)
CARD       = 'F5EEEF'  # conclusion cards / alternating rows / light-pink fills
WHITE      = 'FFFFFF'
BODY       = '333333'  # all body / data / table text
CHART_MAIN = '9C5F70'  # chart main series
CHART_AUX  = 'E5C6D0'  # chart auxiliary series (light pink)
POS        = '739E88'  # up / positive  (gray-green)
NEG        = 'D48B76'  # down / warning (gray orange-pink)
BLUE_MUTED = '8FA3B0'  # 潜力培育款 block
GOLD_MUTED = 'C9A06A'  # 观察调整款 block
LINE_SOFT  = 'E5D9DD'  # soft pink-gray lines / borders

# fills that count as "colored" (text on them should be white/light)
COLORED_FILLS = {'#C94F82','#E879A6','#F5A0BC','#4CAF50','#4C8FC9','#E0A83C','#E54C4C'}

# shape/table fill remap (color -> target). F5A0BC & E879A6 need context (handled in code).
FILL_MAP = {
    '#C94F82': MAIN,
    '#FCEAF1': CARD, '#FDF1F5': CARD, '#FBE4EE': CARD, '#F3EFF0': CARD, '#FFF8FB': CARD,
    '#4CAF50': POS, '#4C8FC9': BLUE_MUTED, '#E0A83C': GOLD_MUTED, '#E54C4C': NEG,
}

def set_fill(fill, hexstr):
    try:
        fill.solid(); fill.fore_color.rgb = C(hexstr)
    except Exception:
        pass

def set_line(line, hexstr):
    try:
        if line.color and line.color.type is not None:
            line.color.rgb = C(hexstr)
    except Exception:
        pass

def set_run(run, hexstr):
    try:
        if run.font.color and run.font.color.type is not None:
            run.font.color.rgb = C(hexstr)
    except Exception:
        pass

def text_target(hexcolor, own_colored, slide_colored):
    c = hexcolor
    if c == '#FFFFFF':
        return WHITE
    if c == '#FCEAF1':
        return WHITE if (own_colored or slide_colored) else MAIN
    if c in ('#C94F82', '#E879A6', '#B25B7E'):
        return WHITE if (own_colored or slide_colored) else MAIN
    if c == '#E54C4C':
        return NEG
    if c == '#D6336C':   # vivid magenta down-arrows (P11 table) -> warning
        return NEG
    if c == '#4CAF50':
        return POS
    if c in ('#7A6A70', '#4A3D42', '#4A3A41', '#B8AEB2'):
        return BODY
    return None

def recolor_runs(tf, own_colored, slide_colored, force_white=False):
    for para in tf.paragraphs:
        for run in para.runs:
            try:
                if run.font.color and run.font.color.type is not None:
                    cur = '#' + str(run.font.color.rgb)
                else:
                    continue
            except Exception:
                continue
            if force_white:
                set_run(run, WHITE); continue
            t = text_target(cur, own_colored, slide_colored)
            if t:
                set_run(run, t)

def shape_fill_color(sh):
    try:
        if sh.fill.type == 1:
            return '#' + str(sh.fill.fore_color.rgb)
    except Exception:
        pass
    return None

p = Presentation(SRC)
fc = lc = tc = hc = 0

for s in p.slides:
    # slide background
    try:
        if s.background.fill.type == 1:
            bg = '#' + str(s.background.fill.fore_color.rgb)
            if bg == '#FFF6F9':
                set_fill(s.background.fill, WHITE); fc += 1
            elif bg == '#C94F82':
                set_fill(s.background.fill, MAIN); fc += 1
    except Exception:
        pass
    slide_colored = False
    try:
        if s.background.fill.type == 1 and ('#'+str(s.background.fill.fore_color.rgb)) in COLORED_FILLS:
            slide_colored = True
    except Exception:
        pass

    for sh in s.shapes:
        own_colored = False
        sfc = shape_fill_color(sh)
        if sfc in COLORED_FILLS:
            own_colored = True

        # ---- shape fill ----
        if sfc:
            if sfc in FILL_MAP:
                set_fill(sh.fill, FILL_MAP[sfc]); fc += 1
            elif sfc == '#E879A6':
                set_fill(sh.fill, MAIN_LIGHT); fc += 1
            elif sfc == '#F5A0BC':
                # white-text block -> accent; otherwise divider line
                has_white = False
                if sh.has_text_frame:
                    for para in sh.text_frame.paragraphs:
                        for run in para.runs:
                            try:
                                if run.font.color and run.font.color.type is not None and '#'+str(run.font.color.rgb) == '#FFFFFF':
                                    has_white = True
                            except Exception:
                                pass
                set_fill(sh.fill, MAIN_LIGHT if has_white else LINE_SOFT); fc += 1

        # ---- shape line ----
        try:
            if sh.line and sh.line.color and sh.line.color.type == 1:
                lc0 = '#' + str(sh.line.color.rgb)
                if lc0 in ('#FCD5E3', '#F5A0BC'):
                    set_line(sh.line, LINE_SOFT); lc += 1
        except Exception:
            pass

        # ---- text ----
        if sh.has_text_frame:
            recolor_runs(sh.text_frame, own_colored, slide_colored)
            tc += 1

        # ---- table ----
        if sh.has_table:
            tbl = sh.table
            for ri, row in enumerate(tbl.rows):
                is_header = (ri == 0)
                for cell in row.cells:
                    # cell fill
                    try:
                        if cell.fill.type == 1:
                            cf = '#' + str(cell.fill.fore_color.rgb)
                            if is_header:
                                set_fill(cell.fill, MAIN); hc += 1
                            elif cf in ('#FCEAF1', '#FDF1F5', '#FBE4EE', '#F3EFF0', '#FFF8FB'):
                                set_fill(cell.fill, CARD); fc += 1
                            elif cf in ('#C94F82', '#E879A6', '#F5A0BC', '#4CAF50', '#4C8FC9', '#E0A83C', '#E54C4C'):
                                set_fill(cell.fill, MAIN); fc += 1
                    except Exception:
                        pass
                    # cell borders
                    try:
                        for edge in ('top', 'left', 'bottom', 'right'):
                            b = getattr(cell.borders, edge)
                            if b is not None:
                                try:
                                    if b.color and b.color.type == 1:
                                        b.color.rgb = C(LINE_SOFT)
                                except Exception:
                                    pass
                    except Exception:
                        pass
                    # cell text
                    recolor_runs(cell.text_frame, own_colored=False, slide_colored=False, force_white=is_header)

        # ---- chart ----
        if sh.has_chart:
            ch = sh.chart
            for ser in ch.series:
                try:
                    cur = None
                    if ser.format.fill.type == 1:
                        cur = '#' + str(ser.format.fill.fore_color.rgb)
                except Exception:
                    cur = None
                tgt = None
                if cur == '#C94F82':
                    tgt = CHART_MAIN
                elif cur in ('#FCD5E3', '#F5A0BC', '#B25B7E'):
                    tgt = CHART_AUX
                elif cur == '#E54C4C':
                    tgt = NEG
                else:
                    nm = (ser.name or '')
                    if '2722' in nm or '8月' in nm or '单产' in nm:
                        tgt = CHART_MAIN
                    elif '2757' in nm or '测款' in nm or '覆盖' in nm or '7月' in nm or '样本' in nm:
                        tgt = CHART_AUX
                    else:
                        tgt = CHART_MAIN
                if tgt:
                    try:
                        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = C(tgt); fc += 1
                    except Exception:
                        pass
                # data labels -> body
                try:
                    if ser.data_labels and ser.data_labels.font:
                        ser.data_labels.font.color.rgb = C(BODY)
                except Exception:
                    pass
            # axes font
            for ax in (ch.category_axis, ch.value_axis):
                try:
                    if ax.tick_labels and ax.tick_labels.font:
                        ax.tick_labels.font.color.rgb = C(BODY)
                except Exception:
                    pass
            # chart title
            try:
                if ch.chart_title and ch.chart_title.text_frame:
                    recolor_runs(ch.chart_title.text_frame, own_colored=False, slide_colored=False)
            except Exception:
                pass

p.save(OUT)
print('SAVED:', OUT)
print('fill=%d line=%d text_shapes=%d header_cells=%d' % (fc, lc, tc, hc))
