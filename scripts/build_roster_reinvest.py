# -*- coding: utf-8 -*-
"""
从「每日工作.xlsx」提取名单库「竞品达人+爆单达人复投名单」(模块2) 数据。
来源 sheet:
  - 复投        : 937 条结构化复投记录 (A日期/B类型/C ID/D sku/E邮箱/F复投方式/G是否成功/H原因/I等级L/J备注/K话术/L话术)
  - 复投成功     : 复投成功矩阵 (每4列一组: sku / 视频·出单 / 是否 / ·)，汇总哪些 handle 在哪些 sku 复投成功

合并单元格导致部分字段只在组首行有值 -> 向下填充(forward-fill)还原。
输出: js/reinvest-data.js -> window.DloodaReinvest = { records:[...], success:{sku:[handles]} }
"""
import openpyxl, re, json

SRC = "scripts/source_xlsx/每日工作_2026-10-25.xlsx"
OUT = "js/reinvest-data.js"

HANDLE_RE = re.compile(r'^[A-Za-z0-9_.][A-Za-z0-9._]{1,40}$')
def norm(h):
    if h is None: return None
    s = str(h).strip().lstrip('@').strip().lower()
    s = s.split('\n')[0].split(' ')[0]
    if not s or not HANDLE_RE.match(s):
        return None
    return s

def ffill_col(ws, col, r0, r1):
    """对单列做向下填充, 返回 list(长度=r1-r0+1, 索引0对应r0)"""
    out = []
    last = None
    for r in range(r0, r1 + 1):
        v = ws.cell(row=r, column=col).value
        if v not in (None, ""):
            last = v
        out.append(last)
    return out

def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb['复投']
    nrows = ws.max_row

    # 逐列 forward-fill (日期/类型/是否成功/等级/复投方式/备注 等可能合并)
    fill_cols = {1: 'date', 2: 'type', 4: 'sku', 5: 'email',
                 6: 'method', 7: 'success', 8: 'reason', 9: 'level', 10: 'note',
                 11: 'script', 12: 'script2'}
    filled = {}
    for col, key in fill_cols.items():
        filled[key] = ffill_col(ws, col, 2, nrows)

    records = []
    for r in range(2, nrows + 1):
        cid = ws.cell(row=r, column=3).value  # C=ID
        if cid in (None, ""):
            continue
        h = norm(cid)
        if not h:
            continue
        rec = {"handle": h}
        for col, key in fill_cols.items():
            if key == 'script2':  # L列话术合并到 script
                continue
            v = filled[key][r - 2]
            if v not in (None, ""):
                if key in ('script',):
                    rec[key] = str(v)
                else:
                    rec[key] = str(v).strip()
        # script2 (L) 若独立非空也追加
        v2 = filled['script2'][r - 2]
        if v2 not in (None, "") and str(v2).strip():
            rec['script'] = (rec.get('script', '') + "\n" + str(v2).strip()).strip()
        records.append(rec)

    # 加稳定 id (r0, r1, ...) 供本机标记定位
    for i, rec in enumerate(records):
        rec["id"] = "r%d" % i

    # 复投成功汇总
    success = {}
    if '复投成功' in wb.sheetnames:
        wss = wb['复投成功']
        c = 1
        while c <= wss.max_column:
            sku = wss.cell(row=1, column=c).value
            if sku in (None, ""):
                c += 4
                continue
            sku = str(sku).strip()
            if not re.match(r'^[A-Za-z0-9]{2,10}$', sku):  # 只保留 SKU 列(排除"已添加WhatsApp"等说明列)
                c += 4
                continue
            hs = []
            for r in range(2, wss.max_row + 1):
                h = norm(wss.cell(row=r, column=c).value)
                if h:
                    hs.append(h)
            if hs:
                success[sku] = sorted(set(hs))
            c += 4

    total_script_chars = sum(len(r.get('script', '')) for r in records)
    out = ("// 自动生成：竞品达人+爆单达人复投名单（来源：每日工作.xlsx 复投 + 复投成功）\n"
           "// 结构: records[] = {id, date, type, sku, email, method, success, reason, level, note, script}\n"
           "//       success{} = { sku: [handles] } 复投成功汇总\n"
           "window.DloodaReinvest = " +
           json.dumps({"records": records, "success": success,
                       "meta": {"recordCount": len(records), "successSku": len(success)}},
                      ensure_ascii=False, indent=1) + ";\n")
    with open(OUT, "w", encoding="utf-8") as f:
        f.write(out)
    print("写出 %s" % OUT)
    print("  复投记录: %d 条" % len(records))
    print("  复投成功 sku 组: %d" % len(success))
    for k, v in list(success.items())[:6]:
        print("    %s : %d 人" % (k, len(v)))
    print("  话术总字符: %d (≈%.1f KB)" % (total_script_chars, total_script_chars / 1024))
    # 是否成功分布
    from collections import Counter
    cnt = Counter(r.get('success', '') for r in records)
    print("  是否成功分布:", dict(cnt))
    cntm = Counter(r.get('method', '') for r in records)
    print("  复投方式Top:", cntm.most_common(5))

if __name__ == "__main__":
    main()
