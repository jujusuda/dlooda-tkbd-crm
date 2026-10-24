from pptx import Presentation
from pptx.dml.color import RGBColor

SRC = r'E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-09\8be9710a0a0dd58ff4cb419316577cc7\8月余慧玲8月月度汇报(2).pptx'
OUT = r'C:\Users\LENOVO\Desktop\8月余慧玲8月月度汇报_淡粉最终版.pptx'

p = Presentation(SRC)

# ---- 填充映射：整体压成浅粉，保留层级，不雾 ----
FILL_MAP = {
    '#C94F82': '#DDA0BD',   # 标题栏/强色 -> 浅玫瑰
    '#F5A0BC': '#F7B6C8',   # 卡片/图表柱 -> 柔粉(仍可见)
    '#E879A6': '#F2AEC0',   # 次粉 -> 柔粉
    '#B25B7E': '#C587A1',   # 深玫变体 -> 柔化
    '#FCEAF1': '#FCEFF4',   # 浅粉底 -> 更淡一档
    '#FCD5E3': '#FCE3EB',   # 浅粉 -> 柔
}
# 这些分类/中性色不动
KEEP_FILL = {'#4C8FC9', '#E54C4C', '#4CAF50', '#E0A83C', '#FFFFFF',
             '#FFF8FB', '#FDF1F5', '#FBE4EE', '#F3EFF0'}

# ---- 文字映射：白/浅粉字 -> 深玫瑰(防雾关键)；浅灰加深 ----
TEXT_MAP = {
    '#FFFFFF': '#8E3E5C',
    '#FCEAF1': '#8E3E5C',
    '#B8AEB2': '#968A90',
}
# 保留：#C94F82 强调字、#4A3D42/#7A6A70 正文、#E54C4C 红、#4CAF50 绿

def recolor_fill(c):
    h = '#' + str(c)
    if h in FILL_MAP:
        return RGBColor.from_string(FILL_MAP[h][1:])
    return None

def recolor_text(c):
    h = '#' + str(c)
    if h in TEXT_MAP:
        return RGBColor.from_string(TEXT_MAP[h][1:])
    return None

fc = tc = pc = 0

for idx, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        # 1) 形状填充
        try:
            if sh.fill.type == 1:
                nc = recolor_fill(sh.fill.fore_color.rgb)
                if nc is not None:
                    sh.fill.fore_color.rgb = nc
                    fc += 1
        except Exception:
            pass
        # 2) 表格单元格填充
        if sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    try:
                        if cell.fill.type == 1:
                            nc = recolor_fill(cell.fill.fore_color.rgb)
                            if nc is not None:
                                cell.fill.fore_color.rgb = nc
                                fc += 1
                    except Exception:
                        pass
        # 3) 图表系列填充
        if sh.has_chart:
            for ser in sh.chart.series:
                try:
                    if ser.format.fill.type == 1:
                        nc = recolor_fill(ser.format.fill.fore_color.rgb)
                        if nc is not None:
                            ser.format.fill.fore_color.rgb = nc
                            fc += 1
                except Exception:
                    pass
        # 4) 文字颜色（防雾：白/浅粉 -> 深玫瑰）
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        if run.font.color and run.font.color.type is not None:
                            nc = recolor_text(run.font.color.rgb)
                            if nc is not None:
                                run.font.color.rgb = nc
                                tc += 1
                    except Exception:
                        pass
        # 5) 右下角页码：统一为真实页序
        top = sh.top / 914400 if sh.top is not None else None
        left = sh.left / 914400 if sh.left is not None else None
        if sh.has_text_frame and top is not None and left is not None and top > 6.0 and left > 9.0:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if run.text.strip().isdigit():
                        run.text = str(idx)
                        pc += 1

p.save(OUT)
print('SAVED:', OUT)
print('填充改 %d 处 | 文字改 %d 处 | 页码改 %d 处' % (fc, tc, pc))

# ---- 校验：打印最终页码 + 残留粉/白 ----
p2 = Presentation(OUT)
print('\n=== 校验：右下角页码 ===')
for i, s in enumerate(p2.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        top = sh.top / 914400 if sh.top is not None else None
        left = sh.left / 914400 if sh.left is not None else None
        if top is not None and left is not None and top > 6.0 and left > 9.0:
            t = ' '.join(pp.text for pp in sh.text_frame.paragraphs if pp.text.strip())
            print('  P%d -> [%s]' % (i, t))
