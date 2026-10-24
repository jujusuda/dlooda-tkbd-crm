# -*- coding: utf-8 -*-
"""Slide 9 (2189 + 28173) of v9_1(9): replace three pasted images with
   native, fully-editable PowerPoint objects.

   - merged combo chart: clustered COLUMNS = 出单达人数量 (LEFT axis, 位)
                         LINES = 达人出单率 (RIGHT axis, %)
   - two causal-chain cards rebuilt as real shapes + text
"""
import copy
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(9).pptx"
OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_v9_1(9)_P9修改.pptx"

ROSE   = "B06579"   # 2189
SAND   = "D9A5B5"   # 28173
INK    = "1F1F1F"
LBL    = "8A4A5C"   # data-label text (dark rose)
FONT   = "微软雅黑"

def c(t): return qn('c:' + t)
def a(t): return qn('a:' + t)
def p(t): return qn('p:' + t)
def el(x): return parse_xml(x)


# ----------------------------------------------------------------- chart pieces
def _dlbls(fmt, pos, sz, color):
    f = fmt.replace('"', '&quot;')
    return el(
        '<c:dLbls ' + nsdecls('c', 'a') + '>'
        f'<c:numFmt formatCode="{f}" sourceLinked="0"/>'
        '<c:spPr>'
        '<a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>'
        '<a:ln w="9525"><a:solidFill><a:srgbClr val="E2C9D2"/></a:solidFill></a:ln>'
        '</c:spPr>'
        '<c:txPr>'
        '<a:bodyPr rot="0" spcFirstLastPara="0" vertOverflow="ellipsis" vert="horz" wrap="square" '
        'lIns="38100" tIns="9525" rIns="38100" bIns="9525" anchor="ctr" anchorCtr="1"/>'
        '<a:lstStyle/>'
        f'<a:p><a:pPr><a:defRPr lang="zh-CN" sz="{sz}" b="0" i="0" u="none" strike="noStrike" kern="1200" baseline="0">'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:latin typeface="{FONT}" charset="-122"/><a:ea typeface="{FONT}"/><a:cs typeface="+mn-cs"/>'
        '</a:defRPr></a:pPr></a:p>'
        '</c:txPr>'
        f'<c:dLblPos val="{pos}"/>'
        '<c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/>'
        '<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/>'
        '<c:showLeaderLines val="0"/>'
        '</c:dLbls>'
    )


def _bar_spPr(fill):
    return el(
        '<c:spPr ' + nsdecls('c', 'a') + '>'
        f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>'
        '<a:ln><a:noFill/></a:ln>'
        '</c:spPr>'
    )


def _line_spPr(color, dash):
    return el(
        '<c:spPr ' + nsdecls('c', 'a') + '>'
        '<a:ln w="28575" cap="rnd" cmpd="sng" algn="ctr">'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:prstDash val="{dash}"/><a:round/>'
        '</a:ln></c:spPr>'
    )


def _marker(color, symbol, size=7):
    return el(
        '<c:marker ' + nsdecls('c', 'a') + '>'
        f'<c:symbol val="{symbol}"/><c:size val="{size}"/>'
        '<c:spPr>'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        '<a:ln w="12700" cap="flat" cmpd="sng" algn="ctr">'
        '<a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>'
        '<a:prstDash val="solid"/><a:round/></a:ln>'
        '</c:spPr></c:marker>'
    )


def _numfmt(code):
    return el(f'<c:numFmt {nsdecls("c")} formatCode="{code.replace(chr(34), "&quot;")}" sourceLinked="0"/>')


def _axis_title(text, sz=800, color="5A5A5A", rot="-5400000"):
    return el(
        '<c:title ' + nsdecls('c', 'a') + '>'
        '<c:tx><c:rich>'
        f'<a:bodyPr rot="{rot}" spcFirstLastPara="0" vertOverflow="ellipsis" vert="horz" wrap="square" anchor="ctr" anchorCtr="0"/>'
        '<a:lstStyle/>'
        '<a:p><a:pPr><a:defRPr lang="zh-CN" sz="' + str(sz) + '" b="0">'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:latin typeface="{FONT}" charset="-122"/><a:ea typeface="{FONT}"/><a:cs typeface="+mn-cs"/>'
        '</a:defRPr></a:pPr>'
        f'<a:r><a:rPr lang="zh-CN" sz="{sz}" b="0"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:latin typeface="{FONT}" charset="-122"/><a:ea typeface="{FONT}"/></a:rPr>'
        f'<a:t>{text}</a:t></a:r></a:p>'
        '</c:rich></c:tx>'
        '<c:layout/><c:overlay val="0"/>'
        '</c:title>'
    )


def build_combo(slide):
    plan = [
        ('2189 出单达人数量',  (14, 8, 14),            'bar'),
        ('28173 出单达人数量', (16, 25, 24),           'bar'),
        ('2189 达人出单率',    (56.00, 22.86, 31.11),  'line'),
        ('28173 达人出单率',   (30.19, 33.78, 31.58),  'line'),
    ]
    cd = CategoryChartData()
    cd.categories = ['7月', '8月', '9月']
    for nm, v, k in plan:
        cd.add_series(nm, v)

    L, T, W, H = 127000, 1191260, 5934075, 2684145
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, L, T, W, H, cd)
    chart = gf.chart

    # --- native title (inside the chart, editable)
    chart.has_title = True
    ct = chart.chart_title.text_frame
    ct.text = '出单率与出单达人数量（7–9月）'
    r0 = ct.paragraphs[0].runs[0]
    r0.font.size = Pt(11)
    r0.font.bold = True
    r0.font.name = FONT
    r0.font.color.rgb = RGBColor.from_string(LBL)

    # --- native legend, one row at the bottom
    chart.has_legend = True
    lg = chart.legend
    lg.position = XL_LEGEND_POSITION.BOTTOM
    lg.include_in_layout = False
    lg.font.size = Pt(8)
    lg.font.name = FONT

    plotArea = chart._chartSpace.plotArea
    lineChart = plotArea.find(c('lineChart'))
    catAx = plotArea.find(c('catAx'))
    valAx = plotArea.find(c('valAx'))          # primary = LEFT = 位 (columns)
    catId = catAx.find(c('axId')).get('val')
    valId = valAx.find(c('axId')).get('val')
    secId = str(int(valId) + 23)

    sers = lineChart.findall(c('ser'))
    bar_sers, line_sers = sers[0:2], sers[2:4]

    for s in bar_sers:
        lineChart.remove(s)
        sm = s.find(c('smooth'))
        if sm is not None:
            s.remove(sm)

    # columns plot on the PRIMARY (left) axis, drawn behind the lines
    bc = el('<c:barChart ' + nsdecls('c') + '>'
            '<c:barDir val="col"/>'
            '<c:grouping val="clustered"/>'
            '<c:varyColors val="0"/>'
            f'<c:axId val="{catId}"/><c:axId val="{valId}"/>'
            '</c:barChart>')
    first_ax = bc.findall(c('axId'))[0]
    for s in bar_sers:
        first_ax.addprevious(s)
    first_ax.addprevious(el(f'<c:gapWidth {nsdecls("c")} val="90"/>'))
    first_ax.addprevious(el(f'<c:overlap {nsdecls("c")} val="-20"/>'))
    lineChart.addprevious(bc)

    # lines ride the SECONDARY (right) axis
    lineChart.findall(c('axId'))[1].set('val', secId)

    # --- secondary value axis (RIGHT, %) carried by the lines
    sec = copy.deepcopy(valAx)
    sec.find(c('axId')).set('val', secId)
    sec.find(c('axPos')).set('val', 'r')
    mg = sec.find(c('majorGridlines'))
    if mg is not None:
        sec.remove(mg)
    sec.find(c('crossAx')).set('val', catId)
    sec.find(c('delete')).set('val', '0')
    mtm = sec.find(c('majorTickMark'))
    mtm.addprevious(_axis_title('达人出单率（%）'))
    mtm.addprevious(_numfmt('0'))
    valAx.addnext(sec)

    # --- primary axis: counts, keeps its gridlines, gets a title
    vmtm = valAx.find(c('majorTickMark'))
    vmtm.addprevious(_axis_title('出单达人数量（位）', color="5A5A5A"))
    vmtm.addprevious(_numfmt('0'))

    # --- series formatting
    for ser, fill in zip(bar_sers, (ROSE, SAND)):
        cat = ser.find(c('cat'))
        cat.addprevious(_bar_spPr(fill))
        cat.addprevious(_dlbls('0', 'outEnd', 900, INK))

    for ser, color, dash, symbol, pos in zip(
            line_sers,
            (ROSE, SAND),
            ('solid', 'dash'),
            ('circle', 'square'),
            ('t', 'b')):
        cat = ser.find(c('cat'))
        cat.addprevious(_line_spPr(color, dash))
        cat.addprevious(_marker(color, symbol))
        cat.addprevious(_dlbls('0.00"%"', pos, 900, LBL))
    return gf


# ----------------------------------------------------------------- card pieces
def _rect(slide, l, t, w, h, fill, line=None, rounded=None, dash=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded is not None else MSO_SHAPE.RECTANGLE,
        l, t, w, h)
    if rounded is not None:
        try:
            shp.adjustments[0] = rounded
        except Exception:
            pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        shp.line.color.rgb = RGBColor.from_string(line)
        shp.line.width = Pt(0.75)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _tri(slide, cx, top, size=115000):
    shp = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                 int(cx - size / 2), top, size, int(size * 0.72))
    shp.rotation = 180
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string("C97F95")
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _runs(tf, parts, size, color=INK):
    para = tf.paragraphs[0]
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    for text, bold in parts:
        run = para.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = FONT
        run.font.color.rgb = RGBColor.from_string(color)
    para.alignment = PP_ALIGN.LEFT


def build_card(slide, left, top, width, height, title_parts, rows):
    PADX = 130000
    # panel
    _rect(slide, left, top, width, height, "FBF1F4", line="EFD3DC", rounded=0.045)

    # title
    tb = slide.shapes.add_textbox(left + PADX, top + 70000, width - 2 * PADX, 280000)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _runs(tf, title_parts, 11.5, "7A3B4E")

    # rows
    n = len(rows)
    area_top = top + 400000
    area_bot = top + height - 90000
    gap = 120000
    total = area_bot - area_top
    rh = int((total - gap * (n - 1)) / n)
    box_l = left + 95000
    box_w = width - 190000
    for i, parts in enumerate(rows):
        ry = area_top + i * (rh + gap)
        is_last = (i == n - 1)
        shp = _rect(slide, box_l, ry, box_w, rh,
                    "FCEFF3" if is_last else "FEFAFB",
                    line="C97F95" if is_last else "E7C4D0",
                    rounded=0.22, dash=is_last)
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Emu(120000)
        tf.margin_right = Emu(90000)
        tf.margin_top = Emu(20000)
        tf.margin_bottom = Emu(20000)
        _runs(tf, parts, 9)
        if not is_last:
            _tri(slide, left + width // 2, ry + rh + int(gap * 0.08))


# ----------------------------------------------------------------- main
prs = Presentation(SRC)
slide = prs.slides[8]

pics = [sh for sh in slide.shapes if sh.shape_type == 13]     # 图片0/1/2
pic_els = [sh._element for sh in pics]

# 1) one-decimal click / conversion (values + label format)
def patch_one_decimal(chart, va, vb):
    pa = chart._chartSpace.plotArea
    lc = pa.find(c('lineChart'))
    for ser, vals in zip(lc.findall(c('ser')), (va, vb)):
        cache = ser.find(c('val')).find(c('numRef')).find(c('numCache'))
        for pt, v in zip(cache.findall(c('pt')), vals):
            pt.find(c('v')).text = str(v)
        for holder in (lc, ser):
            for dl in holder.findall(c('dLbls')):
                nf = dl.find(c('numFmt'))
                if nf is not None:
                    nf.set('formatCode', '0.0"%"')
                    nf.set('sourceLinked', '0')

patch_one_decimal(slide.shapes[5].chart, (3.5, 3.0, 5.2), (4.2, 3.3, 4.1))    # 点击率
patch_one_decimal(slide.shapes[7].chart, (2.4, 2.1, 1.1), (2.1, 1.7, 0.9))    # 转化率

# 2) native combo chart (replaces 图片 0)
build_combo(slide)

# 3) native cards (replace 图片 1 / 图片 2)
build_card(
    slide, 475615, 3895725, 5454650, 2851150,
    [("2189", True), (" → 内容转化型（缺供给）", False)],
    [
        [("结果：", True), ("月度出单 ", False), ("176→723→999", True), (" 三连升", False)],
        [("内容端：", True), ("点击率 ", False), ("2.96→5.17", True), (" 升、点击量涨 → 内容在拉动", False)],
        [("承接端：", True), ("转化率 ", False), ("2.08→1.13", True), (" 降 → 点击多、成交没跟上", False)],
        [("达人端：", True), ("出单达人 ", False), ("14→8→14", True), ("、寄样缺 → 供给不足", False)],
        [("判断：", True), ("内容把量带起来了，但承接弱、供给缺 → 补视频/达人 + 盯承接", False)],
    ])
build_card(
    slide, 6142990, 4249420, 5686425, 2362200,
    [("28173", True), (" → 达人承接型（可放量）", False)],
    [
        [("结果：", True), ("月度出单 ", False), ("321→846→761", True), ("（8月冲高、9月小回落）", False)],
        [("达人端：", True), ("寄样 ", False), ("67→86→95", True), (" 加码、出单达人 ", False),
         ("16→25→24", True), ("、出单率稳30-34% → 多寄不稀释", False)],
        [("内容端：", True), ("9月视频 ", False), ("172→152", True), (" 降、播放 ", False),
         ("537k→308k", True), (" 降 → 流量少了", False)],
        [("判断：", True), ("达人端一直稳，9月回落是视频少/沉淀期，不是转化崩 → 可继续放量", False)],
    ])

# 4) drop the three pasted images
for e in pic_els:
    for sub in e.iter():
        for attr in ('r:embed', 'r:link', 'r:id'):
            rid = sub.get(qn(attr))
            if rid:
                try:
                    slide.part.drop_rel(rid)
                except Exception:
                    pass
    e.getparent().remove(e)

prs.save(OUT)
print('saved ->', OUT)
