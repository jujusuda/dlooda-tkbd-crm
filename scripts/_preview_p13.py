# -*- coding: utf-8 -*-
"""按成品 slide14 真实几何 + 图表数据生成 SVG 预览（校验版式/重叠/换行）。"""
from pptx import Presentation
from pptx.util import Pt

OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_v9_1(11)_达人素材供给修改v2.pptx"
SVG = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\_preview_p13_v2.svg"

prs = Presentation(OUT)
sl = prs.slides[13]
W = prs.slide_width / 914400
H = prs.slide_height / 914400
S = 78


def esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def cw(ch):
    """字符视觉宽度（CJK=1，半角=0.5）"""
    n = 0
    for c in ch:
        n += 0.5 if ord(c) < 0x2E80 else 1.0
    return n


def wrap(text, width_in, size_pt):
    """按可用宽度（英寸）与字号（pt）粗略折行"""
    per_line = max(1, int(width_in * 72 / size_pt))
    out, cur, curw = [], "", 0.0
    for c in text:
        w = cw(c)
        if curw + w > per_line and cur:
            out.append(cur)
            cur, curw = c, w
        else:
            cur += c
            curw += w
    if cur:
        out.append(cur)
    return out


parts = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {int(W*S)} {int(H*S)}" width="{int(W*S)}" height="{int(H*S)}">']
parts.append(f'<rect width="{int(W*S)}" height="{int(H*S)}" fill="#FFFFFF"/>')


def draw_text(x, y, w, text, size, color, bold=False, align="start", lh=1.22):
    lines = wrap(text, w, size)
    ox = x
    for ln in lines:
        anchor = "middle" if align == "center" else ("end" if align == "right" else "start")
        tx = ox + w / 2 if align == "center" else (ox + w if align == "right" else ox)
        parts.append(f'<text x="{tx*S:.1f}" y="{y*S:.1f}" text-anchor="{anchor}" '
                     f'font-family="Microsoft YaHei" font-size="{size*S/72:.1f}" fill="{color}" '
                     f'{"font-weight=\"bold\" " if bold else ""}dominant-baseline="hanging">{esc(ln)}</text>')
        y += size * lh / 72
    return y


for sh in sl.shapes:
    x, y, w, h = [v / 914400 for v in (sh.left, sh.top, sh.width, sh.height)]
    if sh.has_chart:
        ch = sh.chart
        parts.append(f'<rect x="{x*S:.1f}" y="{y*S:.1f}" width="{w*S:.1f}" height="{h*S:.1f}" '
                     f'fill="#FFFFFF" stroke="#DADADA"/>')
        ty = y + 0.05
        if ch.has_title:
            ty = draw_text(x + 0.02, ty, w - 0.04, ch.chart_title.text_frame.text, 11, "#1F1F1F", True) + 0.04
        plot_top = ty + 0.10
        plot_bot = y + h - 0.30
        plot_left = x + 0.42
        plot_right = x + w - 0.12
        ph = plot_bot - plot_top
        cats = [str(c) for c in ch.plots[0].categories]
        series = list(ch.plots[0].series)
        vmax = ch.value_axis.maximum_scale or max(max(s.values) for s in series)
        vmaj = ch.value_axis.major_unit or (vmax / 4)
        gv = 0.0
        while gv <= vmax + 1e-9:
            gy = plot_bot - (gv / vmax) * ph
            parts.append(f'<line x1="{plot_left*S:.1f}" y1="{gy*S:.1f}" x2="{plot_right*S:.1f}" '
                         f'y2="{gy*S:.1f}" stroke="#EFEFEF"/>')
            parts.append(f'<text x="{(x+0.02)*S:.1f}" y="{(gy-0.07)*S:.1f}" font-family="Microsoft YaHei" '
                         f'font-size="{9*S/72:.1f}" fill="#9A9A9A" dominant-baseline="hanging">{gv:g}</text>')
            gv += vmaj
        gw = (plot_right - plot_left) / len(cats)
        for ci, cat in enumerate(cats):
            cx0 = plot_left + ci * gw
            nb = len(series)
            bw = gw * 0.34 if nb > 1 else gw * 0.18
            start = cx0 + (gw - bw * nb) / 2
            for si, s in enumerate(series):
                v = s.values[ci]
                if v is None:
                    continue
                bh = (v / vmax) * ph
                bx = start + si * bw
                by = plot_bot - bh
                col = "#D9A5B5" if si == 0 else "#B06579"
                if nb == 1:
                    col = "#D9A5B5" if ci == 0 else "#B06579"
                parts.append(f'<rect x="{bx*S:.1f}" y="{by*S:.1f}" width="{bw*S:.1f}" '
                             f'height="{bh*S:.1f}" fill="{col}"/>')
                fmt = f"{v:.2f}" if vmax < 10 else f"{v:.1f}%"
                parts.append(f'<text x="{(bx+bw/2)*S:.1f}" y="{(by-0.19)*S:.1f}" text-anchor="middle" '
                             f'font-family="Microsoft YaHei" font-size="{10*S/72:.1f}" fill="#3D3D3D" '
                             f'dominant-baseline="hanging">{fmt}</text>')
            parts.append(f'<text x="{(cx0+gw/2)*S:.1f}" y="{(plot_bot+0.06)*S:.1f}" text-anchor="middle" '
                         f'font-family="Microsoft YaHei" font-size="{11*S/72:.1f}" fill="#3D3D3D" '
                         f'dominant-baseline="hanging">{esc(cat)}</text>')
        if ch.has_legend:
            lx, ly = x + w - 0.05, ty - 0.28
            for i, it in enumerate([s.name for s in series][::-1]):
                ox = lx - (i + 1) * 0.72
                parts.append(f'<rect x="{ox*S:.1f}" y="{ly*S:.1f}" width="{0.12*S:.1f}" '
                             f'height="{0.12*S:.1f}" fill="{"#B06579" if i==0 else "#D9A5B5"}"/>')
                parts.append(f'<text x="{(ox+0.15)*S:.1f}" y="{(ly-0.02)*S:.1f}" font-family="Microsoft YaHei" '
                             f'font-size="{10*S/72:.1f}" fill="#3D3D3D" dominant-baseline="hanging">{esc(it)}</text>')
    elif sh.has_text_frame and sh.text_frame.text.strip():
        yy = y
        for p in sh.text_frame.paragraphs:
            segs = [(r.text, r.font.size.pt if r.font.size else 12,
                     f"#{r.font.color.rgb}" if (r.font.color and r.font.color.type is not None) else "#3D3D3D",
                     bool(r.font.bold)) for r in p.runs]
            if not segs:
                continue
            line = "".join(s[0] for s in segs)
            size = segs[0][1]
            lh = p.line_spacing or 1.22
            yy = draw_text(x, yy, w, line, size, segs[0][2], segs[0][3],
                           "center" if p.alignment == 2 else "start", lh)
            if p.space_after:
                yy += p.space_after.pt / 72
    else:
        try:
            fill = f"#{sh.fill.fore_color.rgb}"
        except Exception:
            fill = "none"
        parts.append(f'<rect x="{x*S:.1f}" y="{y*S:.1f}" width="{w*S:.1f}" height="{h*S:.1f}" fill="{fill}"/>')

parts.append("</svg>")
open(SVG, "w", encoding="utf-8").write("\n".join(parts))
print("wrote", SVG)
