# -*- coding: utf-8 -*-
"""达人建联页（PART 03·达人端·slide idx 15）v2.1 重做：左右分栏 + 改进漏斗/字号/渠道数量。
左栏（~45%）倒三角梯形漏斗（27→20→17），转化率放梯形层间（单行清晰）；
右栏（~55%）各品明细表 + WhatsApp/付费达人独立小卡（大字写明数量）；
底部通栏落点策略。基于 v9_13，仅改第 16 页。输出 v9_14。
数据均为用户已核对真实值，不改动、不编造。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INK   = RGBColor.from_string('1F1F1F')
BODY  = RGBColor.from_string('3D3D3D')
MUTE  = RGBColor.from_string('9A9A9A')
ACC   = RGBColor.from_string('B06579')   # 玫红强调
GRAY_D= RGBColor.from_string('6B6B6B')
FONT  = '微软雅黑'

WHITE = RGBColor.from_string('FFFFFF')
CARD  = RGBColor.from_string('F7F4F5')
HDR   = RGBColor.from_string('EFE7EA')
ZEB   = RGBColor.from_string('FAF7F8')
LINE  = RGBColor.from_string('E6E0E2')
FUN1  = RGBColor.from_string('C97F95')   # 漏斗上层（浅）
FUN2  = RGBColor.from_string('B06579')   # 漏斗中层
FUN3  = RGBColor.from_string('9A4F66')   # 漏斗下层（深）

SRC = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_13.pptx'
OUT = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_14.pptx'

prs = Presentation(SRC)
s = prs.slides[15]

for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)

def add_text(x, y, w, h, lines, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    first = True
    for para in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(3); p.space_before = Pt(0)
        for (t, sz, b, col) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = b; r.font.name = FONT
            r.font.color.rgb = col
    return tb

def style_cell(cell, text, size, bold, color, fill):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.04); cell.margin_right = Inches(0.04)
    cell.margin_top = Inches(0.004); cell.margin_bottom = Inches(0.004)
    tf = cell.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
    r.font.color.rgb = color
    cell.fill.solid(); cell.fill.fore_color.rgb = fill

# ---------- 1. 顶部标题区 ----------
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.34), Inches(0.1), Inches(0.74)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = ACC; s.shapes[-1].line.fill.background()
add_text(0.84, 0.26, 11.9, 0.66, [[('达人建联：建联 27 → 已发视频 20 → 已出单 17', 24, True, INK)]], anchor=MSO_ANCHOR.MIDDLE)
add_text(9.3, 0.40, 3.43, 0.5, [[('PART 03 · 达人端', 12, False, MUTE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
add_text(0.84, 1.0, 11.89, 0.36, [[('抖音号 / 官号 / 邮件私信建联；四款爆品共成功建联 27 位，其中 17 位已出单', 15, False, GRAY_D)]], anchor=MSO_ANCHOR.MIDDLE)
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.5), Inches(12.13), Inches(0.012)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = LINE; s.shapes[-1].line.fill.background()

# ---------- 2. 左栏：倒三角梯形漏斗（无描边、连续锥形） ----------
CX_L = 3.425
SCALE = 0.159259   # 27→4.3"
def half(v): return (v * SCALE) / 2.0

layers = [
    (27, 20, 1.66, 2.61, FUN1, '成功建联'),
    (20, 17, 2.89, 3.84, FUN2, '已发视频'),
    (17, 11, 4.12, 5.07, FUN3, '已出单'),
]
for (vt, vb, yt, yb, fill, label) in layers:
    ht = half(vt); hb = half(vb)
    builder = s.shapes.build_freeform(Inches(CX_L - ht), Inches(yt))
    builder.add_line_segments([
        (Inches(CX_L + ht), Inches(yt)),
        (Inches(CX_L + hb), Inches(yb)),
        (Inches(CX_L - hb), Inches(yb)),
    ], close=True)
    shp = builder.convert_to_shape()
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background()   # 无描边，漏斗更连续干净
    mid_w = (ht + hb)
    ym = (yt + yb) / 2.0
    add_text(CX_L - mid_w, ym - 0.42, mid_w * 2, 0.84, [
        [(str(vt), 32, True, WHITE)],
        [(label, 13, True, WHITE)],
    ], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

# 层间转化率（单行清晰，居中于层间空隙）
add_text(0.6, 2.61, 5.65, 0.28, [[('建联→发视频  74.1%（流失 26% · 最大环节）', 12, True, ACC)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
add_text(0.6, 3.84, 5.65, 0.28, [[('发视频→出单  85.0%（转化健康）', 12, True, ACC)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

# 整体小字 + 分析点（左栏底部）
add_text(0.6, 5.16, 5.65, 0.30, [[('整体建联→出单  63.0%', 13, True, ACC)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
add_text(0.7, 5.50, 5.45, 0.62, [[
    ('分析：', 11.5, True, ACC),
    ('建联→发视频流失最大（约 26%），下一步重点催发视频、筛选高意愿达人。', 11.5, False, BODY),
]], anchor=MSO_ANCHOR.MIDDLE)

# ---------- 3. 右栏：各品明细表 ----------
add_text(6.45, 1.66, 6.25, 0.26, [[('各品建联明细', 13, True, ACC)]], anchor=MSO_ANCHOR.MIDDLE)
rows = [
    ('2178',  '9',  '7',  '7'),
    ('2190',  '13', '9',  '6'),
    ('28173', '3',  '2',  '2'),
    ('2189',  '2',  '2',  '2'),
    ('合计',  '27', '20', '17'),
]
gf = s.shapes.add_table(6, 4, Inches(6.45), Inches(1.94), Inches(6.25), Inches(1.8))
tb = gf.table
for i, cw_i in enumerate([1.55, 1.55, 1.95, 1.20]):
    tb.columns[i].width = Inches(cw_i)
hdr = ['品', '建联', '已发视频', '已出单']
for ci, h in enumerate(hdr):
    style_cell(tb.cell(0, ci), h, 12, True, INK, HDR)
for ri, row in enumerate(rows, start=1):
    tb.rows[ri].height = Inches(0.30) if ri <= 4 else Inches(0.32)
    is_total = (ri == 5)
    for ci, val in enumerate(row):
        fill = WHITE if ri % 2 == 1 else ZEB
        if is_total:
            fill = HDR
        bold = is_total or (ci >= 1)
        col = ACC if (ci >= 1 and not is_total) else (INK if is_total else BODY)
        style_cell(tb.cell(ri, ci), val, 12, bold, col, fill)

# ---------- 4. 右栏：渠道 / 深绑（独立小卡，大字写明数量，不进主漏斗） ----------
add_text(6.45, 3.90, 6.25, 0.26, [[('渠道 / 深绑（独立，不进主漏斗）', 13, True, ACC)]], anchor=MSO_ANCHOR.MIDDLE)

def channel_card(cy, num, title, sub):
    s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.45), Inches(cy), Inches(6.25), Inches(0.85))
    sh = s.shapes[-1]; sh.fill.solid(); sh.fill.fore_color.rgb = CARD
    sh.line.color.rgb = LINE; sh.line.width = Pt(0.75)
    s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.65), Inches(cy + 0.14), Inches(0.58), Inches(0.58))
    ov = s.shapes[-1]; ov.fill.solid(); ov.fill.fore_color.rgb = ACC; ov.line.fill.background()
    add_text(6.65, cy + 0.14, 0.58, 0.58, [[(num, 22, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    add_text(7.4, cy + 0.10, 5.15, 0.66, [
        [(title, 15, True, INK)],
        [(sub, 11.5, False, GRAY_D)],
    ], anchor=MSO_ANCHOR.MIDDLE)

channel_card(4.20, '10', 'WhatsApp 深绑 10 位', '绑定更深，出单更稳定')
channel_card(5.15, '2',  '付费达人 2 位',       '已寄样，待发视频（推进中）')

# ---------- 5. 底部通栏：落点策略 ----------
s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.30), Inches(12.13), Inches(0.62))
sh = s.shapes[-1]; sh.fill.solid(); sh.fill.fore_color.rgb = CARD
sh.line.color.rgb = LINE; sh.line.width = Pt(0.75)
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(6.44), Inches(0.08), Inches(0.34)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = ACC; s.shapes[-1].line.fill.background()
add_text(0.84, 6.30, 11.7, 0.62, [[
    ('落点策略　', 13.5, True, ACC),
    ('① 扩大建联规模　② WhatsApp 加深绑定　③ 推进付费达人合作落地', 13.5, True, ACC),
]], anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print('SAVED:', OUT, 'slides:', len(prs.slides._sldIdLst))
