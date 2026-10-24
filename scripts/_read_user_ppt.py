# -*- coding: utf-8 -*-
import sys
from pptx import Presentation
from pptx.util import Emu

path = r"E:/xwechat_files/wxid_od6e7r1klxs422_54c2/temp/RWTemp/2026-09/8be9710a0a0dd58ff4cb419316577cc7/余慧玲8月月度汇报.pptx"
prs = Presentation(path)

print("=== 幻灯片尺寸 ===")
print("width EMU:", prs.slide_width, "height EMU:", prs.slide_height)
print("width inch:", prs.slide_width / 914400, "height inch:", prs.slide_height / 914400)
print("总页数:", len(prs.slides))
print()

for idx, slide in enumerate(prs.slides, start=1):
    print("=" * 60)
    print(f"### Slide {idx} ###")
    for shp in slide.shapes:
        kind = shp.shape_type
        name = shp.name
        # 文本
        text = ""
        if shp.has_text_frame:
            text = shp.text_frame.text.strip()
        # 图表
        chart_info = ""
        if shp.has_chart:
            ch = shp.chart
            chart_info = f"[CHART type={ch.chart_type}]"
            try:
                cats = list(ch.plots[0].categories)
                chart_info += f" cats={cats}"
            except Exception:
                pass
        # 表格
        table_info = ""
        if shp.has_table:
            table_info = "[TABLE]"
        if text:
            # 限制每段长度
            lines = text.split("\n")
            shown = "\n  | ".join(l.strip() for l in lines if l.strip())
            if len(shown) > 1500:
                shown = shown[:1500] + "...(截断)"
            print(f"  [{kind}] {name}: {chart_info}{table_info}\n  | {shown}")
        elif chart_info or table_info:
            print(f"  [{kind}] {name}: {chart_info}{table_info}")
    print()
