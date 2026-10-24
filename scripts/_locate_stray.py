# -*- coding: utf-8 -*-
from pptx import Presentation
OUT = r'C:\Users\LENOVO\Desktop\8月余慧玲8月月度汇报_商务淡粉版.pptx'
p = Presentation(OUT)
WATCH = {'#4A3A41', '#D6336C'}
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        if run.font.color and run.font.color.type is not None:
                            col = '#'+str(run.font.color.rgb)
                        else: continue
                    except Exception: continue
                    if col in WATCH:
                        print('P%d TEXT %s -> %r' % (i, col, run.text.strip()[:25]))
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    for para in c.text_frame.paragraphs:
                        for run in para.runs:
                            try:
                                if run.font.color and run.font.color.type is not None:
                                    col = '#'+str(run.font.color.rgb)
                                else: continue
                            except Exception: continue
                            if col in WATCH:
                                print('P%d CELL %s -> %r' % (i, col, run.text.strip()[:25]))
        if sh.has_chart:
            ch = sh.chart
            for ser in ch.series:
                try:
                    if ser.data_labels and ser.data_labels.font and ser.data_labels.font.color and ser.data_labels.font.color.type is not None:
                        col = '#'+str(ser.data_labels.font.color.rgb)
                        if col in WATCH: print('P%d CHART-DL %s' % (i, col))
                except Exception: pass
            for axname, ax in (('cat',ch.category_axis),('val',ch.value_axis)):
                try:
                    if ax.tick_labels and ax.tick_labels.font and ax.tick_labels.font.color and ax.tick_labels.font.color.type is not None:
                        col='#'+str(ax.tick_labels.font.color.rgb)
                        if col in WATCH: print('P%d AXIS-%s %s' % (i, axname, col))
                except Exception: pass
