# -*- coding: utf-8 -*-
"""第 14 页（达人素材供给）v2：上=两张人均素材图，中=发布次数分布对比，下=因果/结论/动作。
全部 PPT 原生可编辑，禁止图片。
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(11).pptx"
OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_v9_1(11)_达人素材供给修改v2.pptx"

FONT = "微软雅黑"
ROSE = RGBColor(0xB0, 0x65, 0x79)   # #B06579
PINK = RGBColor(0xD9, 0xA5, 0xB5)   # #D9A5B5
INK = RGBColor(0x1F, 0x1F, 0x1F)
BODY = RGBColor(0x3D, 0x3D, 0x3D)
MUTE = RGBColor(0x9A, 0x9A, 0x9A)
CARD_BG = RGBColor(0xF7, 0xF4, 0xF5)


def force_ea(font):
    try:
        rPr = font._rPr
    except Exception:
        return
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        return
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = copy.deepcopy(latin)
        ea.tag = qn("a:ea")
        latin.addnext(ea)
    else:
        ea.set("typeface", latin.get("typeface"))


def style_font(font, size=None, bold=None, color=None, name=FONT):
    if name:
        font.name = name
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color
    force_ea(font)


def add_para(tf, parts, align=PP_ALIGN.LEFT, space_after=0, first=False, line_spacing=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if space_after:
        p.space_after = Pt(space_after)
    if line_spacing:
        p.line_spacing = line_spacing
    for t, size, bold, color in parts:
        r = p.add_run()
        r.text = t
        style_font(r.font, size, bold, color)
    return p


def new_tb(slide, x, y, w, h, wrap=True, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def ensure_numfmt(dlbls, code):
    nf = dlbls.find(qn("c:numFmt"))
    if nf is None:
        nf = parse_xml(f'<c:numFmt {nsdecls("c")} formatCode="{code}" sourceLinked="0"/>')
        dlbls.insert(0, nf)
    else:
        nf.set("formatCode", code)
        nf.set("sourceLinked", "0")


def setup_labels(plot, fmt="0.00", size=10):
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.show_value = True
    try:
        dl.number_format = fmt
        dl.number_format_is_linked = False
    except Exception as e:
        print("  [warn] number_format api:", e)
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    style_font(dl.font, size, False, BODY)
    for d in plot._element.iter(qn("c:dLbls")):
        ensure_numfmt(d, fmt)


def style_axis(ax, size=10, color=MUTE, fmt=None, maximum=None, major=None, grid=False):
    ax.has_major_gridlines = grid
    tl = ax.tick_labels
    style_font(tl.font, size, False, color)
    if fmt is not None:
        try:
            tl.number_format = fmt
            tl.number_format_is_linked = False
        except Exception as e:
            print("  [warn] axis numfmt:", e)
    if maximum is not None:
        ax.maximum_scale = maximum
        ax.minimum_scale = 0
    if major is not None:
        ax.major_unit = major


def set_chart_title(chart, text, size=11):
    chart.has_title = True
    tf = chart.chart_title.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    style_font(p.runs[0].font, size, True, INK)


def set_legend(chart, size=10, in_layout=True):
    chart.has_legend = True
    lg = chart.legend
    lg.position = XL_LEGEND_POSITION.TOP
    lg.include_in_layout = in_layout
    style_font(lg.font, size, False, BODY)


def set_point_color(series, idx, rgb):
    pt = series.points[idx]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb
    pt.format.line.fill.background()


def add_col_chart(slide, x, y, w, h, cats, series_defs, gap=60, vmax=None, vmajor=None,
                  label_fmt="0.00", axis_fmt="0.0", legend=False, title=None, grid=True,
                  legend_in_layout=True):
    """series_defs: [(name, values, rgb), ...]"""
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals, _ in series_defs:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                                Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    style_font(ch.font, 10, False, BODY)
    if title:
        set_chart_title(ch, title)
    else:
        ch.has_title = False
    pl = ch.plots[0]
    pl.gap_width = gap
    if len(series_defs) > 1:
        pl.overlap = -12
    for ser, (_, _, rgb) in zip(pl.series, series_defs):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb
        ser.format.line.fill.background()
    setup_labels(pl, label_fmt, 10)
    style_axis(ch.category_axis, 11, BODY)
    style_axis(ch.value_axis, 10, MUTE, fmt=axis_fmt, maximum=vmax, major=vmajor, grid=grid)
    if legend:
        set_legend(ch, in_layout=legend_in_layout)
    else:
        ch.has_legend = False
    return ch


# =====================================================================
prs = Presentation(SRC)
slide = prs.slides[13]  # 第 14 页

keep = {0, 1, 2, 3}
for idx in sorted([i for i in range(len(slide.shapes)) if i not in keep], reverse=True):
    el = slide.shapes[idx]._element
    el.getparent().remove(el)
print("shapes left:", len(slide.shapes))

# ---------- 上排：两张人均素材图 ----------
TOP_Y, TOP_H = 1.78, 1.80

ch1 = add_col_chart(slide, 0.6, TOP_Y, 5.4, TOP_H,
                    ["8月", "9月"],
                    [("人均素材", (1.48, 1.90), ROSE)],
                    gap=130, vmax=2.5, vmajor=0.5,
                    title="整月人均素材（条/人）", legend=False)
set_point_color(ch1.plots[0].series[0], 0, PINK)
set_point_color(ch1.plots[0].series[0], 1, ROSE)

tb, tf = new_tb(slide, 2.95, 2.03, 0.70, 0.26)
add_para(tf, [("+28%", 15, True, ROSE)], align=PP_ALIGN.CENTER, first=True)

tb, tf = new_tb(slide, 0.6, 3.62, 5.4, 0.26)
add_para(tf, [("9月人均素材较8月回升（1.48→1.90，", 11, False, BODY),
              ("+28%", 11, True, ROSE),
              ("）", 11, False, BODY)], first=True)

add_col_chart(slide, 6.35, TOP_Y, 5.98, TOP_H,
              ["28173", "2190", "2178", "2189"],
              [("8月", (1.52, 1.42, 1.21, 1.47), PINK),
               ("9月", (1.79, 1.69, 1.50, 1.48), ROSE)],
              gap=55, vmax=2.4, vmajor=0.5,
              title="四款爆品人均素材（条/人）", legend=True)

# ---------- 中排：发布次数分布对比 ----------
add_col_chart(slide, 0.6, 3.92, 12.13, 1.50,
              ["发 1 条", "发 2 条", "发 3 条及以上"],
              [("8月", (83.8, 9.3, 6.9), PINK),
               ("9月", (69.9, 17.6, 12.5), ROSE)],
              gap=400, vmax=100, vmajor=25,
              label_fmt='0.0"%"', axis_fmt='0"%"',
              title="达人发布次数分布（按 SKU·达人组合）", legend=True,
              legend_in_layout=False)

tb, tf = new_tb(slide, 0.6, 5.47, 12.13, 0.26)
add_para(tf, [("多次发布达人占比翻倍（", 11, False, BODY),
              ("16.2%→30.1%", 11, True, ROSE),
              ("），人均素材上升的动因", 11, False, BODY),
              ("　　口径：达人 × SKU 组合、去重视频 ID（一个达人合作多款 SKU 各计一次）", 9, False, MUTE)],
         first=True)

# ---------- 下排：因果 / 结论 / 动作 ----------
CX, CY, CW, CH, GAP = 0.6, 5.76, 3.87, 1.40, 0.26
CARDS = [
    ("因果", [("9月素材总量下降（1070→752）、达人更精炼（723→395），人均反升；", 11.5, False, BODY),
              ("多次发布达人占比翻倍（16.2%→30.1%）", 11.5, True, BODY),
              ("，更多人从只发1条转向多发，供给效率整体提升", 11.5, False, BODY)]),
    ("结论", [("9月人均素材 1.90，较8月回升（+28%），", 11.5, False, BODY),
              ("动因是多次发布达人占比提升", 11.5, True, BODY),
              ("，达人端素材供给更高效，核心款人均普遍 1.5 以上", 11.5, False, BODY)]),
    ("动作", None),
]
ACTION = [
    "① 延续复投核心达人（产出更高效）",
    "② 重点补 2189 素材（唯一持平款）",
    "③ 稳定多次发布达人占比，让人均往 2+ 走",
]

for i, (title, parts) in enumerate(CARDS):
    x = CX + i * (CW + GAP)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(CY), Inches(CW), Inches(CH))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.fill.background()
    card.shadow.inherit = False
    card.text_frame.text = ""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(CY + 0.12), Inches(0.055), Inches(CH - 0.24))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ROSE
    bar.line.fill.background()
    bar.shadow.inherit = False
    bar.text_frame.text = ""
    tb, tf = new_tb(slide, x + 0.22, CY + 0.13, CW - 0.36, 0.28)
    add_para(tf, [(title, 13.5, True, ROSE)], first=True)
    tb, tf = new_tb(slide, x + 0.22, CY + 0.46, CW - 0.36, CH - 0.56)
    if parts is not None:
        add_para(tf, parts, first=True, line_spacing=1.22)
    else:
        for j, a in enumerate(ACTION):
            if a.startswith("②"):
                add_para(tf, [(a[:2], 11.5, True, ROSE), (a[2:], 11.5, False, BODY)],
                         first=(j == 0), space_after=2, line_spacing=1.12)
            else:
                add_para(tf, [(a, 11.5, False, BODY)],
                         first=(j == 0), space_after=2, line_spacing=1.12)

prs.save(OUT)
print("saved:", OUT)
