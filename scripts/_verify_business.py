# -*- coding: utf-8 -*-
from pptx import Presentation
import collections

OUT = r'C:\Users\LENOVO\Desktop\8月余慧玲8月月度汇报_商务淡粉版.pptx'
p = Presentation(OUT)

fills = collections.Counter(); texts = collections.Counter(); lines = collections.Counter()
white_on_light = []  # text white/#FCEAF1 but shape fill white/none and slide white -> invisible risk
for i, s in enumerate(p.slides, 1):
    slide_bg = None
    try:
        if s.background.fill.type == 1: slide_bg = '#'+str(s.background.fill.fore_color.rgb)
    except Exception: pass
    for sh in s.shapes:
        sfc = None
        try:
            if sh.fill.type == 1: sfc = '#'+str(sh.fill.fore_color.rgb)
        except Exception: pass
        if sfc: fills[sfc] += 1
        try:
            if sh.line and sh.line.color and sh.line.color.type == 1:
                lines['#'+str(sh.line.color.rgb)] += 1
        except Exception: pass
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    try:
                        if c.fill.type == 1: fills['#'+str(c.fill.fore_color.rgb)] += 1
                    except Exception: pass
        if sh.has_chart:
            for ser in sh.chart.series:
                try:
                    if ser.format.fill.type == 1: fills['#'+str(ser.format.fill.fore_color.rgb)] += 1
                except Exception: pass
        # text runs + invisibility check
        def check_tf(tf):
            for para in tf.paragraphs:
                for run in para.runs:
                    try:
                        if run.font.color and run.font.color.type is not None:
                            col = '#'+str(run.font.color.rgb)
                        else:
                            continue
                    except Exception:
                        continue
                    texts[col] += 1
                    # invisible risk: light text on white/light shape, white slide
                    if col in ('#FFFFFF', '#FCEAF1'):
                        light_bg = (sfc in (None,'#FFFFFF') or (sfc and sfc.upper() in ('#FFFFFF','#F5EEEF')))
                        if light_bg and (slide_bg in (None,'#FFFFFF')):
                            white_on_light.append((i, col, sfc, slide_bg, run.text.strip()[:20]))
        if sh.has_text_frame:
            check_tf(sh.text_frame)
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    check_tf(c.text_frame)

print('=== REMAINING FILL COLORS ===')
for c,n in fills.most_common(): print('  %s : %d' % (c,n))
print('\n=== REMAINING TEXT COLORS ===')
for c,n in texts.most_common(): print('  %s : %d' % (c,n))
print('\n=== REMAINING LINE COLORS ===')
for c,n in lines.most_common(): print('  %s : %d' % (c,n))
print('\n=== WHITE/FCEAF1 TEXT ON LIGHT BG (invisible risk) ===')
if not white_on_light: print('  NONE')
else:
    for x in white_on_light: print('  ', x)
