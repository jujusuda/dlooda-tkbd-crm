# -*- coding: utf-8 -*-
from pptx import Presentation
import re
F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_9.pptx'
prs = Presentation(F)
for idx, s in enumerate(prs.slides):
    texts = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip().split('\n')[0][:60])
    joined = ' | '.join(texts)
    if '2L06' in joined or '2722' in joined:
        print(f"=== SLIDE {idx+1} ===")
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                print("  TXT:", repr(sh.text_frame.text[:120]))
            if sh.has_table:
                tb = sh.table
                for r in range(len(tb.rows)):
                    print("   R%d:"%r, " | ".join(repr(tb.cell(r,c).text) for c in range(len(tb.columns))))
        print()
