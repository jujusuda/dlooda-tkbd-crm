# -*- coding: utf-8 -*-
"""P10（2189 + 28173｜同周期两款）整改：基于 v9_5（已含 P5+P8+P9 修复），仅重建第 10 页（slide idx 9）。
- 标题/副标题：差异分析视角（一个靠内容、一个靠达人），副标题不出现具体数字
- 三张图表：图一 达人出单率趋势(折线)、图二 出单达人数量趋势(折线)、图三 9月关键指标对比(横向条形)
- 分析文字：2189(内容转化型)/28173(达人承接型)，讲变化/原因/卡点/策略（数据在图三）
- 落点策略：仅两条（28173放量 / 2189补供给）
配色/字号沿用 v9 规范。折线图标签 ABOVE/BELOW（禁用 outEnd/inEnd）；横向条形用 BAR_CLUSTERED+OUTSIDE_END（合法）。
产物 Desktop/余慧玲9月月度汇报_v9_6.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE

INK   = RGBColor.from_string('1F1F1F')
BODY  = RGBColor.from_string('3D3D3D')
MUTE  = RGBColor.from_string('9A9A9A')
ACC   = RGBColor.from_string('B06579')   # 2189 线/强调
ACC_L = RGBColor.from_string('C97F95')
CARD  = RGBColor.from_string('F9F5F6')
LINE  = RGBColor.from_string('EAE2E5')
GRAY_D= RGBColor.from_string('5A5A5A')   # 28173 线
WHITE = RGBColor.from_string('FFFFFF')
FONT  = '微软雅黑'

SRC = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_5.pptx'
OUT = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_6.pptx'

prs = Presentation(SRC)
s = prs.slides[9]  # P10

for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)

def rect(x, y, w, h, fill, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp

def add_text(x, y, w, h, paras, anchor=None, align=None, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    if anchor is not None: tf.vertical_anchor = anchor
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if align is not None: p.alignment = align
        for (t, sz, bold, color) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = FONT
    return tb

# ---------- 标题区 ----------
rect(0.6, 0.34, 0.1, 0.74, ACC)
add_text(0.84, 0.30, 11.7, 0.66,
         [[('2189 + 28173｜同周期两款：一个靠内容、一个靠达人', 24, True, INK)]], anchor=MSO_ANCHOR.MIDDLE)
add_text(9.3, 0.42, 3.43, 0.56,
         [[('PART 02 · 各款 SKU 分析', 12, False, MUTE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
add_text(0.84, 1.02, 11.8, 0.34,
         [[('外形相似、7 月同推；2189 内容一转就转化，28173 靠达人承接放量', 15, False, BODY)]])
rect(0.6, 1.44, 12.13, Pt(1)/914400, LINE)

# ---------- 折线图通用 ----------
def add_line_chart(x, y, w, h, cats, series, colors, vmax, vfmt='0.00"%"'):
    cd = CategoryChartData(); cd.categories = cats
    for (nm, vals) in series: cd.add_series(nm, vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(10); ch.legend.font.name = FONT
    pl = ch.plots[0]
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.show_value = True; dl.number_format = vfmt; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.ABOVE
    dl.font.size = Pt(9); dl.font.bold = False; dl.font.name = FONT; dl.font.color.rgb = INK
    for i, c in enumerate(colors):
        se = pl.series[i]
        se.format.line.color.rgb = c
        se.format.line.width = Pt(2.25)
        se.marker.style = XL_MARKER_STYLE.CIRCLE
        se.marker.size = 6
        se.marker.format.fill.solid(); se.marker.format.fill.fore_color.rgb = c
        se.marker.format.line.color.rgb = WHITE; se.marker.format.line.width = Pt(1)
        sd = se.data_labels
        sd.show_value = True; sd.number_format = vfmt; sd.number_format_is_linked = False
        # 2189 上方(t)、28173 下方(b)：竖直错开避免叠字（9月两款出单率仅差 0.47pp）
        sd.position = XL_LABEL_POSITION.BELOW if i == 1 else XL_LABEL_POSITION.ABOVE
        sd.font.size = Pt(9); sd.font.bold = False; sd.font.name = FONT; sd.font.color.rgb = INK
    va = ch.value_axis
    va.minimum_scale = 0; va.maximum_scale = vmax
    va.major_unit = 15 if vmax >= 60 else 10
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.major_gridlines.format.line.width = Pt(0.5)
    va.tick_labels.font.size = Pt(9); va.tick_labels.font.name = FONT
    va.tick_labels.font.color.rgb = MUTE; va.tick_labels.number_format = vfmt
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(11); ca.tick_labels.font.name = FONT; ca.tick_labels.font.color.rgb = GRAY_D
    ca.format.line.color.rgb = LINE
    return ch

# ===== 左列：图一 / 图二 =====
add_text(0.6, 1.54, 6.0, 0.26, [[('达人出单率趋势（7–9 月）', 14, True, INK)]])
add_line_chart(0.6, 1.80, 6.0, 2.0,
               ['7月', '8月', '9月'],
               [('2189', [56.00, 22.86, 31.11]),
                ('28173', [30.19, 33.78, 31.58])],
               [ACC, GRAY_D], 60)
add_text(0.6, 3.88, 6.0, 0.26, [[('出单达人数量趋势（7–9 月）', 14, True, INK)]])
add_line_chart(0.6, 4.14, 6.0, 2.36,
               ['7月', '8月', '9月'],
               [('2189', [14, 8, 14]),
                ('28173', [16, 25, 24])],
               [ACC, GRAY_D], 30, vfmt='0')

# ===== 右列：图三（横向条形） =====
add_text(6.8, 1.54, 5.93, 0.26, [[('两款 9 月关键指标对比', 14, True, INK)]])
metrics = ['寄样','履约','履约率','出单达人','达人出单率','视频数量','出单视频',
           '视频出单率','月度出单','自然单占比','点击率','转化率']
v2189  = [52, 46, 88.46, 14, 31.11, 79, 22, 27.85, 999, 71.35, 5.17, 1.13]
v28173 = [95, 77, 81.05, 24, 31.58, 152, 35, 23.03, 761, 76.85, 4.06, 0.90]
cd = CategoryChartData(); cd.categories = metrics
cd.add_series('2189', v2189); cd.add_series('28173', v28173)
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(6.8), Inches(1.80), Inches(5.93), Inches(2.15), cd)
ch3 = gf.chart
ch3.has_title = False
ch3.has_legend = True; ch3.legend.position = XL_LEGEND_POSITION.TOP
ch3.legend.include_in_layout = False
ch3.legend.font.size = Pt(10); ch3.legend.font.name = FONT
pl3 = ch3.plots[0]
pl3.has_data_labels = True
pl3.data_labels.show_value = True; pl3.data_labels.number_format = '0.##'; pl3.data_labels.number_format_is_linked = False
pl3.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl3.data_labels.font.size = Pt(8); pl3.data_labels.font.bold = False; pl3.data_labels.font.name = FONT; pl3.data_labels.font.color.rgb = INK
pl3.series[0].format.fill.solid(); pl3.series[0].format.fill.fore_color.rgb = ACC
pl3.series[1].format.fill.solid(); pl3.series[1].format.fill.fore_color.rgb = GRAY_D
ca3 = ch3.category_axis
ca3.tick_labels.font.size = Pt(9); ca3.tick_labels.font.name = FONT; ca3.tick_labels.font.color.rgb = BODY
va3 = ch3.value_axis
va3.has_major_gridlines = True; va3.major_gridlines.format.line.color.rgb = LINE; va3.major_gridlines.format.line.width = Pt(0.5)
va3.tick_labels.font.size = Pt(8); va3.tick_labels.font.name = FONT; va3.tick_labels.font.color.rgb = MUTE
va3.has_title = False

# ===== 右列下：分析文字（两款） =====
def analysis_card(x, y, w, h, title, color, bullets):
    rect(x, y, w, h, CARD)
    rect(x, y, 0.07, h, color)
    add_text(x+0.22, y+0.06, w-0.3, 0.28, [[(title, 13.5, True, INK)]])
    paras = [[('• ', 11, True, color), (b, 11, False, BODY)] for b in bullets]
    add_text(x+0.22, y+0.36, w-0.35, h-0.42, paras)

analysis_card(6.8, 4.03, 5.93, 1.32, '2189｜内容转化型（缺供给）', ACC,
    ['变化：月度 176→723→999 三连升；达人出单率 56→22.9→31.1 波动大',
     '原因：内容一转就转化，点击/转化/视频出单率均高、单条效率强',
     '卡点：视频仅 79 条、寄样偏小、出单达人少（14→8→14），出单率受单笔影响大',
     '策略：① 补视频 / 达人供给　② 释放最高出单率'])
analysis_card(6.8, 5.43, 5.93, 1.32, '28173｜达人承接型（可放量）', GRAY_D,
    ['变化：寄样 67→86→95 加码，出单达人稳 16→25→24；自然单占比 60.7→81.6→76.85% 高位',
     '原因：多寄不稀释效率、达人愿带；9 月回落因新增达人沉淀 + 视频量降（172→152），非转化崩',
     '卡点：放量型，靠达人带量、点击率 4.06% 偏低不靠内容点击',
     '策略：① 放大寄样　② 复投核心达人　③ 加大素材拉回订单'])

# ---------- 落点策略（仅两条） ----------
SYB, SHB = 6.84, 0.56
rect(0.6, SYB, 12.13, SHB, CARD)
rect(0.6, SYB, 0.09, SHB, ACC)
add_text(0.95, SYB, 11.65, SHB,
         [[('落点策略　', 13, True, ACC),
           ('① ', 13, True, INK), ('28173 放量：继续放大寄样、复投核心达人、拉回订单', 13, False, BODY),
           ('　② ', 13, True, INK), ('2189 补供给：补视频 / 达人，释放转化力', 13, False, BODY)]],
         anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print('SAVED:', OUT, 'slides:', len(prs.slides._sldIdLst))
