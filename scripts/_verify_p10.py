# -*- coding: utf-8 -*-
import zipfile, re, posixpath
from pptx import Presentation

F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_6.pptx'
prs = Presentation(F)
s = prs.slides[9]

# --- chart data ---
print("=== P10 charts ===")
for sh in s.shapes:
    if sh.has_chart:
        ch = sh.chart
        print("chart type:", ch.chart_type)
        for se in ch.plots[0].series:
            print("   ", se.name, "->", [round(v,3) if isinstance(v,float) else v for v in se.values])

# --- dLblPos legality on ALL line charts in deck ---
z = zipfile.ZipFile(F)
print("\n=== line-chart dLblPos (legal=t/b/ctr/l/r) ===")
bad = []
for n in sorted(z.namelist()):
    if re.match(r'ppt/(slides/)?charts/chart\d+\.xml$', n):
        x = z.read(n).decode('utf-8','ignore')
        if '<c:lineChart>' in x:
            pos = re.findall(r'<c:dLblPos val="(\w+)"/>', x)
            ill = [p for p in pos if p not in ('t','b','ctr','l','r')]
            print("  %s -> %s %s" % (n, pos, "  ILLEGAL!" if ill else "OK"))
            if ill: bad.append((n, ill))

# --- package integrity ---
names = set(z.namelist())
ct = z.read('[Content_Types].xml').decode('utf-8','ignore')
overrides = set(re.findall(r'PartName="([^"]+)"', ct))
defaults = set(re.findall(r'Extension="([^"]+)"', ct))
pkg = []
for o in overrides:
    if o.lstrip('/') not in names: pkg.append('CT override missing: '+o)
for nm in names:
    if nm == '[Content_Types].xml': continue
    ext = nm.rsplit('.',1)[-1].lower()
    if ('/'+nm) not in overrides and ext not in defaults: pkg.append('part not declared: '+nm)
    if nm.endswith('.rels'):
        base = posixpath.dirname(posixpath.dirname(nm))
        for rid, tgt in re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read(nm).decode('utf-8','ignore')):
            if tgt.startswith('http'): continue
            full = tgt.lstrip('/') if tgt.startswith('/') else posixpath.normpath(posixpath.join(base, tgt))
            if full not in names: pkg.append('%s: %s->%s MISSING' % (nm, rid, full))
print("\nPACKAGE PROBLEMS:", pkg if pkg else "none")
print("ILLEGAL LABELS:", bad if bad else "none")

# --- forbidden words + text dump ---
print("\n=== P10 text scan ===")
full = ""
for sh in s.shapes:
    if sh.has_text_frame: full += sh.text_frame.text + "\n"
print("contains '播放量':", '播放量' in full)
print("contains 'vs 强弱' style '内容转化强 vs':", '内容转化强 vs' in full)
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print("  ", repr(sh.text_frame.text[:60]))
