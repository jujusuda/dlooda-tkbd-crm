# -*- coding: utf-8 -*-
from pptx import Presentation
import re
F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_1.pptx'
prs = Presentation(F)
s = prs.slides[7]
def emu(v): return round(v/914400,2) if v is not None else None
print("=== P8 shapes (count=%d) ===" % len(s.shapes))
for i, sh in enumerate(s.shapes):
    L=emu(sh.left); T=emu(sh.top); W=emu(sh.width); H=emu(sh.height)
    kind = sh.shape_type
    extra=""
    if sh.has_text_frame:
        txt = sh.text_frame.text.replace("\n"," / ")
        extra = "TXT="+repr(txt[:140])
    if sh.has_chart:
        extra = "CHART="+str(sh.chart.chart_type)
    if sh.has_table:
        tb = sh.table
        extra = "TABLE rows=%d cols=%d"%(len(tb.rows), len(tb.columns))
    print(f"[{i}] L={L} T={T} W={W} H={H} {kind} {extra}")
for sh in s.shapes:
    if sh.has_table:
        tb = sh.table
        print("\n=== TABLE CONTENT ===")
        for r in range(len(tb.rows)):
            cells=[]
            for c in range(len(tb.columns)):
                cells.append(repr(tb.cell(r,c).text))
            print("R%d:"%(r), " | ".join(cells))
