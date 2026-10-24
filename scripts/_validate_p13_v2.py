# -*- coding: utf-8 -*-
import zipfile
from pptx import Presentation
from pptx.oxml.ns import qn

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(11).pptx"
OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_v9_1(11)_达人素材供给修改v2.pptx"
IDX = 13

ok = True


def chk(cond, msg):
    global ok
    print(("  PASS  " if cond else "  FAIL  ") + msg)
    if not cond:
        ok = False


z = zipfile.ZipFile(OUT)
print("ZIP testzip:", z.testzip())
a, b = Presentation(SRC), Presentation(OUT)
chk(len(a.slides) == len(b.slides), f"页数一致 {len(b.slides)}")

src_zip, out_zip = zipfile.ZipFile(SRC), zipfile.ZipFile(OUT)
CR, LF = chr(13), chr(10)


def norm(x):
    s = x.decode("utf-8")
    return s.replace(CR + LF, LF).replace(CR, LF).encode("utf-8")


same = diff = 0
for n in src_zip.namelist():
    parts = n.split("/")
    if len(parts) == 3 and parts[0] == "ppt" and parts[1] == "slides" and parts[2].startswith("slide") and n.endswith(".xml"):
        num = int(parts[2][5:-4])
        if num == IDX + 1:
            continue
        ra = src_zip.read(n).split(b"?>")[-1]
        rb = out_zip.read(n).split(b"?>")[-1]
        if norm(ra) == norm(rb):
            same += 1
        else:
            diff += 1
            print("   DIFF", n)
chk(diff == 0, f"其余 {same} 页与源内容完全一致（忽略 XML 声明/换行差异）")

sl = b.slides[IDX]
pics = [s for s in sl.shapes if s.shape_type is not None and str(s.shape_type).startswith("PICTURE")]
chk(len(pics) == 0, f"本页图片数 = {len(pics)}（要求 0）")
charts = [s for s in sl.shapes if s.has_chart]
chk(len(charts) == 3, f"本页原生图表数 = {len(charts)}（要求 3）")
for i, s in enumerate(charts):
    chk(str(s.chart.chart_type).startswith("COLUMN_CLUSTERED"), f"图{i+1} 类型 = {s.chart.chart_type}")

# --- 图1 整月人均
c1 = charts[0].chart
chk([str(x) for x in c1.plots[0].categories] == ["8月", "9月"], "图1 横轴 8月/9月")
chk(list(c1.plots[0].series[0].values) == [1.48, 1.90], f"图1 数值 {list(c1.plots[0].series[0].values)}")
chk(c1.chart_title.text_frame.text.strip() == "整月人均素材（条/人）", "图1 标题")
dp = {}
for d in c1.plots[0]._element.iter(qn("c:dPt")):
    i = d.find(qn("c:idx")).get("val")
    clr = d.find(".//" + qn("a:srgbClr"))
    dp[i] = clr.get("val") if clr is not None else None
chk(dp.get("0") == "D9A5B5" and dp.get("1") == "B06579", f"图1 逐点配色 {dp}")
chk(not c1.has_legend, "图1 无图例")

# --- 图2 四款
c2 = charts[1].chart
chk([str(x) for x in c2.plots[0].categories] == ["28173", "2190", "2178", "2189"], "图2 横轴四款")
v2 = [list(s.values) for s in c2.plots[0].series]
chk(v2[0] == [1.52, 1.42, 1.21, 1.47], f"图2 8月 {v2[0]}")
chk(v2[1] == [1.79, 1.69, 1.50, 1.48], f"图2 9月 {v2[1]}")
chk(c2.chart_title.text_frame.text.strip() == "四款爆品人均素材（条/人）", "图2 标题")
cols2 = [s._element.find(".//" + qn("a:srgbClr")).get("val") for s in c2.plots[0].series]
chk(cols2 == ["D9A5B5", "B06579"], f"图2 配色 {cols2}")
chk(c2.has_legend and str(c2.legend.position).startswith("TOP"), "图2 图例 TOP")

# --- 图3 分布
c3 = charts[2].chart
chk([str(x) for x in c3.plots[0].categories] == ["发 1 条", "发 2 条", "发 3 条及以上"], "图3 横轴三档")
v3 = [list(s.values) for s in c3.plots[0].series]
chk(v3[0] == [83.8, 9.3, 6.9], f"图3 8月 {v3[0]}")
chk(v3[1] == [69.9, 17.6, 12.5], f"图3 9月 {v3[1]}")
chk(c3.chart_title.text_frame.text.strip() == "达人发布次数分布（按 SKU·达人组合）", "图3 标题")
cols3 = [s._element.find(".//" + qn("a:srgbClr")).get("val") for s in c3.plots[0].series]
chk(cols3 == ["D9A5B5", "B06579"], f"图3 配色 {cols3}")
chk(c3.has_legend and str(c3.legend.position).startswith("TOP"), "图3 图例 TOP")

# --- 数据标签
for tag, ch, want in (("图1", c1, "0.00"), ("图2", c2, "0.00"), ("图3", c3, '0.0"%"')):
    el = ch.plots[0]._element
    nf = [d.find(qn("c:numFmt")).get("formatCode") for d in el.iter(qn("c:dLbls")) if d.find(qn("c:numFmt")) is not None]
    sv = [d.find(qn("c:showVal")).get("val") for d in el.iter(qn("c:dLbls")) if d.find(qn("c:showVal")) is not None]
    pos = [d.find(qn("c:dLblPos")).get("val") for d in el.iter(qn("c:dLbls")) if d.find(qn("c:dLblPos")) is not None]
    print(f"   {tag} dLbls numFmt={nf} showVal={sv} pos={pos}")
    chk(all(f == want for f in nf) and nf, f"{tag} 标签格式 {want}")
    chk(all(v == "1" for v in sv) and sv, f"{tag} 显示数值")
    chk(all(p == "outEnd" for p in pos), f"{tag} 标签 outEnd（柱状合法）")

# --- 文字
alltext = "\n".join(s.text_frame.text for s in sl.shapes if s.has_text_frame)
for kw in ["+28%", "1.48", "1.90", "16.2%→30.1%",
           "达人素材供给", "因果", "结论", "动作",
           "1070→752", "723→395", "多次发布达人占比",
           "延续复投核心达人", "重点补 2189 素材", "稳定多次发布达人占比",
           "去重视频 ID"]:
    chk(kw in alltext, f"含文案 {kw!r}")
for bad in ["靠人头铺量", "近七成"]:
    chk(bad not in alltext, f"无旧措辞 {bad!r}")

# --- 边界
W, H = b.slide_width / 914400, b.slide_height / 914400
for s in sl.shapes:
    if s.left is None:
        continue
    x, y, w, h = [v / 914400 for v in (s.left, s.top, s.width, s.height)]
    if x < -0.01 or y < -0.01 or x + w > W + 0.01 or y + h > H + 0.01:
        ok = False
        print(f"  FAIL  越界 {s.shape_type} {x:.2f},{y:.2f},{w:.2f},{h:.2f}")
print("  边界检查完成（%.2f x %.2f）" % (W, H))

print("\nRESULT:", "ALL PASS" if ok else "HAS FAILURES")
