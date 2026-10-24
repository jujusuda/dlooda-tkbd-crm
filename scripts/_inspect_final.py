from pptx import Presentation
from pptx.util import Emu
import collections

SRC = r'E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-09\8be9710a0a0dd58ff4cb419316577cc7\8月余慧玲8月月度汇报(2).pptx'

p = Presentation(SRC)
print('SLIDE COUNT:', len(p.slides))
print('SIZE: %.2f x %.2f in' % (p.slide_width/914400, p.slide_height/914400))
print()

fills = collections.Counter()
textcols = collections.Counter()
for s in p.slides:
    for sh in s.shapes:
        try:
            if sh.fill.type == 1:
                fills['#'+str(sh.fill.fore_color.rgb)] += 1
        except Exception:
            pass
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    try:
                        if c.fill.type == 1:
                            fills['#'+str(c.fill.fore_color.rgb)] += 1
                    except Exception:
                        pass
        if sh.has_chart:
            for ser in sh.chart.series:
                try:
                    if ser.format.fill.type == 1:
                        fills['#'+str(ser.format.fill.fore_color.rgb)] += 1
                except Exception:
                    pass
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        if run.font.color and run.font.color.type is not None:
                            textcols['#'+str(run.font.color.rgb)] += 1
                    except Exception:
                        pass

print('=== FILL COLORS ===')
for c, n in fills.most_common():
    print('  %s : %d' % (c, n))
print()
print('=== TEXT COLORS ===')
for c, n in textcols.most_common():
    print('  %s : %d' % (c, n))
print()

print('=== BOTTOM-RIGHT PAGE NUMBERS (top>6.0, left>9.0) ===')
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        top = sh.top/914400 if sh.top is not None else None
        left = sh.left/914400 if sh.left is not None else None
        txt = ' '.join(pp.text for pp in sh.text_frame.paragraphs if pp.text.strip())
        if top is not None and left is not None and top > 6.0 and left > 9.0:
            print('  P%d  L=%.2f T=%.2f  txt=[%s]' % (i, left, top, txt))
print()
print('=== ALL text boxes with digits near bottom (top>6.3) ===')
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        top = sh.top/914400 if sh.top is not None else None
        txt = ' '.join(pp.text for pp in sh.text_frame.paragraphs if pp.text.strip())
        if top is not None and top > 6.3 and any(ch.isdigit() for ch in txt):
            left = sh.left/914400 if sh.left is not None else None
            print('  P%d  L=%.2f T=%.2f  txt=[%s]' % (i, left, top, txt))
