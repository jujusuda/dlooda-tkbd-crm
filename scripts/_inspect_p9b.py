# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.oxml.ns import qn

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(9).pptx"

def c(t): return qn('c:' + t)
def a(t): return qn('a:' + t)

prs = Presentation(SRC)
print("slides:", len(prs.slides))
sl = prs.slides[8]
print("=" * 90)
for i, sh in enumerate(sl.shapes):
    print(f"\n[{i:2d}] {sh.name:<22} type={sh.shape_type} L={sh.left} T={sh.top} W={sh.width} H={sh.height}")
    if sh.has_chart:
        ch = sh.chart
        pa = ch._chartSpace.plotArea
        print("     plotArea:", [e.tag.split('}')[1] for e in pa])
        for nm in ('barChart', 'lineChart'):
            p = pa.find(c(nm))
            if p is None:
                continue
            print(f"     {nm} axIds={[x.get('val') for x in p.findall(c('axId'))]}")
            for s in p.findall(c('ser')):
                tx = s.find(c('tx'))
                nmv = tx.find(c('strRef')).find(c('strCache')).find(c('pt')).find(c('v')).text
                vals = [pt.find(c('v')).text for pt in
                        s.find(c('val')).find(c('numRef')).find(c('numCache')).findall(c('pt'))]
                print(f"        {nmv:<24} {vals}")
        for va in pa.findall(c('valAx')):
            nf = va.find(c('numFmt'))
            print(f"     valAx {va.find(c('axId')).get('val')} pos={va.find(c('axPos')).get('val')} "
                  f"numFmt={nf.get('formatCode') if nf is not None else None}")
    if sh.has_table:
        print("     TABLE", len(sh.table.rows), "x", len(sh.table.columns))
        for r in range(len(sh.table.rows)):
            print("      ", [sh.table.cell(r, cc).text for cc in range(len(sh.table.columns))])
    if sh.has_text_frame:
        tf = sh.text_frame
        print("     TEXT:", repr(tf.text[:200]))
        for pi, para in enumerate(tf.paragraphs):
            runs = para.runs
            sizes = sorted({(r.font.size.pt if r.font.size else None) for r in runs})
            print(f"       p{pi}: sz={sizes} n_runs={len(runs)} text={para.text[:60]!r}")
