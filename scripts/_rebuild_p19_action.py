# -*- coding: utf-8 -*-
"""下月行动页（PART 04·下月行动·slide idx 18）v2 重排：按 SKU 逐行 + 删不寄样 + 加新增测品板块。
基于 v9_14，仅改第 19 页（idx 18）。输出 v9_15。
数据均为用户已核对真实值，不改动、不编造。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INK   = RGBColor.from_string('1F1F1F')
BODY  = RGBColor.from_string('3D3D3D')
MUTE  = RGBColor.from_string('9A9A9A')
ACC   = RGBColor.from_string('B06579')   # 玫红强调
GRAY_D= RGBColor.from_string('6B6B6B')
FONT  = '微软雅黑'

WHITE = RGBColor.from_string('FFFFFF')
CARD  = RGBColor.from_string('F7F4F5')
HDR   = RGBColor.from_string('EFE7EA')
ZEB   = RGBColor.from_string('FAF7F8')
LINE  = RGBColor.from_string('E6E0E2')

SRC = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_14.pptx'
OUT = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_15.pptx'

prs = Presentation(SRC)
s = prs.slides[18]

for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)

def add_text(x, y, w, h, lines, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    first = True
    for para in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(3); p.space_before = Pt(0)
        for (t, sz, b, col) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = b; r.font.name = FONT
            r.font.color.rgb = col
    return tb

def style_cell(cell, text, size, bold, color, fill, align=PP_ALIGN.CENTER):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.05); cell.margin_right = Inches(0.05)
    cell.margin_top = Inches(0.004); cell.margin_bottom = Inches(0.004)
    tf = cell.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
    r.font.color.rgb = color
    cell.fill.solid(); cell.fill.fore_color.rgb = fill

# ---------- 1. 顶部标题区 ----------
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.34), Inches(0.1), Inches(0.74)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = ACC; s.shapes[-1].line.fill.background()
add_text(0.84, 0.26, 11.9, 0.66, [[('下月行动：爆品守量 + 测品扩测', 23, True, INK)]], anchor=MSO_ANCHOR.MIDDLE)
add_text(9.3, 0.42, 3.43, 0.5, [[('PART 04 · 下月行动', 12, False, MUTE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
add_text(0.84, 1.12, 11.89, 0.36, [[('把 9 月数据结论翻译成动作；结构从「爆品为主」转向「爆品守量 + 测品扩测」', 15, False, GRAY_D)]], anchor=MSO_ANCHOR.MIDDLE)
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.6), Inches(12.13), Inches(0.012)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = LINE; s.shapes[-1].line.fill.background()

# ---------- 2. 数字卡（560 件 / +27.0%） ----------
s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.7), Inches(12.13), Inches(0.92))
sh = s.shapes[-1]; sh.fill.solid(); sh.fill.fore_color.rgb = CARD
sh.line.color.rgb = LINE; sh.line.width = Pt(0.75)
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.84), Inches(0.09), Inches(0.64)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = ACC; s.shapes[-1].line.fill.background()
add_text(0.95, 1.74, 4.6, 0.84, [
    [('560 件', 38, True, ACC)],
], anchor=MSO_ANCHOR.MIDDLE)
add_text(0.97, 2.30, 4.6, 0.26, [[('10 月寄样目标', 14, False, BODY)]], anchor=MSO_ANCHOR.MIDDLE)
# +27.0% pill (right)
s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.75), Inches(1.9), Inches(2.7), Inches(0.5))
pill = s.shapes[-1]; pill.fill.solid(); pill.fill.fore_color.rgb = ACC; pill.line.fill.background()
add_text(9.75, 1.9, 2.7, 0.5, [[('+27.0%', 19, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
add_text(9.75, 2.42, 2.7, 0.26, [[('9 月 441 → 10 月 560', 11, False, MUTE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

# ---------- 3. 主表：下月寄样 SKU（爆品） ----------
add_text(0.6, 2.62, 6.0, 0.26, [[('下月寄样 SKU（爆品）', 13, True, ACC)]], anchor=MSO_ANCHOR.MIDDLE)
main_rows = [
    ('2190',  '扩量', '保持并适度扩寄样',            '40'),
    ('28173', '放量', '继续放大寄样、稳住出单率 30-34%', '70'),
    ('2178',  '守量', '加寄样兑现出单率',            '40'),
    ('2189',  '补量', '补视频与达人供给',            '40'),
    ('合计',  '',     '',                           '190'),
]
gf = s.shapes.add_table(6, 4, Inches(0.6), Inches(2.9), Inches(6.0), Inches(2.52))
tb = gf.table
for i, cw_i in enumerate([1.2, 0.95, 2.95, 0.9]):
    tb.columns[i].width = Inches(cw_i)
hdr = ['SKU', '方向', '动作', '下月寄样目标']
for ci, h in enumerate(hdr):
    style_cell(tb.cell(0, ci), h, 12, True, INK, HDR)
for ri, row in enumerate(main_rows, start=1):
    tb.rows[ri].height = Inches(0.38) if ri == 5 else Inches(0.44)
    is_total = (ri == 5)
    fill = HDR if is_total else (WHITE if ri % 2 == 1 else ZEB)
    for ci, val in enumerate(row):
        bold = is_total or (ci in (0, 3))
        col = INK if is_total else (ACC if ci == 3 else (BODY if ci != 1 else ACC))
        if ci == 0 and not is_total:
            col = INK
        if ci == 1 and not is_total:
            col = ACC; bold = True
        align = PP_ALIGN.LEFT if ci == 2 else PP_ALIGN.CENTER
        style_cell(tb.cell(ri, ci), val, 12, bold, col, fill, align)

# ---------- 4. 新增板块：下月新增测品寄样目标 ----------
add_text(6.9, 2.62, 6.0, 0.26, [[('下月新增测品寄样目标', 13, True, ACC)]], anchor=MSO_ANCHOR.MIDDLE)
test_rows = [
    ('测款', '2A16', '100'),
    ('测款', '2F08', '20'),
    ('测款', '2A51', '50'),
    ('测款', '2A50', '50'),
    ('测款', '2H43', '50'),
    ('测款', '2Q02', '50'),
    ('测款', '2C33', '50'),
    ('合计', '',    '370'),
]
gf2 = s.shapes.add_table(9, 3, Inches(6.9), Inches(2.9), Inches(6.0), Inches(2.86))
tb2 = gf2.table
for i, cw_i in enumerate([1.3, 2.0, 2.7]):
    tb2.columns[i].width = Inches(cw_i)
hdr2 = ['定位', 'SKU', '月目标']
for ci, h in enumerate(hdr2):
    style_cell(tb2.cell(0, ci), h, 12, True, INK, HDR)
for ri, row in enumerate(test_rows, start=1):
    is_total = (ri == 8)
    tb2.rows[ri].height = Inches(0.38) if is_total else Inches(0.30)
    fill = HDR if is_total else (WHITE if ri % 2 == 1 else ZEB)
    for ci, val in enumerate(row):
        bold = is_total or (ci in (1, 2))
        col = INK if is_total else (ACC if ci == 2 else (BODY if ci == 0 else INK))
        align = PP_ALIGN.LEFT if ci != 2 else PP_ALIGN.CENTER
        style_cell(tb2.cell(ri, ci), val, 12, bold, col, fill, align)

# 测品板块小结小字（合计行下方右侧）
add_text(6.9, 5.80, 6.0, 0.3, [[('测品合计 370 件　|　总目标 = 爆品 190 + 测品 370 = 560 件', 11.5, True, ACC)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)

# ---------- 5. 口径备注（保留，底部小字） ----------
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(6.32), Inches(0.09), Inches(0.34)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = ACC; s.shapes[-1].line.fill.background()
add_text(0.84, 6.32, 11.7, 0.34, [[
    ('※ 目标口径同「本月新增寄样记录」，与概览表寄样端口径不同', 11, False, MUTE),
]], anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print('SAVED:', OUT, 'slides:', len(prs.slides._sldIdLst))
