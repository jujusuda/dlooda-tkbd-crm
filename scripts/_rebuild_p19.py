# -*- coding: utf-8 -*-
"""重排第19页：左问题 / 右策略 3行，沿用原主题配色。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC = r'E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-09\8be9710a0a0dd58ff4cb419316577cc7\8月余慧玲8月月度汇报.pptx'
OUT = r'C:\Users\LENOVO\Desktop\8月余慧玲8月月度汇报_第19页改.pptx'

# 原主题色
TITLE_BAR   = 'C94F82'
PROB_ACC    = 'C94F82'   # 左：问题 强调色
STRAT_ACC   = 'F5A0BC'   # 右：策略 强调色
CARD_FILL   = 'FFFFFF'
CARD_LINE   = 'FCD5E3'
HEAD_PROB   = 'C94F82'
HEAD_STRAT  = 'E879A6'
BODY        = '4A3D42'

p = Presentation(SRC)
slide = list(p.slides)[18]

# ---- 1. 删除内容卡片，保留标题栏 + 页码 ----
keep = []
for sh in slide.shapes:
    top = sh.top / 914400 if sh.top is not None else 0
    txt = sh.text_frame.text.strip() if sh.has_text_frame else ''
    if top < 1.1 or txt == '20':
        keep.append(sh)

for sh in list(slide.shapes):
    if sh not in keep:
        sh._element.getparent().remove(sh._element)

# ---- 2. 改标题文字（保留原格式）----
title_sh = None
for sh in slide.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip().startswith('下月行动'):
        title_sh = sh
        break
if title_sh:
    para0 = title_sh.text_frame.paragraphs[0]
    r0 = para0.runs[0] if para0.runs else None
    fs = r0.font.size if r0 else Pt(24)
    fb = r0.font.bold if r0 else True
    fc = r0.font.color.rgb if (r0 and r0.font.color and r0.font.color.type is not None) else RGBColor.from_string('FFFFFF')
    fn = r0.font.name if r0 else '微软雅黑'
    for para in list(title_sh.text_frame.paragraphs):
        for run in list(para.runs):
            run._r.getparent().remove(run._r)
    nr = para0.add_run()
    nr.text = '下月行动方向'
    nr.font.size = fs
    nr.font.bold = fb
    nr.font.color.rgb = fc
    nr.font.name = fn

# ---- helpers ----
def rect(l, t, w, h, fill, line=None, lw=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(lw)
    else:
        sp.line.fill.background()
    try:
        sp.shadow.inherit = False
    except Exception:
        pass
    return sp

def textbox(l, t, w, h, paras, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for i, para in enumerate(paras):
        pp = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pp.alignment = para.get('align', PP_ALIGN.LEFT)
        pp.space_after = Pt(para.get('sa', 2))
        pp.space_before = Pt(para.get('sb', 0))
        if para.get('line'):
            pp.line_spacing = para['line']
        for rn in para['runs']:
            r = pp.add_run()
            r.text = rn['t']
            r.font.size = Pt(rn.get('s', 11))
            r.font.bold = rn.get('b', False)
            r.font.name = '微软雅黑'
            if rn.get('c'):
                r.font.color.rgb = RGBColor.from_string(rn['c'])
    return tb

# ---- 3. 列头 ----
rect(0.60, 1.22, 6.00, 0.46, TITLE_BAR)
textbox(0.60, 1.22, 6.00, 0.46,
        [{'runs': [{'t': '普遍问题', 's': 13, 'b': True, 'c': 'FFFFFF'}], 'align': PP_ALIGN.CENTER}],
        anchor=MSO_ANCHOR.MIDDLE)
rect(6.85, 1.22, 6.00, 0.46, STRAT_ACC)
textbox(6.85, 1.22, 6.00, 0.46,
        [{'runs': [{'t': '下月策略', 's': 13, 'b': True, 'c': 'FFFFFF'}], 'align': PP_ALIGN.CENTER}],
        anchor=MSO_ANCHOR.MIDDLE)

# ---- 4. 三行 左问题 / 右策略 ----
rows = [
    ('①', '出单覆盖面在收缩',
     '多数 SKU 出单达人数环比未增、甚至下降，增长主要靠历史达人老视频持续回单撑着。',
     '▸ 扩大出单达人覆盖面',
     '持续拓展新达人，并重新激活历史已出单达人。'),
    ('②', '达人流失风险',
     '高产出达人一旦停更或流失，对应 SKU 订单即明显回落，抗风险能力弱。',
     '▸ 培育第二梯队',
     '把寄样与沟通重心向中腰部达人倾斜，降低对个别达人的依赖。'),
    ('③', '复投效果不明显',
     '复投动作一直有推进，但路径单一、转化有限，已出单达人的复利没充分释放。',
     '▸ 增加复投路径',
     '多轮触达 + 重点款追加激励，提升已出单达人二次寄样与持续出单比例。'),
]

RH = 1.62
GAP = 0.16
TOP0 = 1.82
for i, (num, p_title, p_body, s_title, s_body) in enumerate(rows):
    top = TOP0 + i * (RH + GAP)
    # 左卡
    rect(0.60, top, 6.00, RH, CARD_FILL, line=CARD_LINE, lw=1.0)
    rect(0.60, top, 0.14, RH, PROB_ACC)
    textbox(0.86, top + 0.13, 5.55, 0.44,
            [{'runs': [{'t': '%s %s' % (num, p_title), 's': 13, 'b': True, 'c': HEAD_PROB}]}])
    textbox(0.86, top + 0.60, 5.55, RH - 0.70,
            [{'runs': [{'t': p_body, 's': 10.5, 'c': BODY}], 'line': 1.18}])
    # 右卡
    rect(6.85, top, 6.00, RH, CARD_FILL, line=CARD_LINE, lw=1.0)
    rect(6.85, top, 0.14, RH, STRAT_ACC)
    textbox(7.11, top + 0.13, 5.55, 0.44,
            [{'runs': [{'t': s_title, 's': 13, 'b': True, 'c': HEAD_STRAT}]}])
    textbox(7.11, top + 0.60, 5.55, RH - 0.70,
            [{'runs': [{'t': s_body, 's': 10.5, 'c': BODY}], 'line': 1.18}])

p.save(OUT)
print('SAVED:', OUT)

# 校验
p2 = Presentation(OUT)
s2 = list(p2.slides)[18]
print('--- 第19页 shapes ---')
for sh in s2.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print('  [%s]' % sh.text_frame.text.strip()[:50])
