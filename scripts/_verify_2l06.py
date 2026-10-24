# -*- coding: utf-8 -*-
from pptx import Presentation
import re, zipfile, posixpath
F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_10.pptx'
prs = Presentation(F)
s = prs.slides[11]

print("=== 2L06 page text ===")
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print(repr(sh.text_frame.text[:140]))

print("\n=== COMPARISON TABLE ===")
for sh in s.shapes:
    if sh.has_table:
        tb = sh.table
        for r in range(len(tb.rows)):
            row=[]
            for c in range(len(tb.columns)):
                cell = tb.cell(r,c)
                rose=False
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.bold and run.font.color and run.font.color.type is not None:
                            try:
                                if str(run.font.color.rgb)=='B06579': rose=True
                            except: pass
                row.append((cell.text, '★' if rose else ''))
            print("  ", row)

# package + overflow
z = zipfile.ZipFile(F)
names=set(z.namelist())
ct=z.read('[Content_Types].xml').decode('utf-8','ignore')
overrides=set(re.findall(r'PartName="([^"]+)"',ct)); defaults=set(re.findall(r'Extension="([^"]+)"',ct))
probs=[]
for o in overrides:
    if o.lstrip('/') not in names: probs.append('CT missing '+o)
for n in names:
    if n=='[Content_Types].xml': continue
    ext=n.rsplit('.',1)[-1].lower()
    if ('/'+n) not in overrides and ext not in defaults: probs.append('part '+n)
    if n.endswith('.rels'):
        base=posixpath.dirname(posixpath.dirname(n))
        for rid,tgt in re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"',z.read(n).decode('utf-8','ignore')):
            if tgt.startswith('http'): continue
            fp=tgt.lstrip('/') if tgt.startswith('/') else posixpath.normpath(posixpath.join(base,tgt))
            if fp not in names: probs.append(f'{n} {rid}->{tgt} MISSING {fp}')
print("\nPACKAGE PROBLEMS:", probs if probs else "none")
for sh in s.shapes:
    if sh.left is None: continue
    if (sh.left/914400+sh.width/914400)>13.38 or (sh.top/914400+sh.height/914400)>7.55:
        print("OVERFLOW:", sh.shape_type, round(sh.left/914400,2), round(sh.top/914400,2), round(sh.width/914400,2), round(sh.height/914400,2))
print("overflow check done")
