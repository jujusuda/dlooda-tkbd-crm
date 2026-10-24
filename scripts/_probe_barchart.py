from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(7).pptx"
prs = Presentation(SRC)
slide = prs.slides[8]

cd = CategoryChartData()
cd.categories = ['7月','8月','9月']
cd.add_series('2189 出单达人数量', (14,8,14))
cd.add_series('28173 出单达人数量', (16,25,24))
cd.add_series('2189 达人出单率', (56.00,22.86,31.11))
cd.add_series('28173 达人出单率', (30.19,33.78,31.58))
cd.add_series('2189 视频出单率', (33.33,28.41,27.85))
cd.add_series('28173 视频出单率', (18.58,23.84,23.03))

gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.82),
                            Inches(4.2), Inches(2.66), cd)
ch = gf.chart

plotArea = ch._chartSpace.plotArea
print("plotArea children:")
for el in plotArea:
    tag = el.tag.split('}')[1]
    if tag == 'ser':
        continue
    print("  ", tag, {k:v for k,v in el.attrib.items()} if el.attrib else '')
    if tag.endswith('Chart'):
        for sub in el:
            stag = sub.tag.split('}')[1]
            print("      ", stag, dict(sub.attrib) if sub.attrib else '')
            if stag == 'ser':
                print("          ser children:", [s.tag.split('}')[1] for s in sub])
