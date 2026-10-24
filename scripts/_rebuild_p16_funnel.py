# -*- coding: utf-8 -*-
"""达人建联页（PART 03·达人端·达人建联页，slide idx 15）整改：
建联链路 → 倒三角漏斗（27→20→17）；删重复指标卡；WhatsApp深绑/付费达人独立图形化。
基于 v9_12，仅改第 16 页。输出 v9_13。数据均为用户已核对真实值，不改动、不编造。
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INK   = RGBColor.from_string('1F1F1F')
BODY  = RGBColor.from_string('3D3D3D')
MUTE  = RGBColor.from_string('9A9A9A')
ACC   = RGBColor.from_string('B06579')   # 玫红强调
GRAY_D= RGBColor.from_string('6B6B6B')
FONT  = '微软雅黑'

WHITE = RGBColor.from_string('FFFFFF')
CARD  = RGBColor.from_string('F7F4F5')
HDR   = RGBColor.from_string('EFE7EA')
ZEB   = RGBColor.from_string('FAF7F8')
LINE  = RGBColor.from_string('E6E0E2')
FUN2  = RGBColor.from_string('A85C73')   # 漏斗中层
FUN3  = RGBColor.from_string('9A4F66')   # 漏斗下层

SRC = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_12.pptx'
OUT = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_13.pptx'

prs = Presentation(SRC)
s = prs.slides[15]

# 删除当前页所有 shape
for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)

def add_text(x, y, w, h, lines, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    first = True
    for para in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(3); p.space_before = Pt(0)
        for (t, sz, b, col) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = b; r.font.name = FONT
            r.font.color.rgb = col
    return tb

def style_cell(cell, text, size, bold, color, fill):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.04); cell.margin_right = Inches(0.04)
    cell.margin_top = Inches(0.004); cell.margin_bottom = Inches(0.004)
    tf = cell.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
    r.font.color.rgb = color
    cell.fill.solid(); cell.fill.fore_color.rgb = fill

# ---------- 1. 顶部标题区 ----------
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.34), Inches(0.1), Inches(0.74)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = ACC; s.shapes[-1].line.fill.background()
add_text(0.84, 0.26, 11.9, 0.66, [[('达人建联：建联 27 → 已发视频 20 → 已出单 17', 24, True, INK)]], anchor=MSO_ANCHOR.MIDDLE)
add_text(9.3, 0.40, 3.43, 0.5, [[('PART 03 · 达人端', 12, False, MUTE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
add_text(0.84, 1.0, 11.89, 0.36, [[('抖音号 / 官号 / 邮件私信建联；四款爆品共成功建联 27 位，其中 17 位已出单', 15, False, GRAY_D)]], anchor=MSO_ANCHOR.MIDDLE)
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.5), Inches(12.13), Inches(0.012)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = LINE; s.shapes[-1].line.fill.background()

# ---------- 2. 主漏斗（倒三角，3 层梯形） ----------
CX = 6.6665
SCALE = 0.17778   # 每单位宽度（英寸）：27→4.8
def half(v): return (v * SCALE) / 2.0

# (val_top, val_bot, y_top, y_bot, fill)
layers = [
    (27, 20, 1.64, 2.30, ACC),
    (20, 17, 2.30, 2.96, FUN2),
    (17, 11, 2.96, 3.62, FUN3),
]
for (vt, vb, yt, yb, fill) in layers:
    ht = half(vt); hb = half(vb)
    builder = s.shapes.build_freeform(Inches(CX - ht), Inches(yt))
    builder.add_line_segments([
        (Inches(CX + ht), Inches(yt)),   # 上右
        (Inches(CX + hb), Inches(yb)),   # 下右
        (Inches(CX - hb), Inches(yb)),   # 下左
    ], close=True)
    shp = builder.convert_to_shape()
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = WHITE; shp.line.width = Pt(1.5)
    # 层内文字（数字 + 标签），居中
    mid_w = (ht + hb)  # 半宽均值的 2 倍 ≈ 中部宽度
    ym = (yt + yb) / 2.0
    label = {27: '成功建联', 20: '已发视频', 17: '已出单'}[vt]
    num = str(vt)
    add_text(CX - mid_w, ym - 0.32, mid_w * 2, 0.64, [
        [(num, 30, True, WHITE)],
        [(label, 12.5, True, WHITE)],
    ], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

# ---------- 3. 层间转化率标注（漏斗右侧，3 张小卡） ----------
callouts = [
    (1.80, '建联→发视频  74.1%', '流失 26%（最大环节）'),
    (2.55, '发视频→出单  85.0%', '转化健康'),
    (3.30, '整体建联→出单  63.0%', '建联→出单约六成'),
]
for (cy, big, small) in callouts:
    s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.2), Inches(cy), Inches(3.5), Inches(0.62))
    sh = s.shapes[-1]; sh.fill.solid(); sh.fill.fore_color.rgb = CARD
    sh.line.color.rgb = LINE; sh.line.width = Pt(0.75)
    add_text(9.35, cy + 0.02, 3.25, 0.6, [
        [(big, 13.5, True, ACC)],
        [(small, 10.5, False, GRAY_D)],
    ], anchor=MSO_ANCHOR.MIDDLE)

# ---------- 4. 漏斗下方：转化分析 caption ----------
add_text(0.6, 3.74, 12.13, 0.34, [[
    ('转化分析：', 12.5, True, ACC),
    ('建联→发视频 流失最大（约 26%），是下一步要重点抓的环节——催发视频、筛选高意愿达人；发视频→出单 85% 较健康。', 12.5, False, BODY),
]], anchor=MSO_ANCHOR.MIDDLE)

# ---------- 5. 各品明细表（左侧） ----------
add_text(0.6, 4.06, 6.0, 0.26, [[('各品建联明细', 12.5, True, ACC)]], anchor=MSO_ANCHOR.MIDDLE)
rows = [
    ('2178',  '9',  '7',  '7'),
    ('2190',  '13', '9',  '6'),
    ('28173', '3',  '2',  '2'),
    ('2189',  '2',  '2',  '2'),
    ('合计',  '27', '20', '17'),
]
gf = s.shapes.add_table(6, 4, Inches(0.6), Inches(4.32), Inches(6.0), Inches(1.74))
tb = gf.table
for i, cw_i in enumerate([1.5, 1.5, 1.8, 1.2]):
    tb.columns[i].width = Inches(cw_i)
hdr = ['品', '建联', '已发视频', '已出单']
for ci, h in enumerate(hdr):
    style_cell(tb.cell(0, ci), h, 11.5, True, INK, HDR)
for ri, row in enumerate(rows, start=1):
    tb.rows[ri].height = Inches(0.30) if ri <= 4 else Inches(0.32)
    is_total = (ri == 5)
    for ci, val in enumerate(row):
        fill = WHITE if ri % 2 == 1 else ZEB
        if is_total:
            fill = HDR
        bold = is_total or (ci >= 1)  # 数值列加粗
        col = ACC if (ci >= 1 and not is_total) else (INK if is_total else BODY)
        style_cell(tb.cell(ri, ci), val, 11.5, bold, col, fill)

# ---------- 6. 渠道 / 深绑（右侧，独立图形化小卡，不进漏斗） ----------
add_text(7.0, 4.06, 5.7, 0.26, [[('渠道 / 深绑（独立，不进主漏斗）', 12.5, True, ACC)]], anchor=MSO_ANCHOR.MIDDLE)

def channel_card(cy, num, title, sub):
    s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(cy), Inches(5.7), Inches(0.82))
    sh = s.shapes[-1]; sh.fill.solid(); sh.fill.fore_color.rgb = CARD
    sh.line.color.rgb = LINE; sh.line.width = Pt(0.75)
    # 数字圆
    s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.2), Inches(cy + 0.12), Inches(0.58), Inches(0.58))
    ov = s.shapes[-1]; ov.fill.solid(); ov.fill.fore_color.rgb = ACC; ov.line.fill.background()
    add_text(7.2, cy + 0.12, 0.58, 0.58, [[(num, 22, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    add_text(7.95, cy + 0.08, 4.6, 0.66, [
        [(title, 14, True, INK)],
        [(sub, 11, False, GRAY_D)],
    ], anchor=MSO_ANCHOR.MIDDLE)

channel_card(4.32, '10', 'WhatsApp 深绑达人', '绑定更深，出单更稳定')
channel_card(5.24, '2',  '付费达人', '已寄样，待发视频（推进中）')

# ---------- 7. 落点策略 ----------
s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.30), Inches(12.13), Inches(0.62))
sh = s.shapes[-1]; sh.fill.solid(); sh.fill.fore_color.rgb = CARD
sh.line.color.rgb = LINE; sh.line.width = Pt(0.75)
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(6.44), Inches(0.08), Inches(0.34)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = ACC; s.shapes[-1].line.fill.background()
add_text(0.84, 6.30, 11.7, 0.62, [[
    ('落点策略　', 13, True, ACC),
    ('① 扩大建联规模　② WhatsApp 加深绑定　③ 推进付费达人合作落地', 13, True, ACC),
]], anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print('SAVED:', OUT, 'slides:', len(prs.slides._sldIdLst))
