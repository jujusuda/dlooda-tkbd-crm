# -*- coding: utf-8 -*-
"""P10（2189 + 28173｜同周期两款）二次整改：基于 v9_6，仅重建第 10 页（slide idx 9）。
新布局：
- 顶部一排：4 张折线图（达人出单率 / 出单达人数量(小一号) / 点击率 / 转化率）+ 右侧 1 个表格（流量与转化）
- 底部两列：2189（内容转化型）/ 28173（达人承接型）各自内容，含数据/变化/原因/卡点/策略，无独立策略块
配色/字号沿用 v9 规范。折线图标签 ABOVE/BELOW（禁用 outEnd/inEnd），横向条形已弃用。
产物 Desktop/余慧玲9月月度汇报_v9_7.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE

INK   = RGBColor.from_string('1F1F1F')
BODY  = RGBColor.from_string('3D3D3D')
MUTE  = RGBColor.from_string('9A9A9A')
ACC   = RGBColor.from_string('B06579')   # 2189 线/强调
ACC_L = RGBColor.from_string('C97F95')
CARD  = RGBColor.from_string('F9F5F6')
LINE  = RGBColor.from_string('EAE2E5')
GRAY_D= RGBColor.from_string('5A5A5A')   # 28173 线
WHITE = RGBColor.from_string('FFFFFF')
FONT  = '微软雅黑'
ZEBRA = RGBColor.from_string('F2ECEE')

SRC = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_6.pptx'
OUT = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_7.pptx'

prs = Presentation(SRC)
s = prs.slides[9]  # P10

for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)

def rect(x, y, w, h, fill, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp

def add_text(x, y, w, h, paras, anchor=None, align=None, wrap=True, spacing=1.0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.margin_left = Pt(2); tf.margin_right = Pt(2); tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    if anchor is not None: tf.vertical_anchor = anchor
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = spacing
        if align is not None: p.alignment = align
        for (t, sz, bold, color) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = FONT
    return tb

# ---------- 标题区 ----------
rect(0.6, 0.34, 0.1, 0.74, ACC)
add_text(0.84, 0.30, 11.7, 0.66,
         [[('2189 + 28173｜同周期两款：一个靠内容、一个靠达人', 24, True, INK)]], anchor=MSO_ANCHOR.MIDDLE)
add_text(9.3, 0.42, 3.43, 0.56,
         [[('PART 02 · 各款 SKU 分析', 12, False, MUTE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
add_text(0.84, 1.02, 11.8, 0.34,
         [[('外形相似、7 月同推；2189 内容一转就转化，28173 靠达人承接放量', 15, False, BODY)]])
rect(0.6, 1.44, 12.13, Pt(1)/914400, LINE)

# ---------- 折线图（顶部一排） ----------
def add_line_chart(x, y, w, h, cats, series, colors, vmax, vfmt='0.00"%"'):
    cd = CategoryChartData(); cd.categories = cats
    for (nm, vals) in series: cd.add_series(nm, vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(8); ch.legend.font.name = FONT
    pl = ch.plots[0]
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.show_value = True; dl.number_format = vfmt; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.ABOVE
    dl.font.size = Pt(8); dl.font.bold = False; dl.font.name = FONT; dl.font.color.rgb = INK
    for i, c in enumerate(colors):
        se = pl.series[i]
        se.format.line.color.rgb = c
        se.format.line.width = Pt(2.0)
        se.marker.style = XL_MARKER_STYLE.CIRCLE
        se.marker.size = 5
        se.marker.format.fill.solid(); se.marker.format.fill.fore_color.rgb = c
        se.marker.format.line.color.rgb = WHITE; se.marker.format.line.width = Pt(1)
        sd = se.data_labels
        sd.show_value = True; sd.number_format = vfmt; sd.number_format_is_linked = False
        # 2189 上方(t)、28173 下方(b)：竖直错开避免叠字
        sd.position = XL_LABEL_POSITION.BELOW if i == 1 else XL_LABEL_POSITION.ABOVE
        sd.font.size = Pt(8); sd.font.bold = False; sd.font.name = FONT; sd.font.color.rgb = INK
    va = ch.value_axis
    va.minimum_scale = 0; va.maximum_scale = vmax
    va.major_unit = 15 if vmax >= 60 else (10 if vmax >= 30 else 1)
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.major_gridlines.format.line.width = Pt(0.5)
    va.tick_labels.font.size = Pt(8); va.tick_labels.font.name = FONT
    va.tick_labels.font.color.rgb = MUTE; va.tick_labels.number_format = vfmt
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(10); ca.tick_labels.font.name = FONT; ca.tick_labels.font.color.rgb = GRAY_D
    ca.format.line.color.rgb = LINE
    return ch

CHY = 1.84; CHH = 2.22
# 图一 达人出单率（正常）
add_text(0.60, 1.56, 2.30, 0.24, [[('达人出单率', 12.5, True, INK)]])
add_line_chart(0.60, CHY, 2.30, CHH,
               ['7月', '8月', '9月'],
               [('2189', [56.00, 22.86, 31.11]),
                ('28173', [30.19, 33.78, 31.58])],
               [ACC, GRAY_D], 60)
# 图二 出单达人数量（小一号）
add_text(2.945, 1.62, 1.95, 0.24, [[('出单达人数量', 12.5, True, INK)]])
add_line_chart(2.945, CHY+0.22, 1.95, CHH-0.22,
               ['7月', '8月', '9月'],
               [('2189', [14, 8, 14]),
                ('28173', [16, 25, 24])],
               [ACC, GRAY_D], 30, vfmt='0')
# 图三 点击率（正常）
add_text(4.94, 1.56, 2.30, 0.24, [[('点击率', 12.5, True, INK)]])
add_line_chart(4.94, CHY, 2.30, CHH,
               ['7月', '8月', '9月'],
               [('2189', [3.49, 2.96, 5.17]),
                ('28173', [4.22, 3.33, 4.06])],
               [ACC, GRAY_D], 6)
# 图四 转化率（正常）
add_text(7.285, 1.56, 2.30, 0.24, [[('转化率', 12.5, True, INK)]])
add_line_chart(7.285, CHY, 2.30, CHH,
               ['7月', '8月', '9月'],
               [('2189', [2.37, 2.08, 1.13]),
                ('28173', [2.11, 1.73, 0.90])],
               [ACC, GRAY_D], 3)

# ---------- 右侧表格（流量与转化） ----------
def set_cell(cell, text, size, bold, color, fill=None, align=PP_ALIGN.CENTER):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Pt(2); cell.margin_right = Pt(2); cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
    if fill is not None:
        cell.fill.solid(); cell.fill.fore_color.rgb = fill
    else:
        cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
    tf = cell.text_frame
    for p in tf.paragraphs:
        p.alignment = align
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT

add_text(9.63, 1.56, 3.10, 0.24, [[('流量与转化（7–9 月）', 12.5, True, INK)]])
TX, TY, TW, TH = 9.63, CHY, 3.10, CHH
gt = s.shapes.add_table(7, 5, Inches(TX), Inches(TY), Inches(TW), Inches(TH))
tb = gt.table
col_w = [0.52, 0.5, 0.78, 0.78, 0.62]
for i, wd in enumerate(col_w):
    tb.columns[i].width = Inches(wd)
for r in range(7):
    tb.rows[r].height = Inches(TH/7)
hdr = ['款', '月份', '总播放量', '总点击次数', '转化单数']
for c, htxt in enumerate(hdr):
    set_cell(tb.cell(0, c), htxt, 9.5, True, WHITE, fill=ACC_L)
data = [
    ('2189', '7月', '75,037', '2,618', '62'),
    ('2189', '8月', '496,850', '14,682', '306'),
    ('2189', '9月', '330,070', '17,074', '193'),
    ('28173', '7月', '133,700', '5,643', '119'),
    ('28173', '8月', '537,860', '17,904', '310'),
    ('28173', '9月', '307,580', '12,482', '112'),
]
for ri, (a, b, c, d, e) in enumerate(data, start=1):
    zf = ZEBRA if a == '28173' else WHITE
    set_cell(tb.cell(ri, 0), a, 9, False, BODY, fill=zf)
    set_cell(tb.cell(ri, 1), b, 9, False, BODY, fill=zf)
    set_cell(tb.cell(ri, 2), c, 9, False, BODY, fill=zf)
    set_cell(tb.cell(ri, 3), d, 9, False, BODY, fill=zf)
    set_cell(tb.cell(ri, 4), e, 9, False, BODY, fill=zf)

# ---------- 底部两列：2189 / 28173 ----------
COL_Y = 4.24; COL_H = 2.96
def build_column(x, color, header, blocks):
    rect(x, COL_Y, 5.97, COL_H, CARD)
    rect(x, COL_Y, 0.08, COL_H, color)
    add_text(x+0.22, COL_Y+0.08, 5.7, 0.3, [[(header, 13.5, True, INK)]])
    paras = []
    for (lab, txt) in blocks:
        paras.append([(lab, 12, True, color), (txt, 12, False, BODY)])
    add_text(x+0.22, COL_Y+0.42, 5.66, COL_H-0.5, paras, anchor=MSO_ANCHOR.TOP, spacing=1.04)

build_column(0.60, ACC,
    '2189｜内容转化型（缺供给）',
    [('数据', '：寄样 52（履约 46 / 88.46%）· 出单达人 14 · 达人出单率 31.11% · 视频 79（出单 22 / 27.85%）· 月度 999 单 · 自然单占比 71.35%'),
     ('变化', '：月度出单 176→723→999 三连升；达人出单率 56.00%→22.86%→31.11% 波动大'),
     ('原因', '：内容一转就转化（点击率、转化率、视频出单率均高、单条效率强）；但视频量不足（仅 79 条）、寄样基数小、出单达人少（14→8→14），出单率受单笔达人影响大'),
     ('卡点', '：有转化力、缺供给（视频量少、寄样偏小）'),
     ('策略', '：① 补视频 / 达人供给　② 释放最高出单率')])

build_column(6.77, GRAY_D,
    '28173｜达人承接型（可放量）',
    [('数据', '：寄样 95（履约 77 / 81.05%）· 出单达人 24 · 达人出单率 31.58% · 视频 152（出单 35 / 23.03%）· 月度 761 单 · 自然单占比 76.85%'),
     ('变化', '：寄样逐月加码（67→86→95），出单达人稳定（16→25→24），达人出单率稳在 30–34%；自然单占比 60.70%→81.60%→76.85%（高位）'),
     ('原因', '：多寄不稀释出单效率、达人真心愿带；9 月出单率小幅回落（33.78→31.58）是因寄样加码新增达人多未出单（沉淀期）+ 视频量下降（172→152），非转化崩；点击率 4.06% 偏低、不靠内容点击，靠达人带量'),
     ('卡点', '：放量型，需继续放大寄样、稳住达人承接'),
     ('策略', '：① 放大寄样　② 复投承接强的核心达人　③ 加大素材拉回订单')])

prs.save(OUT)
print('SAVED:', OUT, 'slides:', len(prs.slides._sldIdLst))
