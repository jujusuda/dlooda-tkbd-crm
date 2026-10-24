# -*- coding: utf-8 -*-
import zipfile, re
from pptx import Presentation

p = "C:/Users/LENOVO/Desktop/余慧玲9月月度汇报_v9_12.pptx"

z = zipfile.ZipFile(p)
print("ZIP integrity:", "OK" if z.testzip() is None else "CORRUPT")

rels = z.read("ppt/_rels/presentation.xml.rels").decode("utf-8", "ignore")
targets = re.findall(r'Target="([^"]+)"', rels)
names = z.namelist()
miss = [t for t in targets if not t.startswith("http") and t not in names and ("ppt/" + t.lstrip("/")) not in names]
print("rels missing targets:", miss if miss else "none")

prs = Presentation(p)
print("slides:", len(prs.slides._sldIdLst))

s = prs.slides[14]
blob_lines = []
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        blob_lines.append(sh.text_frame.text.strip())
    if sh.has_table:
        tb = sh.table
        rows = [" | ".join(c.text for c in r.cells) for r in tb.rows]
        blob_lines.append("[T] " + " || ".join(rows))
    if sh.has_chart:
        ch = sh.chart
        cats = list(ch.plots[0].categories) if ch.plots else []
        vals = list(ch.series[0].values) if ch.series else []
        blob_lines.append("[CHART %s] cats=%s vals=%s title=%s" % (ch.chart_type, cats, vals, ch.chart_title.text_frame.text))

blob = "\n".join(blob_lines)
print("\n===== SLIDE 15 DUMP =====")
for t in blob_lines:
    print("  ", t[:160])

print("\n===== CHECKLIST =====")
checks = {
    "主标题": "达人素材端：供给稳定，四款爆品人均近 2 条" in blob,
    "副标题": "达人投稿活跃、素材持续积累，为内容侧健康提供支撑" in blob,
    "1.92 条/人": "1.92 条/人" in blob,
    "743 条": "743 条" in blob,
    "386 位": "386 位" in blob,
    "1.5–1.8 条": ("1.5" in blob and "1.8" in blob),
    "28173 1.79": "1.79" in blob,
    "2190 1.80": "1.80" in blob,
    "2178 1.53": "1.53" in blob,
    "2189 1.52": "1.52" in blob,
    "策略无付费达人素材合作": "付费达人素材合作" not in blob,
    "正向结论": ("供给稳定" in blob and "持续积累" in blob),
}
for k, v in checks.items():
    print(("  [OK] " if v else "  [!!] ") + k)

bad = ["近七成", "只发 1 条", "目标 3 条", "靠人头铺量", "单达人产出不足", "人均素材目标", "付费达人素材合作", "产出不足"]
found = [b for b in bad if b in blob]
print("  负面/禁用措辞命中:", found if found else "无（pass）")

charts = [n for n in names if "charts/chart" in n and n.endswith(".xml")]
for n in charts:
    d = z.read(n).decode("utf-8", "ignore")
    ctype = "line" if "<c:lineChart>" in d else ("bar" if "<c:barChart>" in d else ("pie/doughnut" if ("<c:pieChart>" in d or "<c:doughnutChart>" in d) else "other"))
    if ctype == "line" and ("outEnd" in d or "inEnd" in d):
        print("  [!!] line chart illegal label:", n)
print("  图表合法性扫描完成（仅折线图检查 outEnd/inEnd）")
