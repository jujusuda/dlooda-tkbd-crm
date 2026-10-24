import sys
from pptx import Presentation
from lxml import etree

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(7).pptx"
prs = Presentation(SRC)
slide = prs.slides[8]

def dump(idx, label):
    sh = slide.shapes[idx]
    print("="*90)
    print(f"CHART shape[{idx}] {label}  name={sh.name}")
    xml = etree.tostring(sh.chart._chartSpace, pretty_print=True).decode('utf-8')
    print(xml)

# 6 = 达人出单率, 8 = 出单达人数量, 10 = 点击率
for idx, lab in [(6,"达人出单率"),(8,"出单达人数量"),(10,"点击率")]:
    dump(idx, lab)
