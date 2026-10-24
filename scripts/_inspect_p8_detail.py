# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Emu, Pt
F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_1.pptx'
prs = Presentation(F)
s = prs.slides[7]
# header cell run style
tb = None
for sh in s.shapes:
    if sh.has_table:
        tb = sh.table; break
def hx(c):
    try: return '%06X'%c.rgb
    except: return 'none'
print("=== header row run fonts ===")
for c in range(len(tb.columns)):
    cell = tb.cell(0,c)
    tf = cell.text_frame
    runs = tf.paragraphs[0].runs
    txt = cell.text
    col = hx(runs[0].font.color) if runs and runs[0].font.color and runs[0].font.color.type is not None else '?'
    sz = runs[0].font.size.pt if runs and runs[0].font.size else '?'
    print(f"col{c} '{txt}' font sz={sz} bold={runs[0].font.bold} color={col} name={runs[0].font.name}")
print("\n=== a data cell (R1 col5 达人出单率) run detail ===")
cell = tb.cell(1,5)
for p in cell.text_frame.paragraphs:
    for r in p.runs:
        print(" run:", repr(r.text), "sz=", r.font.size.pt if r.font.size else '?', "bold=", r.font.bold, "color=", hx(r.font.color) if r.font.color and r.font.color.type is not None else '?')
print("\n=== cell fills (header R0 + data R1) ===")
for (rr,cc) in [(0,0),(1,5),(1,6),(5,4)]:
    cell = tb.cell(rr,cc)
    try:
        fill = cell.fill
        print(f"R{rr}C{cc} fill.type={fill.type}")
    except Exception as e:
        print(f"R{rr}C{cc} fill err {e}")
print("\n=== chart [10] data ===")
for sh in s.shapes:
    if sh.has_chart:
        ch = sh.chart
        print("type", ch.chart_type, "cats", list(ch.plots[0].categories))
        for se in ch.plots[0].series:
            print(" ", se.name, [round(v) for v in se.values])
