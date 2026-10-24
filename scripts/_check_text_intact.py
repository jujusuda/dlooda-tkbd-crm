# -*- coding: utf-8 -*-
from pptx import Presentation
SRC = r'E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-09\8be9710a0a0dd58ff4cb419316577cc7\8月余慧玲8月月度汇报(3).pptx'
OUT = r'C:\Users\LENOVO\Desktop\8月余慧玲8月月度汇报_商务淡粉版.pptx'

def dump(path):
    p = Presentation(path)
    out = []
    for s in p.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = ''.join(r.text for r in para.runs)
                    if t.strip(): out.append(t)
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        t = c.text
                        if t.strip(): out.append(t)
            if sh.has_chart:
                try:
                    for ser in sh.chart.series:
                        out.append('SER:%s' % ser.name)
                        # data values
                        try:
                            vals = list(ser.values)
                            out.append('VALS:%s' % (vals if len(vals) < 20 else 'len%d' % len(vals)))
                        except Exception:
                            pass
                except Exception:
                    pass
    return out

a = dump(SRC); b = dump(OUT)
print('SRC text/chart items:', len(a), '| OUT:', len(b))
diff = 0
for i, (x, y) in enumerate(zip(a, b)):
    if x != y:
        diff += 1
        if diff <= 20:
            print('DIFF@%d\n  SRC=%r\n  OUT=%r' % (i, x, y))
print('Total differing items:', diff)
print('Identical length & content:', a == b)
