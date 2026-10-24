# -*- coding: utf-8 -*-
"""P9（2178 + 2190｜同款同打法）整改：基于 v9_4（已含 P5+P8 修复），仅重建第 9 页（slide idx 8）。
- 标题/副标题：共性视角（同款同打法），副标题不出现具体数字
- 两张折线趋势图（达人出单率 / 视频出单率），各两条线 2190/2178，7-9 月，逐点标数
- 趋势一致性分析（先结论后分析）
- 共同原因（达人/产品/季节 三角度）
- 单独原因（2190 爆单带动 / 2178 无爆单）
- 落点策略（共性 + 差异微调，去重复）
配色与字号沿用 v9 规范。折线图标签用 ABOVE/BELOW（t/b），禁用 outEnd/inEnd 以免损坏。
产物 Desktop/余慧玲9月月度汇报_v9_5.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE

INK   = RGBColor.from_string('1F1F1F')   # 标题/主文字（近黑）
BODY  = RGBColor.from_string('3D3D3D')   # 正文
MUTE  = RGBColor.from_string('9A9A9A')   # 注释/辅助
ACC   = RGBColor.from_string('B06579')   # 玫红（大数字/强调）
ACC_L = RGBColor.from_string('C97F95')   # 柔雾玫瑰（装饰/副线）
CARD  = RGBColor.from_string('F9F5F6')   # 极浅暖粉灰卡片底
LINE  = RGBColor.from_string('EAE2E5')   # 浅分隔线
GRAY_D= RGBColor.from_string('5A5A5A')   # 图表深灰（2178 线）
WHITE = RGBColor.from_string('FFFFFF')
FONT  = '微软雅黑'

SRC = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_4.pptx'
OUT = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_5.pptx'

prs = Presentation(SRC)
s = prs.slides[8]  # P9

# ---------- 删除 P9 全部现有 shape ----------
for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)

def rect(x, y, w, h, fill, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp

def add_text(x, y, w, h, paras, anchor=None, align=None, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    if anchor is not None: tf.vertical_anchor = anchor
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if align is not None: p.alignment = align
        for (t, sz, bold, color) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = FONT
    return tb

# ---------- 标题区 ----------
rect(0.6, 0.34, 0.1, 0.74, ACC)                                  # 竖 accent
add_text(0.84, 0.30, 11.7, 0.66,
         [[('2178 + 2190｜同款同打法：少寄优寄奏效', 24, True, INK)]], anchor=MSO_ANCHOR.MIDDLE)
add_text(9.3, 0.42, 3.43, 0.56,
         [[('PART 02 · 各款 SKU 分析', 12, False, MUTE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
add_text(0.84, 1.02, 11.8, 0.34,
         [[('两款款式相同、趋势同步；上月同用“少寄优寄”，达人出单率与视频出单率双双回升', 15, False, BODY)]])
rect(0.6, 1.44, 12.13, Pt(1)/914400, LINE)                       # 分隔线

# ---------- 折线趋势图 ----------
def add_line_chart(x, y, w, h, cats, series, colors, vmax):
    cd = CategoryChartData(); cd.categories = cats
    for (nm, vals) in series: cd.add_series(nm, vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(10); ch.legend.font.name = FONT
    pl = ch.plots[0]
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.show_value = True; dl.number_format = '0.00"%"'; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.ABOVE
    dl.font.size = Pt(9); dl.font.bold = False; dl.font.name = FONT; dl.font.color.rgb = INK
    for i, c in enumerate(colors):
        se = pl.series[i]
        se.format.line.color.rgb = c
        se.format.line.width = Pt(2.25)
        se.marker.style = XL_MARKER_STYLE.CIRCLE
        se.marker.size = 6
        se.marker.format.fill.solid(); se.marker.format.fill.fore_color.rgb = c
        se.marker.format.line.color.rgb = WHITE; se.marker.format.line.width = Pt(1)
        sd = se.data_labels
        sd.show_value = True
        sd.number_format = '0.00"%"'; sd.number_format_is_linked = False
        # 2190 置于点上方（ABOVE），2178 置于点下方（BELOW），竖直错开避免叠字
        sd.position = XL_LABEL_POSITION.BELOW if i == 1 else XL_LABEL_POSITION.ABOVE
        sd.font.size = Pt(9)
        sd.font.bold = False; sd.font.name = FONT; sd.font.color.rgb = INK
    va = ch.value_axis
    va.minimum_scale = 0; va.maximum_scale = vmax
    va.major_unit = 10
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.major_gridlines.format.line.width = Pt(0.5)
    va.tick_labels.font.size = Pt(9); va.tick_labels.font.name = FONT
    va.tick_labels.font.color.rgb = MUTE; va.tick_labels.number_format = '0"%"'
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(11); ca.tick_labels.font.name = FONT; ca.tick_labels.font.color.rgb = GRAY_D
    ca.format.line.color.rgb = LINE
    return ch

# 图一：达人出单率
add_text(0.6, 1.60, 5.84, 0.30, [[('达人出单率（7–9 月）', 15, True, INK)]])
add_line_chart(0.6, 1.92, 5.84, 2.40,
               ['7月', '8月', '9月'],
               [('2190', [20.83, 7.50, 35.71]),
                ('2178', [16.00, 9.80, 17.95])],
               [ACC, GRAY_D], 40)
# 图二：视频出单率
add_text(6.89, 1.60, 5.84, 0.30, [[('视频出单率（7–9 月）', 15, True, INK)]])
add_line_chart(6.89, 1.92, 5.84, 2.40,
               ['7月', '8月', '9月'],
               [('2190', [19.23, 17.76, 23.36]),
                ('2178', [11.84, 12.62, 21.95])],
               [ACC, GRAY_D], 40)

# ---------- 趋势一致性分析（先结论后分析，整体带） ----------
CYB, CHB = 4.45, 0.97
rect(0.6, CYB, 12.13, CHB, CARD)
rect(0.6, CYB, 0.09, CHB, ACC)
add_text(0.95, CYB+0.06, 11.6, CHB-0.12,
         [[('结论　', 14, True, ACC),
           ('两款出单端与视频端走势均同步，共性成立。', 14, False, BODY)],
          [('达人出单率：两款同步「7 月→8 月降、8 月→9 月升」，方向一致；视频出单率：两款 9 月均升至 20% 以上，走势同步向上（2190 先抑后扬、2178 连续走高）。', 13, False, BODY)]],
         anchor=MSO_ANCHOR.MIDDLE)

# ---------- 共同原因（左）/ 单独原因（右） ----------
CARD_T, CARD_H = 5.52, 1.10
# 共同原因
rect(0.6, CARD_T, 5.84, CARD_H, CARD)
rect(0.6, CARD_T, 0.07, CARD_H, ACC)
add_text(0.82, CARD_T+0.06, 5.5, 0.30, [[('共同原因（三个角度）', 14, True, INK)]])
add_text(0.82, CARD_T+0.40, 5.5, 0.64,
         [[('达人｜小寄样、精选高潜达人（少寄优寄），出单效率提升', 12.5, False, BODY)],
          [('产品｜价格微降，承接力更强、更易转化', 12.5, False, BODY)],
          [('季节｜秋冬款推广期，长裤类目当季放量', 12.5, False, BODY)]])
# 单独原因
rect(6.89, CARD_T, 5.84, CARD_H, CARD)
rect(6.89, CARD_T, 0.07, CARD_H, ACC)
add_text(7.11, CARD_T+0.06, 5.5, 0.30, [[('单独原因（两款差异）', 14, True, INK)]])
add_text(7.11, CARD_T+0.40, 5.5, 0.64,
         [[('2190｜大爆单达人带动跟投，出单率大幅回升 ', 12.5, False, BODY), ('35.71%', 12.5, True, ACC), ('、量价双高', 12.5, False, BODY)],
          [('2178｜无爆单带动，回升幅度较小 ', 12.5, False, BODY), ('17.95%', 12.5, True, ACC), ('，订单量未完全跟上', 12.5, False, BODY)]])

# ---------- 落点策略（共性 + 差异微调） ----------
SYB, SHB = 6.74, 0.60
rect(0.6, SYB, 12.13, SHB, CARD)
rect(0.6, SYB, 0.09, SHB, ACC)
add_text(0.95, SYB, 11.65, SHB,
         [[('落点策略　', 13, True, ACC),
           ('① ', 13, True, INK), ('两款统一沿用“少寄优寄”（量少质高，已验证奏效）', 13, False, BODY),
           ('　② ', 13, True, INK), ('2190 借助爆单达人效应，持续扩量放大', 13, False, BODY),
           ('　③ ', 13, True, INK), ('2178 培育 / 引入爆单达人带动，补寄样与视频供给，兑现回升的出单率', 13, False, BODY)]],
         anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print('SAVED:', OUT, 'slides:', len(prs.slides._sldIdLst))
