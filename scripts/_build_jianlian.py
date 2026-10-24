# -*- coding: utf-8 -*-
"""达人建联页（第 15 页）重排：标题 → 背景 → 漏斗(左)+分品表(右) → 判断横幅 → 落点策略"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from lxml import etree
import copy

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报.pptx"
OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_达人建联修改.pptx"

FONT = "微软雅黑"
INK   = RGBColor(0x1F, 0x1F, 0x1F)
BODY  = RGBColor(0x3D, 0x3D, 0x3D)
MUTE  = RGBColor(0x9A, 0x9A, 0x9A)
ROSE  = RGBColor(0xB0, 0x65, 0x79)   # 玫红
ROSE_D= RGBColor(0x8E, 0x4A, 0x5F)   # 深玫红
ROSE_M= RGBColor(0xC9, 0x81, 0x96)
ROSE_L= RGBColor(0xD9, 0xA5, 0xB5)   # 浅豆沙粉
ROSE_T= RGBColor(0xE8, 0xC6, 0xD0)
PINK_BG = RGBColor(0xFB, 0xEE, 0xF2)
HEAD_BG = RGBColor(0xF3, 0xE4, 0xE9)
CARD_BG = RGBColor(0xF7, 0xF4, 0xF5)
SUM_BG  = RGBColor(0xF2, 0xF0, 0xF1)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
PINK_TX = RGBColor(0xFF, 0xE3, 0xEA)

ML = 0.60          # 左边距
CW = 12.13         # 内容宽度 (12.73 - 0.60)


def _ea(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn('a:latin'))
    ea = rPr.find(qn('a:ea'))
    cs = rPr.find(qn('a:cs'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        if cs is not None:
            cs.addprevious(ea)
        elif latin is not None:
            latin.addnext(ea)
        else:
            rPr.append(ea)
    ea.set('typeface', name)
    if cs is None:
        cs = rPr.makeelement(qn('a:cs'), {})
        ea.addnext(cs)
    cs.set('typeface', name)


def style(run, size=None, bold=None, color=None, name=FONT):
    _ea(run, name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb


def put(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        line_spacing=None, space_after=0):
    """runs: [[(text, size, bold, color), ...], ...] 每个子列表 = 一个段落"""
    tb = textbox(slide, x, y, w, h, anchor)
    tf = tb.text_frame
    for i, para_runs in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        for (t, sz, bd, col) in para_runs:
            r = p.add_run()
            r.text = t
            style(r, sz, bd, col)
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=None):
    from pptx.enum.shapes import MSO_SHAPE
    if radius is not None:
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(x), Inches(y), Inches(w), Inches(h))
        try:
            sh.adjustments[0] = min(0.5, radius / min(w, h))
        except Exception:
            pass
    else:
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                    Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    if sh.has_text_frame:
        sh.text_frame.text = ""
    return sh


def oval(slide, x, y, w, h, fill, text, size, color):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    style(r, size, True, color)
    return sh


def trapezoid(slide, cx, y0, y1, wtop, wbot, fill):
    fb = slide.shapes.build_freeform(Inches(cx - wtop), Inches(y0))
    fb.add_line_segments([
        (Inches(cx + wtop), Inches(y0)),
        (Inches(cx + wbot), Inches(y1)),
        (Inches(cx - wbot), Inches(y1)),
    ], close=True)
    sh = fb.convert_to_shape()
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


# ---------------------------------------------------------------- build
prs = Presentation(SRC)
slide = prs.slides[14]

for sh in list(slide.shapes):
    sh._element.getparent().remove(sh._element)

# ① 标题
rect(slide, ML, 0.34, 0.10, 0.74, fill=ROSE)
put(slide, 0.84, 0.28, 7.70, 0.62,
    [[("达人建联 · 主推款老达人二次合作", 28, True, INK)]],
    anchor=MSO_ANCHOR.MIDDLE)
put(slide, 8.60, 0.44, 4.13, 0.34,
    [[("PART 03 · 达人端", 13, False, MUTE)]], align=PP_ALIGN.RIGHT)

# ② 背景
put(slide, 0.84, 1.00, 11.89, 0.34,
    [[("主推款老品达人趋饱和，转向激活出单老达人二次合作，出单确定性更高", 15, False, BODY)]],
    anchor=MSO_ANCHOR.MIDDLE)
rect(slide, ML, 1.42, CW, 0.012, fill=RGBColor(0xE8, 0xE2, 0xE4))

# ③ 漏斗（左）
FX0, FY0 = ML, 1.62
FCX = FX0 + 2.55
trapezoid(slide, FCX, 1.62, 2.40, 2.45, 2.20, ROSE)
trapezoid(slide, FCX, 2.73, 3.51, 2.00, 1.78, ROSE_M)
trapezoid(slide, FCX, 3.84, 4.62, 1.58, 1.35, ROSE_T)

put(slide, FX0, 1.72, 5.10, 0.36,
    [[("27", 30, True, WHITE), ("  成功建联", 12.5, False, RGBColor(0xF6, 0xE4, 0xE9))]],
    align=PP_ALIGN.CENTER)
put(slide, FX0, 2.10, 5.10, 0.28,
    [[("出单老达人同意再合作", 11, False, RGBColor(0xF6, 0xE4, 0xE9))]],
    align=PP_ALIGN.CENTER)

put(slide, FX0, 2.44, 5.10, 0.28,
    [[("▼  建联 → 发视频 ", 12.5, False, MUTE), ("74.1%", 13.5, True, ROSE)]],
    align=PP_ALIGN.CENTER)

put(slide, FX0, 2.83, 5.10, 0.36,
    [[("20", 30, True, WHITE), ("  已发视频", 12.5, False, RGBColor(0xF9, 0xEE, 0xF2))]],
    align=PP_ALIGN.CENTER)

put(slide, FX0, 3.55, 5.10, 0.28,
    [[("▼  发视频 → 出单 ", 12.5, False, MUTE), ("85.0%", 13.5, True, ROSE)]],
    align=PP_ALIGN.CENTER)

put(slide, FX0, 3.94, 5.10, 0.36,
    [[("17", 30, True, ROSE_D), ("  再出单", 12.5, False, RGBColor(0x7A, 0x3B, 0x50))]],
    align=PP_ALIGN.CENTER)

put(slide, FX0, 4.72, 5.10, 0.52,
    [[("说明：本页「建联」指主动向已出单老达人发起二次合作、", 10.5, False, MUTE)],
     [("成功同意合作的人数，不是新拓达人。", 10.5, False, MUTE)]],
    line_spacing=1.15)

# ③ 分品明细表（右）
TX, TY = 6.05, 1.62
rows, cols = 6, 5
shp = slide.shapes.add_table(rows, cols, Inches(TX), Inches(TY),
                             Inches(6.68), Inches(3.00))
tbl = shp.table
widths = [1.30, 0.90, 1.15, 1.00, 2.33]
for i, w in enumerate(widths):
    tbl.columns[i].width = Inches(w)
for i in range(rows):
    tbl.rows[i].height = Inches(0.50)

data = [
    ["sku", "建联", "已发视频", "已出单", "发视频→出单率"],
    ["2178", "9", "7", "7", "100% ★"],
    ["2190", "13", "9", "6", "67%"],
    ["28173", "3", "2", "2", "100%"],
    ["2189", "2", "2", "2", "100%"],
    ["合计", "27", "20", "17", "85%"],
]
for ri, row in enumerate(data):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.04)
        cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = val
        if ri == 0:
            style(r, 12, True, ROSE)
            cell.fill.solid(); cell.fill.fore_color.rgb = HEAD_BG
        elif ri == 1:
            style(r, 12.5, True, ROSE)
            cell.fill.solid(); cell.fill.fore_color.rgb = PINK_BG
        elif ri == 5:
            style(r, 12.5, True, INK)
            cell.fill.solid(); cell.fill.fore_color.rgb = SUM_BG
        else:
            style(r, 12.5, False, BODY)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 == 0 else RGBColor(0xFA, 0xF8, 0xF9)
        # 发视频→出单率 列：非 2178 行也用玫红强调百分比
        if ci == 4 and ri not in (0, 1):
            style(r, 12.5, True, ROSE)

# 2178 高亮说明
rect(slide, TX, 4.72, 6.68, 0.46, fill=PINK_BG, radius=0.10)
put(slide, TX + 0.18, 4.72, 6.32, 0.46,
    [[("★ 2178：7 位已发视频全部再出单", 12.5, True, ROSE)]],
    anchor=MSO_ANCHOR.MIDDLE)

# ④ 判断横幅
BY, BH = 5.28, 0.80
rect(slide, ML, BY, CW, BH, fill=ROSE, radius=0.08)
rect(slide, ML, BY, 0.075, BH, fill=ROSE_D)
put(slide, ML + 0.28, BY, CW - 0.56, BH,
    [[("判断　", 14, True, RGBColor(0xF0, 0xCB, 0xD5)),
      ("老达人二次合作出单确定性高（发视频→出单 ", 17, True, WHITE),
      ("85%", 18, True, PINK_TX),
      ("），", 17, True, WHITE),
      ("2178 达 100%", 18, True, PINK_TX),
      ("，是稳定出单的高效路径", 17, True, WHITE)]],
    anchor=MSO_ANCHOR.MIDDLE)

# ⑤ 落点策略（三块并列）
AY, AH = 6.34, 0.76
GAP = 0.22
BW = (CW - 2 * GAP) / 3
acts = [
    ("1", "扩大出单老达人的二次合作覆盖面"),
    ("2", "未合作过的达人，付费合作补充"),
    ("3", "巩固 2178 等高确定款的老达人池"),
]
for i, (num, txt) in enumerate(acts):
    x = ML + i * (BW + GAP)
    rect(slide, x, AY, BW, AH, fill=CARD_BG, radius=0.10)
    rect(slide, x, AY, 0.055, AH, fill=ROSE)
    oval(slide, x + 0.20, AY + AH / 2 - 0.14, 0.28, 0.28, ROSE, num, 11, WHITE)
    put(slide, x + 0.58, AY, BW - 0.72, AH,
        [[(txt, 13, False, BODY)]], anchor=MSO_ANCHOR.MIDDLE)

# 落点策略 小标签
put(slide, ML, 6.12, 3.00, 0.22,
    [[("落点策略", 12, True, MUTE)]])

prs.save(OUT)
print("saved:", OUT)
