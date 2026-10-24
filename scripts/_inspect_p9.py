import sys
from pptx import Presentation
from pptx.util import Emu
import copy

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报_v9_1(7).pptx"

prs = Presentation(SRC)
print("slide count:", len(prs.slides))
print("slide width:", prs.slide_width, "height:", prs.slide_height)
# slide 9 -> index 8
slide = prs.slides[8]
print("="*80)
print("SLIDE 9 shapes:")
for i, sh in enumerate(slide.shapes):
    print(f"\n--- shape[{i}] name={sh.name!r} type={sh.shape_type}")
    try:
        print(f"    L={sh.left} T={sh.top} W={sh.width} H={sh.height}")
    except Exception as e:
        print("    pos err", e)
    if sh.has_chart:
        ch = sh.chart
        print("    CHART type=", ch.chart_type)
        try:
            print("    categories=", list(ch.plots[0].categories))
        except Exception as e:
            print("    cat err", e)
        for s in ch.series:
            print("    SERIES name=", repr(s.name), "vals=", list(s.values))
        # axes
        try:
            for ax in ch.axes:
                print("    AXIS", ax.axis_type, "id?", ax.axId if hasattr(ax,'axId') else '?')
        except Exception as e:
            print("    axes err", e)
    elif sh.has_table:
        tbl = sh.table
        print("    TABLE rows=", len(tbl.rows), "cols=", len(tbl.columns))
        for r in range(len(tbl.rows)):
            row = [tbl.cell(r,c).text for c in range(len(tbl.columns))]
            print("    ROW", r, row)
    elif sh.has_text_frame:
        txt = sh.text_frame.text
        print("    TEXT:", repr(txt[:300]))
