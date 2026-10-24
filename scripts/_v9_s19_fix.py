# -*- coding: utf-8 -*-
"""s19（下月行动）重排：表格铺满+动作分行、下半区精简为一张拆分卡、删小图"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INK  = RGBColor.from_string('1F1F1F')
BODY = RGBColor.from_string('3D3D3D')
MUTE = RGBColor.from_string('9A9A9A')
ACC  = RGBColor.from_string('B06579')
FONT = '微软雅黑'

F = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9.pptx'
prs = Presentation(F)
s = prs.slides[18]  # s19

# ---------- 1. 删掉下半区旧内容（5.0-5.2 之间的矩形/文本/图表） ----------
for sh in list(s.shapes):
    if sh.left is None: continue
    T = sh.top/914400
    if 4.95 < T < 5.20:
        sh._element.getparent().remove(sh._element)

# ---------- 2. 改造动作清单表：铺满全宽 + ①②③ 分行 ----------
tbl = None
for sh in s.shapes:
    if sh.has_table:
        tbl = sh.table; tblsh = sh; break
assert tbl is not None

# 列宽铺满：方向 1.75 / 动作 10.38
tbl.columns[0].width = Inches(1.75)
tbl.columns[1].width = Inches(10.38)
# 行高：表头 0.42，数据行 0.62（三条各占一行）
tbl.rows[0].height = Inches(0.42)
for ri in range(1, 5):
    tbl.rows[ri].height = Inches(0.62)

ACTIONS = {
    1: ['① 2190 保持并适度扩寄样', '② 2189 补视频与达人供给', '③ 2178 加寄样兑现出单率'],
    2: ['① 28173 继续放大寄样', '② 稳住出单率 30%-34%', '③ 拉回月度订单'],
    3: ['① 2881 加投 + 履约跟进', '② 2L06 保持健康节奏', '③ 付费达人合作推进'],
    4: ['① 28102 停止寄样', '② 止损退出', '③ 资源腾给新测品'],
}
for ri, lines in ACTIONS.items():
    c = tbl.cell(ri, 1)
    tf = c.text_frame
    tf.clear()
    tf.word_wrap = True
    c.vertical_anchor = MSO_ANCHOR.MIDDLE
    c.margin_left = Inches(0.14); c.margin_right = Inches(0.06)
    c.margin_top = Inches(0.02); c.margin_bottom = Inches(0.02)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(1); p.line_spacing = 1.0
        r = p.add_run(); r.text = ln
        r.font.size = Pt(12.5); r.font.name = FONT; r.font.color.rgb = BODY
# 方向列垂直居中
for ri in range(1, 5):
    tbl.cell(ri, 0).vertical_anchor = MSO_ANCHOR.MIDDLE
# 表头也居中
for ci in range(2):
    tbl.cell(0, ci).vertical_anchor = MSO_ANCHOR.MIDDLE

# ---------- 3. 下半区：寄样拆分轻量卡 ----------
card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.60), Inches(5.62), Inches(12.13), Inches(0.80))
card.fill.solid(); card.fill.fore_color.rgb = RGBColor.from_string('F9F5F6')
card.line.fill.background(); card.shadow.inherit = False
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.60), Inches(5.62), Inches(0.09), Inches(0.80))
bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string('C97F95')
bar.line.fill.background(); bar.shadow.inherit = False

tb = s.shapes.add_textbox(Inches(0.88), Inches(5.70), Inches(11.60), Inches(0.66))
tf = tb.text_frame
tf.word_wrap = True
tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
# 行 1：拆分明细
p1 = tf.paragraphs[0]
p1.alignment = PP_ALIGN.LEFT; p1.space_after = Pt(3); p1.line_spacing = 1.05
r = p1.add_run(); r.text = '10 月寄样拆分　'
r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = ACC; r.font.name = FONT
r = p1.add_run(); r.text = '爆品 190 件：2178=40 · 2190=40 · 28173=70 · 2189=40　｜　测品 350 件：2A16=100 · 2F08=100 · 2A51=50 · 2A50=50 · 2H43=50'
r.font.size = Pt(13); r.font.bold = False; r.font.color.rgb = BODY; r.font.name = FONT
# 行 2：口径注
p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.LEFT; p2.line_spacing = 1.0
r = p2.add_run(); r.text = '※ 目标口径同「本月新增寄样记录」，与概览表寄样端口径不同'
r.font.size = Pt(10.5); r.font.bold = False; r.font.color.rgb = MUTE; r.font.name = FONT

OUT = r'C:\Users\LENOVO\Desktop\余慧玲9月月度汇报_v9_1.pptx'
prs.save(OUT)
print('SAVED s19 rework ->', OUT)
