# -*- coding: utf-8 -*-
import zipfile, posixpath
from pptx import Presentation
from pptx.util import Emu

SRC = r"E:\xwechat_files\wxid_od6e7r1klxs422_54c2\temp\RWTemp\2026-10\8be9710a0a0dd58ff4cb419316577cc7\余慧玲9月月度汇报.pptx"
OUT = r"C:\Users\LENOVO\WorkBuddy\2026-08-02-13-13-32\余慧玲9月月度汇报_达人建联修改.pptx"
IDX = 14  # 第 15 页

ok = True
def chk(cond, msg):
    global ok
    print(("  PASS  " if cond else "  FAIL  ") + msg)
    if not cond:
        ok = False

z = zipfile.ZipFile(OUT)
chk(z.testzip() is None, "ZIP 完整性")
# rels 全部可解析
bad = []
for n in z.namelist():
    if n.endswith('.rels'):
        root = __import__('lxml.etree', fromlist=['etree']).fromstring(z.read(n))
        d = posixpath.dirname(n)
        base = posixpath.dirname(d) if posixpath.basename(d) == '_rels' else d
        for rel in root:
            if rel.get('TargetMode') == 'External':
                continue
            p = posixpath.normpath(posixpath.join(base, rel.get('Target')))
            if p not in z.namelist():
                bad.append((n, rel.get('Target')))
chk(not bad, f"关系链接全部有效 ({len(bad)} 条失效)")

prs = Presentation(OUT)
chk(len(prs.slides) == 17, f"共 17 页（实际 {len(prs.slides)}）")
s = prs.slides[IDX]

# 未改动页与源逐字节一致
sz = zipfile.ZipFile(SRC)
diff = [n for n in sz.namelist()
        if n in z.namelist() and n not in
        ("[Content_Types].xml", f"ppt/slides/slide{IDX+1}.xml",
         f"ppt/slides/_rels/slide{IDX+1}.xml.rels")
        and sz.read(n) != z.read(n)]
chk(not diff, f"其余部件与源逐字节一致（差异 {diff}）")

# 全页文本
alltext = []
for sh in s.shapes:
    if sh.has_text_frame:
        alltext.append(sh.text_frame.text)
    if sh.has_table:
        for r in sh.table.rows:
            for c in r.cells:
                alltext.append(c.text)
T = "\n".join(alltext)

checks = [
    ("达人建联 · 主推款老达人二次合作", "标题"),
    ("主推款老品达人趋饱和，转向激活出单老达人二次合作，出单确定性更高", "背景"),
    ("74.1%", "漏斗 建联→发视频 74.1%"),
    ("85.0%", "漏斗 发视频→出单 85.0%"),
    ("老达人二次合作出单确定性高", "判断横幅"),
    ("扩大出单老达人的二次合作覆盖面", "策略①"),
    ("未合作过的达人，付费合作补充", "策略②"),
    ("巩固 2178 等高确定款的老达人池", "策略③"),
    ("发视频→出单率", "表格新增列"),
    ("7 位已发视频全部再出单", "2178 高亮说明"),
]
for t, lab in checks:
    chk(t in T, lab)

for gone in ["流失", "催发视频", "高意愿达人", "找竞品", "付��达人合作落地"]:
    chk(gone not in T, f"已删除「{gone}」")

# 数字：27 / 20 / 17 都在
for num in ["27", "20", "17"]:
    chk(num in T, f"含数字 {num}")

# 表格内容
tbl = [sh for sh in s.shapes if sh.has_table][0].table
grid = [[c.text for c in row.cells] for row in tbl.rows]
expect = [
    ["sku", "建联", "已发视频", "已出单", "发视频→出单率"],
    ["2178", "9", "7", "7", "100% ★"],
    ["2190", "13", "9", "6", "67%"],
    ["28173", "3", "2", "2", "100%"],
    ["2189", "2", "2", "2", "100%"],
    ["合计", "27", "20", "17", "85%"],
]
chk(grid == expect, f"表格数据正确 -> {grid}")
# 合计校验
chk(sum(int(r[1]) for r in grid[1:-1]) == 27, "建联合计=27")
chk(sum(int(r[2]) for r in grid[1:-1]) == 20, "已发视频合计=20")
chk(sum(int(r[3]) for r in grid[1:-1]) == 17, "已出单合计=17")

# 2178 行玫红
c = tbl.cell(1, 4)
col = c.text_frame.paragraphs[0].runs[0].font.color.rgb
chk(str(col) == "B06579", f"2178 出单率玫红（{col}）")
chk(tbl.cell(1, 0).text_frame.paragraphs[0].runs[0].font.bold, "2178 行加粗")

# 图片数 = 0
pics = [sh for sh in s.shapes if sh.shape_type == 13]
chk(not pics, f"本页无贴图（{len(pics)}）")

# 边界 & 重叠
SW, SH = prs.slide_width, prs.slide_height
oob = []
for sh in s.shapes:
    if sh.left < 0 or sh.top < 0 or sh.left + sh.width > SW or sh.top + sh.height > SH:
        oob.append(sh.name)
chk(not oob, f"无越界（{oob}）")

def box(sh):
    return (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)
def ov(a, b):
    x = min(a[2], b[2]) - max(a[0], b[0])
    y = min(a[3], b[3]) - max(a[1], b[1])
    return x > 0 and y > 0

shapes = [sh for sh in s.shapes]
conf = []
for i in range(len(shapes)):
    for j in range(i + 1, len(shapes)):
        a, b = shapes[i], shapes[j]
        # 忽略无文本的纯装饰底板与叠在其上的文本框
        at = a.has_text_frame and a.text_frame.text.strip()
        bt = b.has_text_frame and b.text_frame.text.strip()
        if not (at and bt):
            continue
        if ov(box(a), box(b)):
            conf.append((a.name, a.text_frame.text[:14], b.name, b.text_frame.text[:14]))
chk(not conf, f"文本块无重叠（{conf}）")

print("\nRESULT:", "ALL PASS" if ok else "HAS FAILURE")
